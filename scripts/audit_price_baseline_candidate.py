"""Audit a candidate workbook; publication is an explicit opt-in."""

from __future__ import annotations

import argparse
import copy
import json
import math
import os
import re
import sys
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, cast

from calc_quote_price_draft import (  # type: ignore[import-not-found]
    SUCCESSOR_COMPONENT_PRICE_MAPPINGS,
    build_governed_mapping_snapshot,
)
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from price_baseline_contract import (  # type: ignore[import-not-found]
    BOOTSTRAP_AUDIT_SCHEMA,
    CANDIDATE_AUDIT_SCHEMA,
    DEFAULT_ACTIVE_SELECTOR,
    SUCCESSOR,
    BaselineContractError,
    PriceBaseline,
    audit_artifact_path,
    canonical_json_bytes,
    load_json_bytes,
    normalize_kzt_literal,
    resolve_active_price_baseline,
    sha256_bytes,
    sha256_file,
    validate_manifest,
)

AUDIT_SCHEMA = CANDIDATE_AUDIT_SCHEMA
BOOTSTRAP_INTENT = "FIRST_CHAIN_BOOTSTRAP"
# Bounded source-table footprints, not calculator mappings. The side tables
# use different columns and include non-price dimension/header blocks.
MATERIAL_REGIONS: dict[str, tuple[int, int]] = {
    "ВРУ250А": (2, 15),
    "ВРУ630А": (2, 15),
    "ВРУ400А": (2, 15),
    "ВРУ-расп.": (2, 10),
    "ВРУ-ВА": (2, 29),
    "КРН": (2, 29),
    "ЩР": (2, 52),
    "АВР Г-Г": (2, 41),
    "АВР Г-Д": (2, 47),
    "АВР-Г-Г-Д": (2, 45),
    "ШРС": (2, 21),
    "ЩЭ": (2, 19),
    "БАУО рассч": (6, 14),
    "Я5111": (2, 20),
    "ЯУО9601": (2, 24),
    "ЯУО9602": (2, 23),
}
WORK_SHEETS = frozenset({"ВРУ-ВА", "КРН", "ЩР", "ЩЭ"})
CABINET_REGIONS: dict[str, tuple[int, int, tuple[tuple[int, int], ...]]] = {
    "ВРУ250А": (10, 11, ((2, 4), (10, 11))),
    "ВРУ630А": (10, 11, ((2, 4),)),
    "ВРУ400А": (10, 11, ((2, 4), (10, 11))),
    "ВРУ-расп.": (10, 11, ((2, 4), (8, 9))),
    "ВРУ-ВА": (12, 13, ((5, 6),)),
    "КРН": (12, 13, ((3, 24),)),
    "ЩР": (12, 13, ((3, 14),)),
    "АВР Г-Г": (11, 12, ((2, 7), (12, 21), (27, 31), (33, 37))),
    "АВР Г-Д": (11, 12, ((2, 7), (10, 19))),
    "АВР-Г-Г-Д": (10, 11, ((2, 8), (11, 20), (25, 29), (31, 35))),
    "ШРС": (10, 11, ((3, 4), (6, 8))),
    "Я5111": (10, 11, ((3, 4), (7, 8))),
    "ЯУО9601": (10, 11, ((3, 6),)),
    "ЯУО9602": (10, 11, ((3, 6),)),
}
BUSBAR_LABEL = re.compile(r"\d+\s*[xх×]\s*\d+\s*мм.*шина\s+АЛ", re.I)


class CandidateAuditError(RuntimeError):
    """The candidate cannot be audited safely."""


@dataclass(frozen=True)
class WorkbookInspection:
    sheet_order: tuple[str, ...]
    price_names: tuple[dict[str, Any], ...]
    formula_cells: tuple[str, ...]
    invalid_price_cells: tuple[str, ...]
    duplicate_names: tuple[str, ...]
    structural_fingerprint: str


def normalize_label(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    normalized = " ".join(value.replace("\xa0", " ").split())
    return normalized or None


def positive_int(value: Any) -> int | None:
    return cast(int | None, normalize_kzt_literal(value))


def _separate_busbar_labor_row(
    sheet: str, row: int, value: Callable[[int, int], Any]
) -> bool:
    if sheet != "ЩР":
        return False
    label = normalize_label(value(row, 1))
    if label is None or BUSBAR_LABEL.search(label) is None:
        return False
    next_row = row + 1
    upper_row = MATERIAL_REGIONS[sheet][1]
    while next_row <= upper_row:
        next_label = normalize_label(value(next_row, 1))
        if next_label is None or BUSBAR_LABEL.search(next_label) is None:
            break
        next_row += 1
    return (
        normalize_label(value(next_row, 1)) == "Шина за работу"
        and value(next_row, 2) is None
        and normalize_kzt_literal(value(next_row, 3)) is not None
        and value(next_row, 6) == f"=D{next_row}*C{next_row}"
        and value(row, 6) == f"=D{row}*C{row}"
    )


def inspected_price_value(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return {"type": "float", "text": repr(value)}
    if value is None or isinstance(value, bool | int | float | str):
        return value
    return {"type": type(value).__name__, "text": str(value)}


def inspect_workbook(path: Path) -> WorkbookInspection:
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (OSError, ValueError) as exc:
        raise CandidateAuditError("candidate workbook could not be opened") from exc
    names: list[dict[str, Any]] = []
    formulas: list[str] = []
    invalid_prices: list[str] = []
    duplicates: list[str] = []
    try:
        sheet_order = tuple(workbook.sheetnames)
        for worksheet in workbook.worksheets:
            sheet = worksheet.title
            last_material = MATERIAL_REGIONS.get(sheet, (0, 0))[1]
            last_cabinet = max(
                (end for _, end in CABINET_REGIONS.get(sheet, (0, 0, ()))[2]),
                default=0,
            )
            last_row = max(last_material, last_cabinet)
            if last_row == 0:
                continue
            row_values = {
                row: tuple(cell.value for cell in cells)
                for row, cells in enumerate(
                    worksheet.iter_rows(min_row=1, max_row=last_row, max_col=13),
                    start=1,
                )
            }

            def cell_value(
                row: int,
                column: int,
                *,
                _rows: dict[int, tuple[Any, ...]] = row_values,
            ) -> Any:
                return _rows.get(row, (None,) * 13)[column - 1]

            regions: list[tuple[str, int, int, tuple[int, ...]]] = []
            if sheet in MATERIAL_REGIONS:
                first, last = MATERIAL_REGIONS[sheet]
                for row in range(first, last + 1):
                    label = normalize_label(cell_value(row, 1))
                    material = cell_value(row, 2)
                    work = cell_value(row, 3) if sheet in WORK_SHEETS else None
                    if label is None or (material is None and work is None):
                        continue
                    work_only = (
                        sheet == "ЩР" and label == "Шина за работу" and material is None
                    )
                    columns: tuple[int, ...] = () if work_only else (2,)
                    if sheet in WORK_SHEETS and (
                        work is not None or sheet in {"КРН", "ЩР", "ЩЭ"}
                    ):
                        columns += (3,)
                    regions.append(("component", row, 1, columns))
            if sheet in CABINET_REGIONS:
                label_col, price_col, intervals = CABINET_REGIONS[sheet]
                for first, last in intervals:
                    for row in range(first, last + 1):
                        label = normalize_label(cell_value(row, label_col))
                        price = cell_value(row, price_col)
                        if label is not None and (price is not None or sheet != "КРН"):
                            regions.append(("cabinet", row, label_col, (price_col,)))
            for kind, row, label_column, price_columns in regions:
                label = normalize_label(cell_value(row, label_column))
                assert label is not None
                cells = [
                    f"{sheet}!{get_column_letter(column)}{row}"
                    for column in price_columns
                ]
                raw_prices = [cell_value(row, column) for column in price_columns]
                for column, cell, raw_value in zip(
                    price_columns, cells, raw_prices, strict=True
                ):
                    if isinstance(raw_value, str) and raw_value.startswith("="):
                        formulas.append(cell)
                        continue
                    allow_zero = (
                        kind == "component"
                        and column == 3
                        and raw_value == 0
                        and not isinstance(raw_value, bool)
                        and _separate_busbar_labor_row(sheet, row, cell_value)
                    )
                    if normalize_kzt_literal(raw_value, allow_zero=allow_zero) is None:
                        rendered = json.dumps(
                            inspected_price_value(raw_value),
                            ensure_ascii=False,
                            sort_keys=True,
                            separators=(",", ":"),
                        )
                        invalid_prices.append(
                            f"{cell}: required integer KZT literal price "
                            f"is missing or invalid: {rendered}"
                        )
                names.append(
                    {
                        "kind": kind,
                        "sheet": sheet,
                        "row": row,
                        "label": label,
                        "label_cell": f"{sheet}!{get_column_letter(label_column)}{row}",
                        "price_cells": cells,
                        "price_values": {
                            cell: inspected_price_value(value)
                            for cell, value in zip(cells, raw_prices, strict=True)
                        },
                    }
                )
        counts = Counter(
            (entry["kind"], entry["sheet"], entry["label"])
            for entry in names
            if entry["sheet"] in {"КРН", "ЩР"}
        )
        duplicates = [
            f"{kind}:{sheet}:{label}"
            for (kind, sheet, label), count in sorted(counts.items())
            if count > 1
        ]
        structural_payload = {
            "sheet_order": list(sheet_order),
            "price_name_coordinates": [
                {
                    "kind": entry["kind"],
                    "sheet": entry["sheet"],
                    "row": entry["row"],
                    "label": entry["label"],
                    "label_cell": entry["label_cell"],
                    "price_cells": entry["price_cells"],
                }
                for entry in names
            ],
        }
        fingerprint = sha256_bytes(canonical_json_bytes(structural_payload))
        return WorkbookInspection(
            sheet_order=sheet_order,
            price_names=tuple(names),
            formula_cells=tuple(sorted(formulas)),
            invalid_price_cells=tuple(sorted(invalid_prices)),
            duplicate_names=tuple(duplicates),
            structural_fingerprint=fingerprint,
        )
    finally:
        workbook.close()


def _price_row_key(entry: Mapping[str, Any]) -> tuple[Any, ...]:
    return (
        entry["kind"],
        entry["sheet"],
        entry["row"],
        entry["label"],
        entry["label_cell"],
        tuple(cast(list[str], entry["price_cells"])),
    )


def _governed_sources(
    snapshot: Sequence[Mapping[str, Any]],
) -> dict[tuple[Any, ...], list[str]]:
    governed: dict[tuple[Any, ...], list[str]] = {}
    for entry in snapshot:
        source = cast(Mapping[str, Any], entry["source"])
        mapping_kind = cast(str, entry["mapping_kind"])
        expected_label = normalize_label(
            cast(Mapping[str, Any], entry["identity"])["expected_label"]
        )
        key = (
            mapping_kind.split("_", maxsplit=1)[0],
            source["sheet"],
            source["row"],
            expected_label,
            source["label_cell"],
            tuple(cast(list[str], source["price_cells"])),
        )
        governed.setdefault(key, []).append(cast(str, entry["mapping_id"]))
    return {key: sorted(values) for key, values in governed.items()}


def diff_all_price_values(
    predecessor: WorkbookInspection,
    candidate: WorkbookInspection,
    governed_snapshot: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    candidate_by_key = {_price_row_key(entry): entry for entry in candidate.price_names}
    governed = _governed_sources(governed_snapshot)
    changes: list[dict[str, Any]] = []
    for predecessor_entry in predecessor.price_names:
        key = _price_row_key(predecessor_entry)
        candidate_entry = candidate_by_key.get(key)
        if candidate_entry is None:
            continue
        before = cast(Mapping[str, Any], predecessor_entry["price_values"])
        after = cast(Mapping[str, Any], candidate_entry["price_values"])
        if dict(before) == dict(after):
            continue
        mapping_ids = governed.get(key, [])
        changes.append(
            {
                "kind": predecessor_entry["kind"],
                "sheet": predecessor_entry["sheet"],
                "row": predecessor_entry["row"],
                "label": predecessor_entry["label"],
                "label_cell": predecessor_entry["label_cell"],
                "price_cells": predecessor_entry["price_cells"],
                "before": dict(before),
                "after": dict(after),
                "currently_governed": bool(mapping_ids),
                "governed_mapping_ids": mapping_ids,
            }
        )
    return changes


def bootstrap_state(selector_path: Path) -> dict[str, Any]:
    selector_file = selector_path.resolve(strict=False)
    root = selector_file.parent
    manifest_dir = root / "manifests"
    manifest_entries: list[str] = []
    if manifest_dir.is_dir():
        try:
            manifest_entries = sorted(entry.name for entry in manifest_dir.iterdir())
        except OSError as exc:
            raise CandidateAuditError(
                "bootstrap manifests directory could not be inspected"
            ) from exc
    return {
        "canonical_root": str(root),
        "root_exists": root.is_dir(),
        "selector_path": str(selector_file),
        "selector_exists": selector_file.exists(),
        "manifest_dir": str(manifest_dir),
        "manifest_entries": manifest_entries,
    }


def build_successor_mapping_snapshot(path: Path) -> list[dict[str, Any]]:
    workbook: Any | None = None
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        return cast(
            list[dict[str, Any]],
            build_governed_mapping_snapshot(
                workbook,
                component_mappings=SUCCESSOR_COMPONENT_PRICE_MAPPINGS,
            ),
        )
    except (KeyError, OSError, ValueError) as exc:
        raise CandidateAuditError(
            f"static successor mapping snapshot could not be built: {exc}"
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def audit_bootstrap_candidate(
    selector_path: Path,
    candidate_workbook: Path,
    *,
    approved_successor: PriceBaseline = SUCCESSOR,
) -> dict[str, Any]:
    selector_file = selector_path.resolve(strict=False)
    candidate_file = candidate_workbook.resolve(strict=False)
    state_before = bootstrap_state(selector_file)
    hold_reasons: list[str] = []
    conflicts: dict[str, list[str]] = {
        "formula": [],
        "missing": [],
        "duplicate": [],
        "mapping_identity_drift": [],
    }
    if state_before["root_exists"] is not True:
        hold_reasons.append("canonical price root does not exist")
    if state_before["selector_exists"] is True:
        hold_reasons.append("active selector already exists")
    if state_before["manifest_entries"]:
        hold_reasons.append("bootstrap manifests directory is not empty")
    if candidate_file != approved_successor.path.resolve(strict=False):
        hold_reasons.append("candidate path differs from static successor binding")
    if (
        not candidate_file.is_relative_to(selector_file.parent)
        or candidate_file.suffix.casefold() != ".xlsx"
    ):
        hold_reasons.append("candidate workbook is outside canonical price root")

    try:
        candidate_sha = sha256_file(candidate_file)
    except OSError:
        candidate_sha = "not-readable"
        hold_reasons.append("candidate workbook could not be read")
    if candidate_sha != approved_successor.sha256:
        hold_reasons.append("candidate SHA-256 differs from static successor binding")

    inspection: WorkbookInspection | None = None
    mapping_snapshot: list[dict[str, Any]] = []
    if (
        candidate_file == approved_successor.path.resolve(strict=False)
        and candidate_sha == approved_successor.sha256
    ):
        try:
            inspection = inspect_workbook(candidate_file)
            conflicts["formula"] = list(inspection.formula_cells)
            conflicts["missing"] = list(inspection.invalid_price_cells)
            conflicts["duplicate"] = list(inspection.duplicate_names)
            mapping_snapshot = build_successor_mapping_snapshot(candidate_file)
        except CandidateAuditError as exc:
            conflicts["mapping_identity_drift"].append(str(exc))
    for category, findings in conflicts.items():
        if findings:
            hold_reasons.append(f"{category} conflicts detected")
    if not mapping_snapshot:
        hold_reasons.append("governed mapping snapshot is incomplete")

    structural_fingerprint = (
        inspection.structural_fingerprint if inspection is not None else "not-validated"
    )
    manifest_id = (
        f"PBM-GENESIS-{candidate_sha[:16].upper()}"
        if candidate_sha != "not-readable"
        else "PBM-GENESIS-NOT-READABLE"
    )
    approval_payload = {
        "intent": BOOTSTRAP_INTENT,
        "manifest_id": manifest_id,
        "candidate_workbook_path": str(candidate_file),
        "candidate_workbook_sha256": candidate_sha,
        "candidate_structural_fingerprint": structural_fingerprint,
        "mapping_snapshot_sha256": sha256_bytes(canonical_json_bytes(mapping_snapshot)),
        "future_cases_only": True,
        "historical_repricing_authorized": False,
    }
    approval_fingerprint = sha256_bytes(canonical_json_bytes(approval_payload))
    audit = {
        "schema_version": BOOTSTRAP_AUDIT_SCHEMA,
        "intent": BOOTSTRAP_INTENT,
        "status": "PASS_CANDIDATE_ONLY" if not hold_reasons else "HOLD",
        "manifest_id": manifest_id,
        "genesis_state": state_before,
        "candidate_workbook": {
            "path": str(candidate_file),
            "sha256": candidate_sha,
            "structural_fingerprint": structural_fingerprint,
        },
        "price_inventory": (
            list(inspection.price_names) if inspection is not None else []
        ),
        "conflicts": conflicts,
        "mapping_snapshot": mapping_snapshot,
        "hold_reasons": hold_reasons,
        "approval_payload": approval_payload,
        "approval_fingerprint": approval_fingerprint,
    }
    if bootstrap_state(selector_file) != state_before:
        raise CandidateAuditError("genesis state changed during bootstrap audit")
    if candidate_sha != "not-readable" and sha256_file(candidate_file) != candidate_sha:
        raise CandidateAuditError("candidate workbook changed during bootstrap audit")
    return audit


def _manifest_data(path: Path) -> Mapping[str, Any]:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise CandidateAuditError("active manifest could not be reread") from exc
    data = cast(Mapping[str, Any], load_json_bytes(raw, "active manifest"))
    validate_manifest(data)
    return data


def _entry_prices(
    workbook: Any,
    entry: Mapping[str, Any],
    conflicts: dict[str, list[str]],
) -> dict[str, int] | None:
    source = cast(Mapping[str, Any], entry["source"])
    identity = cast(Mapping[str, Any], entry["identity"])
    sheet_name = cast(str, source["sheet"])
    try:
        worksheet = workbook[sheet_name]
    except KeyError:
        conflicts["missing"].append(f"missing worksheet: {sheet_name}")
        return None
    label_cell = cast(str, source["label_cell"]).split("!", maxsplit=1)[-1]
    actual_raw_label = worksheet[label_cell].value
    expected_label = identity.get("expected_label")
    strict_label = identity.get("strict_label") is True
    actual_label = (
        actual_raw_label if strict_label else normalize_label(actual_raw_label)
    )
    expected = expected_label if strict_label else normalize_label(expected_label)
    if actual_label != expected:
        conflicts["mapping_identity_drift"].append(
            f"{entry['mapping_id']}: label/source identity changed at "
            f"{sheet_name}!{label_cell}"
        )
        return None
    price_cells = cast(list[str], source["price_cells"])
    values: list[int] = []
    for coordinate in price_cells:
        cell = coordinate.split("!", maxsplit=1)[-1]
        raw_value = worksheet[cell].value
        if isinstance(raw_value, str) and raw_value.startswith("="):
            conflicts["formula"].append(coordinate)
            return None
        parsed = positive_int(raw_value)
        if parsed is None:
            conflicts["missing"].append(
                f"{entry['mapping_id']}: invalid or missing price at {coordinate}"
            )
            return None
        values.append(parsed)
    if cast(str, entry["mapping_kind"]).startswith("component_"):
        if len(values) != 2:
            conflicts["missing"].append(
                f"{entry['mapping_id']}: component requires material and work prices"
            )
            return None
        return {"material_kzt": values[0], "work_kzt": values[1]}
    if len(values) != 1:
        conflicts["missing"].append(
            f"{entry['mapping_id']}: cabinet requires one price"
        )
        return None
    return {"price_kzt": values[0]}


def audit_candidate(
    selector_path: Path,
    candidate_workbook: Path,
) -> dict[str, Any]:
    selector_file = selector_path.resolve(strict=False)
    candidate_file = candidate_workbook.resolve(strict=False)
    if (
        not candidate_file.is_relative_to(selector_file.parent)
        or candidate_file.suffix.casefold() != ".xlsx"
    ):
        raise CandidateAuditError("candidate workbook is outside canonical price root")
    selector_before = selector_file.read_bytes() if selector_file.is_file() else b""
    try:
        active = resolve_active_price_baseline(selector_file)
    except (BaselineContractError, OSError) as exc:
        raise CandidateAuditError(f"active baseline is invalid: {exc}") from exc
    if active.manifest_path is None or active.manifest_sha256 is None:
        raise CandidateAuditError("active baseline has no manifest provenance")
    predecessor = _manifest_data(active.manifest_path)
    predecessor_workbook = cast(Mapping[str, Any], predecessor["workbook"])
    predecessor_snapshot = list(active.mapping_snapshot)

    try:
        candidate_sha_before = sha256_file(candidate_file)
    except OSError as exc:
        raise CandidateAuditError("candidate workbook could not be read") from exc
    inspection = inspect_workbook(candidate_file)
    predecessor_inspection = inspect_workbook(active.path)
    conflicts: dict[str, list[str]] = {
        "formula": list(inspection.formula_cells),
        "missing": list(inspection.invalid_price_cells),
        "duplicate": list(inspection.duplicate_names),
        "mapping_identity_drift": [],
    }
    candidate_snapshot: list[dict[str, Any]] = []
    workbook: Any | None = None
    try:
        workbook = load_workbook(
            candidate_file,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        for raw_entry in predecessor_snapshot:
            entry = copy.deepcopy(dict(raw_entry))
            prices = _entry_prices(workbook, entry, conflicts)
            if prices is None:
                continue
            entry["prices"] = prices
            candidate_snapshot.append(entry)
    finally:
        if workbook is not None:
            workbook.close()

    predecessor_names = {
        (entry["kind"], entry["sheet"], entry["label"])
        for entry in predecessor_inspection.price_names
    }
    candidate_names = {
        (entry["kind"], entry["sheet"], entry["label"])
        for entry in inspection.price_names
    }
    added_names = [list(value) for value in sorted(candidate_names - predecessor_names)]
    removed_names = [
        list(value) for value in sorted(predecessor_names - candidate_names)
    ]
    changed_prices = diff_all_price_values(
        predecessor_inspection,
        inspection,
        predecessor_snapshot,
    )
    structural_compatible = (
        inspection.structural_fingerprint
        == predecessor_workbook["structural_fingerprint"]
    )
    hold_reasons: list[str] = []
    if not structural_compatible:
        hold_reasons.append("workbook structural fingerprint changed")
    if added_names or removed_names:
        hold_reasons.append("price-bearing names were added or removed")
    if len(candidate_snapshot) != len(predecessor_snapshot):
        hold_reasons.append("governed mapping snapshot is incomplete")
    for category, findings in conflicts.items():
        if findings:
            hold_reasons.append(f"{category} conflicts detected")

    manifest_id = f"PBM-{candidate_sha_before[:16].upper()}"
    approval_payload = {
        "manifest_id": manifest_id,
        "predecessor_manifest_id": active.manifest_id,
        "predecessor_manifest_sha256": active.manifest_sha256,
        "candidate_workbook_path": str(candidate_file),
        "candidate_workbook_sha256": candidate_sha_before,
        "candidate_structural_fingerprint": inspection.structural_fingerprint,
        "mapping_snapshot_sha256": sha256_bytes(
            canonical_json_bytes(candidate_snapshot)
        ),
        "future_cases_only": True,
        "historical_repricing_authorized": False,
    }
    approval_fingerprint = sha256_bytes(canonical_json_bytes(approval_payload))
    audit = {
        "schema_version": AUDIT_SCHEMA,
        "status": "PASS_CANDIDATE_ONLY" if not hold_reasons else "HOLD",
        "manifest_id": manifest_id,
        "predecessor": {
            "manifest_id": active.manifest_id,
            "manifest_path": str(active.manifest_path),
            "manifest_sha256": active.manifest_sha256,
        },
        "candidate_workbook": {
            "path": str(candidate_file),
            "sha256": candidate_sha_before,
            "structural_fingerprint": inspection.structural_fingerprint,
        },
        "structural_compatibility": structural_compatible,
        "changed_prices": changed_prices,
        "added_names": added_names,
        "removed_names": removed_names,
        "conflicts": conflicts,
        "mapping_snapshot": candidate_snapshot,
        "hold_reasons": hold_reasons,
        "approval_payload": approval_payload,
        "approval_fingerprint": approval_fingerprint,
    }
    if sha256_file(candidate_file) != candidate_sha_before:
        raise CandidateAuditError("candidate workbook changed during audit")
    if selector_file.read_bytes() != selector_before:
        raise CandidateAuditError("active selector changed during audit")
    if sha256_file(active.manifest_path) != active.manifest_sha256:
        raise CandidateAuditError("active manifest changed during audit")
    return audit


def materialize_audit_artifact(selector_path: Path, audit: Mapping[str, Any]) -> Path:
    """Create only the canonical immutable PASS audit, or verify exact reuse."""
    if audit.get("status") != "PASS_CANDIDATE_ONLY":
        raise CandidateAuditError("only a PASS audit can be materialized")
    try:
        path = cast(Path, audit_artifact_path(selector_path, audit))
    except BaselineContractError as exc:
        raise CandidateAuditError(str(exc)) from exc
    candidate = cast(Mapping[str, Any], audit["candidate_workbook"])
    candidate_path = Path(cast(str, candidate["path"]))
    try:
        if sha256_file(candidate_path) != candidate["sha256"]:
            raise CandidateAuditError(
                "candidate workbook drifted before audit publication"
            )
    except OSError as exc:
        raise CandidateAuditError("candidate workbook could not be rechecked") from exc
    if audit["schema_version"] == BOOTSTRAP_AUDIT_SCHEMA:
        if bootstrap_state(selector_path) != audit["genesis_state"]:
            raise CandidateAuditError("genesis state drifted before audit publication")
    else:
        try:
            active = resolve_active_price_baseline(selector_path)
        except (BaselineContractError, OSError) as exc:
            raise CandidateAuditError("active predecessor is invalid") from exc
        predecessor = cast(Mapping[str, Any], audit["predecessor"])
        if (
            active.manifest_id != predecessor.get("manifest_id")
            or str(active.manifest_path) != predecessor.get("manifest_path")
            or active.manifest_sha256 != predecessor.get("manifest_sha256")
        ):
            raise CandidateAuditError(
                "active predecessor drifted before audit publication"
            )

    directory = path.parent
    if not directory.parent.is_dir() or directory.is_symlink():
        raise CandidateAuditError("canonical audit root is invalid")
    try:
        directory.mkdir(exist_ok=True)
        if directory.is_symlink() or not directory.is_dir():
            raise CandidateAuditError("canonical audit directory is invalid")
        raw = canonical_json_bytes(audit)
        try:
            with path.open("xb") as output:
                output.write(raw)
                output.flush()
                os.fsync(output.fileno())
        except FileExistsError as exc:
            if path.is_symlink() or not path.is_file() or path.read_bytes() != raw:
                raise CandidateAuditError(
                    "existing audit artifact differs from exact bytes"
                ) from exc
        if path.is_symlink() or path.read_bytes() != raw:
            raise CandidateAuditError("published audit artifact bytes differ")
    except OSError as exc:
        raise CandidateAuditError(
            "audit artifact could not be created or verified"
        ) from exc
    return path


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--bootstrap",
        action="store_true",
        help="Audit the exact static successor for first-chain bootstrap",
    )
    parser.add_argument(
        "--active-selector",
        type=Path,
        default=DEFAULT_ACTIVE_SELECTOR,
    )
    parser.add_argument("--candidate-workbook", type=Path, required=True)
    parser.add_argument(
        "--materialize-audit",
        action="store_true",
        help="Create or verify the deterministic immutable PASS audit artifact",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        audit = (
            audit_bootstrap_candidate(args.active_selector, args.candidate_workbook)
            if args.bootstrap
            else audit_candidate(args.active_selector, args.candidate_workbook)
        )
    except (CandidateAuditError, OSError) as exc:
        print(f"AUDIT HOLD: {exc}")
        return 1
    if args.materialize_audit:
        try:
            path = materialize_audit_artifact(args.active_selector, audit)
        except CandidateAuditError as exc:
            print(f"AUDIT HOLD: {exc}", file=sys.stderr)
            return 1
        print(f"AUDIT ARTIFACT: {path}", file=sys.stderr)
    sys.stdout.buffer.write(canonical_json_bytes(audit))
    return 0 if audit["status"] == "PASS_CANDIDATE_ONLY" else 1


if __name__ == "__main__":
    raise SystemExit(main())
