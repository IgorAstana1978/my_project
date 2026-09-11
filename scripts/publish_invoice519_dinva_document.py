"""Build a source-bound DINVA document v0.2 DRAFT or publish its approval."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import stat
import tempfile
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, NoReturn, cast

from openpyxl import load_workbook  # type: ignore[import-untyped]

REPO_ROOT = Path(__file__).resolve().parents[1]
LEDGER_SHA256 = "3391f456ff9a01eed59b455549127a73e46aa97b0f4607c291759b5753959fdc"
YAUO_DECISION_SHA256 = (
    "214a9114c5b676f3754f3220cfe5b3488d9c4ce75325f98072b7e8e9a5f29717"
)
CANONICAL_INVOICE_SHA256 = (
    "17e31d0312f728800d31fd4f125d285edb1114880500d4833261239b87ab58b5"
)
PUBLICATION_AUTHORIZATION_PREFIX = (
    "IGOR_INVOICE519_DINVA_DOCUMENT_V0_2_APPROVAL_PUBLICATION_AUTHORIZED"
)
OUTPUT_FILENAME = "invoice519-dinva-document-v0.2-APPROVED.json"
DRAFT_OUTPUT_FILENAME = "invoice519-dinva-document-v0.2-DRAFT.json"
APPROVED_TOTAL = 19_499_186
PROJECT_ID = "2024/086"
FROZEN_55_SUBTOTAL = 11_963_792
CHECKED_MISSING_33_SUBTOTAL = 7_535_394
APPROVED_LEAD_TIME = "30–40 рабочих дней"
LEAD_TIME_NORMALIZATION_RULE = "APPROVED_ASCII_HYPHEN_RANGE_TO_EN_DASH"
LEAD_TIME_APPROVAL_REFERENCE = "IGOR-INVOICE519-LEAD-TIME-2024-086"
YAUO_POSITION = 87
YAUO_APPROVED_DIMENSIONS = "400×300×250 mm"
YAUO_APPROVED_ENCLOSURE = "Накладной 400х300х250 металл 1,2мм"
YAUO_SOURCE_CELL = "G111"
POSITION_COUNT = 88
SECTION_COUNT = 9
REQUIRED_LEDGER_SOURCE_ROLES = {
    "completed_technical_input",
    "main_price_workbook",
    "custom_sche_metal_workbook",
    "pricing_profile",
    "canonical_invoice_519",
    "ukrm_price_workbook",
    "yarv100_price_workbook",
}
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
UTC_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")
TITLE = "Счёт №519 от 22 июня 2026 года"
PAYER_PREFIX = "Плательщик: "
DELIVERY_SOURCE_PHRASE = "условия поставки EXW г. Астана"
VAT_SOURCE_MARKER = "НДС 0%"
APPARATUS_HEADING = (
    "Применяемые приборы и аппараты согласно схемы, \nпроизводства CHINT"
)


class DocumentPublicationError(ValueError):
    """Publication would violate an authoritative Invoice519 binding."""


class DuplicateJsonKeyError(ValueError):
    """A JSON object contains a duplicate key."""


@dataclass(frozen=True)
class BoundInput:
    path: Path
    expected_sha256: str


@dataclass(frozen=True)
class LoadedInput:
    path: Path
    raw: bytes
    payload: dict[str, Any] | None


@dataclass(frozen=True)
class SourceSet:
    ledger: LoadedInput
    yauo: LoadedInput
    canonical: LoadedInput


@dataclass(frozen=True)
class PublicationResult:
    path: Path
    sha256: str
    size: int
    approved_at: str
    document_fingerprint: str


@dataclass(frozen=True)
class DraftPublicationResult:
    path: Path
    sha256: str
    size: int
    document_fingerprint: str


@dataclass(frozen=True)
class LoadedDraftDocument:
    path: Path
    raw: bytes
    payload: dict[str, Any]
    sha256: str
    document_fingerprint: str
    bound_sources: tuple[tuple[str, LoadedInput], ...]


def fail(message: str) -> NoReturn:
    raise DocumentPublicationError(message)


def require(condition: bool, message: str) -> None:
    if not condition:
        fail(message)


def mapping(value: object, label: str) -> Mapping[str, Any]:
    require(isinstance(value, Mapping), f"{label} must be an object")
    return cast(Mapping[str, Any], value)


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def canonical_json(value: object) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")


def reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise DuplicateJsonKeyError(f"duplicate JSON key: {key}")
        result[key] = value
    return result


def load_json_bytes(raw: bytes, label: str) -> dict[str, Any]:
    try:
        value = json.loads(
            raw.decode("utf-8", errors="strict"),
            object_pairs_hook=reject_duplicate_keys,
        )
    except (UnicodeDecodeError, json.JSONDecodeError, DuplicateJsonKeyError) as exc:
        raise DocumentPublicationError(
            f"{label} is not strict UTF-8 JSON: {exc}"
        ) from exc
    require(isinstance(value, Mapping), f"{label} root must be an object")
    return dict(cast(Mapping[str, Any], value))


def load_bound(source: BoundInput, label: str, *, json_input: bool) -> LoadedInput:
    require(
        SHA256_RE.fullmatch(source.expected_sha256) is not None,
        f"{label} SHA format invalid",
    )
    path = source.path.resolve(strict=True)
    require(
        not path.is_relative_to(REPO_ROOT.resolve(strict=False)),
        f"{label} must be outside Git",
    )
    raw = path.read_bytes()
    require(sha256_bytes(raw) == source.expected_sha256, f"{label} SHA-256 mismatch")
    payload = load_json_bytes(raw, label) if json_input else None
    require(path.read_bytes() == raw, f"{label} changed during initial validation")
    return LoadedInput(path, raw, payload)


def validate_ledger(ledger: Mapping[str, Any], canonical: LoadedInput) -> None:
    require(
        ledger.get("schema_version") == "invoice519_commercial_pricing_ledger.v0.1",
        "ledger schema mismatch",
    )
    require(
        ledger.get("project_id") == PROJECT_ID and ledger.get("invoice_number") == 519,
        "ledger identity mismatch",
    )
    require(
        ledger.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL",
        "ledger authority mismatch",
    )
    require(
        ledger.get("application_status") == "APPLIED", "ledger prices are not applied"
    )
    grain = mapping(ledger.get("price_grain"), "ledger price grain")
    require(
        grain.get("unit_prices_recalculated") is False
        and grain.get("arbitrary_allocation_used") is False,
        "ledger indicates repricing/allocation",
    )
    totals = mapping(ledger.get("ledger_summary"), "ledger summary")
    require(
        totals.get("position_count") == POSITION_COUNT, "ledger position count mismatch"
    )
    require(
        totals.get("approved_total_kzt") == APPROVED_TOTAL
        and totals.get("derived_line_total_kzt") == APPROVED_TOTAL,
        "ledger approved total mismatch",
    )
    require(
        totals.get("frozen_55_subtotal_kzt") == FROZEN_55_SUBTOTAL,
        "frozen 55 subtotal mismatch",
    )
    require(
        totals.get("checked_missing_33_subtotal_kzt") == CHECKED_MISSING_33_SUBTOTAL,
        "checked missing 33 subtotal mismatch",
    )
    require(
        totals.get("price_recalculation_used") is False,
        "ledger price recalculation flag opened",
    )
    bindings = ledger.get("source_input_bindings")
    require(isinstance(bindings, list), "ledger source bindings missing")
    roles: set[str] = set()
    canonical_bound = False
    for raw_binding in cast(list[Any], bindings):
        binding = mapping(raw_binding, "ledger source binding")
        role = binding.get("role")
        require(
            isinstance(role, str) and role not in roles,
            "ledger source role duplicate/invalid",
        )
        roles.add(cast(str, role))
        expected = binding.get("expected_sha256")
        require(
            expected == binding.get("actual_sha256") and isinstance(expected, str),
            "ledger source binding SHA mismatch",
        )
        bound_path = Path(cast(str, binding.get("path"))).resolve(strict=True)
        require(
            sha256_bytes(bound_path.read_bytes()) == expected,
            f"ledger bound source drift: {role}",
        )
        if role == "canonical_invoice_519":
            canonical_bound = (
                bound_path == canonical.path and expected == CANONICAL_INVOICE_SHA256
            )
    require(roles == REQUIRED_LEDGER_SOURCE_ROLES, "ledger source role set mismatch")
    require(canonical_bound, "ledger canonical invoice binding mismatch")
    safety = mapping(ledger.get("safety"), "ledger safety")
    require(
        safety.get("quote_generation_authorized") is False
        and safety.get("invoice_generation_authorized") is False
        and safety.get("client_send_authorized") is False,
        "ledger downstream boundary is open",
    )


def validate_yauo(yauo: Mapping[str, Any], canonical: LoadedInput) -> None:
    require(
        yauo.get("schema_version") == "invoice519_yauo_enclosure_human_decision.v0.1",
        "YAUO decision schema mismatch",
    )
    require(
        yauo.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL", "YAUO authority mismatch"
    )
    require(
        yauo.get("approval_scope") == "ENCLOSURE_DIMENSIONS_ONLY", "YAUO scope mismatch"
    )
    require(
        yauo.get("quote_application_status") == "NOT_APPLIED",
        "YAUO source state mismatch",
    )
    binding = mapping(yauo.get("source_binding"), "YAUO source binding")
    require(
        Path(cast(str, binding.get("path"))).resolve(strict=True) == canonical.path,
        "YAUO canonical path mismatch",
    )
    require(
        binding.get("actual_sha256")
        == CANONICAL_INVOICE_SHA256
        == binding.get("expected_sha256"),
        "YAUO canonical SHA mismatch",
    )
    require(
        binding.get("worksheet") == "Лист1" and binding.get("cell") == YAUO_SOURCE_CELL,
        "YAUO source locator mismatch",
    )
    decision = mapping(yauo.get("technical_decision"), "YAUO technical decision")
    require(
        decision.get("invoice_position_number") == YAUO_POSITION,
        "YAUO position mismatch",
    )
    require(
        decision.get("approved_value") == YAUO_APPROVED_DIMENSIONS,
        "YAUO approved dimensions mismatch",
    )
    require(
        decision.get("change_scope") == f"POSITION_{YAUO_POSITION}_ENCLOSURE_ONLY",
        "YAUO change scope mismatch",
    )


def load_sources(
    ledger: BoundInput, yauo: BoundInput, canonical: BoundInput
) -> SourceSet:
    require(
        ledger.expected_sha256 == LEDGER_SHA256,
        "ledger is not the approved artifact SHA",
    )
    require(
        yauo.expected_sha256 == YAUO_DECISION_SHA256,
        "YAUO decision is not the approved artifact SHA",
    )
    require(
        canonical.expected_sha256 == CANONICAL_INVOICE_SHA256,
        "canonical invoice is not the approved artifact SHA",
    )
    loaded = SourceSet(
        load_bound(ledger, "commercial pricing ledger", json_input=True),
        load_bound(yauo, "YAUO decision", json_input=True),
        load_bound(canonical, "canonical invoice", json_input=False),
    )
    validate_ledger(mapping(loaded.ledger.payload, "ledger"), loaded.canonical)
    validate_yauo(mapping(loaded.yauo.payload, "YAUO decision"), loaded.canonical)
    return loaded


def nonempty(value: object, locator: str) -> str:
    require(
        isinstance(value, str) and bool(value.strip()),
        f"missing source mapping: {locator}",
    )
    return cast(str, value)


def amount_words() -> str:
    return (
        "Девятнадцать миллионов четыреста девяносто девять тысяч "
        "сто восемьдесят шесть тенге 00 тиын"
    )


def build_document_core(sources: SourceSet) -> dict[str, Any]:
    ledger = mapping(sources.ledger.payload, "ledger")
    yauo = mapping(sources.yauo.payload, "YAUO decision")
    positions = ledger.get("positions")
    require(
        isinstance(positions, list) and len(positions) == POSITION_COUNT,
        "ledger positions mismatch",
    )
    workbook = load_workbook(
        sources.canonical.path, data_only=False, read_only=True, keep_links=True
    )
    try:
        require("Лист1" in workbook.sheetnames, "canonical primary sheet missing")
        sheet = workbook["Лист1"]
        require(sheet["C9"].value == TITLE, "canonical document title mismatch")
        payer_text = nonempty(sheet["C10"].value, "Лист1!C10")
        require(payer_text.startswith(PAYER_PREFIX), "canonical payer label mismatch")
        require(
            sheet["F15"].value == APPARATUS_HEADING,
            "canonical apparatus heading mismatch",
        )
        items: list[dict[str, Any]] = []
        sections: list[dict[str, Any]] = []
        current_section: dict[str, Any] | None = None
        for expected_position, raw_position in enumerate(
            cast(list[Any], positions), start=1
        ):
            position = mapping(raw_position, f"ledger position {expected_position}")
            require(
                position.get("invoice_position_number") == expected_position,
                "ledger position order mismatch",
            )
            technical = mapping(
                position.get("technical_description_reference"),
                "technical description reference",
            )
            row = technical.get("row")
            require(
                type(row) is int and technical.get("worksheet") == "Лист1",
                "technical source locator mismatch",
            )
            while (
                current_section is None
                or cast(int, current_section["source_row"]) < cast(int, row) - 1
            ):
                candidate_row = (
                    16
                    if current_section is None
                    else cast(int, current_section["source_row"]) + 1
                )
                while (
                    candidate_row < cast(int, row)
                    and sheet.cell(candidate_row, 2).value is not None
                ):
                    candidate_row += 1
                label = (
                    sheet.cell(candidate_row, 3).value
                    if candidate_row < cast(int, row)
                    else None
                )
                if not isinstance(label, str) or not label.strip():
                    break
                if current_section is not None:
                    current_section["last_position"] = expected_position - 1
                current_section = {
                    "label": label,
                    "first_position": expected_position,
                    "last_position": expected_position,
                    "source_row": candidate_row,
                }
                sections.append(current_section)
            require(
                current_section is not None,
                f"missing section for position {expected_position}",
            )
            cast(dict[str, Any], current_section)["last_position"] = expected_position
            quantity = sheet.cell(cast(int, row), 5).value
            require(
                type(quantity) is int and quantity == position.get("quantity"),
                f"quantity mismatch at position {expected_position}",
            )
            unit_price = position.get("approved_unit_price_kzt")
            line_total = position.get("approved_position_total_kzt")
            require(
                type(unit_price) is int
                and type(line_total) is int
                and quantity * unit_price == line_total,
                f"approved price arithmetic mismatch at position {expected_position}",
            )
            composition = nonempty(sheet.cell(cast(int, row), 6).value, f"Лист1!F{row}")
            enclosure = nonempty(sheet.cell(cast(int, row), 7).value, f"Лист1!G{row}")
            enclosure_reference: str | None = None
            if expected_position == YAUO_POSITION:
                source_binding = mapping(
                    yauo.get("source_binding"), "YAUO source binding"
                )
                require(
                    enclosure == source_binding.get("canonical_cell_value"),
                    "YAUO canonical enclosure drift",
                )
                enclosure = YAUO_APPROVED_ENCLOSURE
                enclosure_reference = cast(str, yauo.get("decision_id"))
            pricing = mapping(position.get("pricing_provenance"), "pricing provenance")
            pricing_sha = sha256_bytes(canonical_json(pricing))
            pricing_reference = (
                f"{ledger['ledger_id']}|POSITION={expected_position}|"
                f"PRICING_SHA256={pricing_sha}"
            )
            locator = f"Лист1!F{row}"
            items.append(
                {
                    "position": expected_position,
                    "name": nonempty(
                        sheet.cell(cast(int, row), 3).value, f"Лист1!C{row}"
                    ),
                    "unit": nonempty(
                        sheet.cell(cast(int, row), 4).value, f"Лист1!D{row}"
                    ),
                    "quantity": quantity,
                    "detailed_technical_composition": composition,
                    "apparatus": {
                        "text": composition,
                        "source_role": "AUTHORITATIVE_DOCUMENT_TEXT_SOURCE",
                        "source_sha256": CANONICAL_INVOICE_SHA256,
                        "source_locator": locator,
                    },
                    "enclosure": enclosure,
                    "approved_unit_price_kzt": unit_price,
                    "approved_line_total_kzt": line_total,
                    "approval_reference": {
                        "pricing": pricing_reference,
                        "technical": f"{CANONICAL_INVOICE_SHA256}|{locator}",
                        "enclosure": enclosure_reference,
                    },
                }
            )
        for section in sections:
            section.pop("source_row")
        require(
            sum(cast(int, item["approved_line_total_kzt"]) for item in items)
            == APPROVED_TOTAL,
            "approved total verification mismatch",
        )
        raw_lines = [
            nonempty(sheet.cell(row, 3).value, f"Лист1!C{row}")
            for row in range(116, 124)
        ]
        require(
            "30-40 рабочих дней" in raw_lines[5], "canonical lead-time text mismatch"
        )
        require(
            DELIVERY_SOURCE_PHRASE in raw_lines[7],
            "canonical delivery term mapping mismatch",
        )
        require(
            VAT_SOURCE_MARKER in nonempty(sheet["C115"].value, "Лист1!C115"),
            "canonical VAT mapping mismatch",
        )
        commercial_lines = [
            {
                "source_order": order,
                "text": source_text,
                "source_sha256": CANONICAL_INVOICE_SHA256,
                "source_locator": f"Лист1!C{row}",
            }
            for order, (row, source_text) in enumerate(
                zip(range(116, 124), raw_lines, strict=True), start=1
            )
        ]
        executor_label = nonempty(sheet["H125"].value, "Лист1!H125").strip()
        executor_full = nonempty(sheet["H126"].value, "Лист1!H126").strip()
        require(
            executor_full == "Инженер-Электрик ПТО Марат А.К.",
            "executor mapping mismatch",
        )
        document = {
            "schema_version": "dinva_quote_invoice_document.v0.2",
            "document_family": "DINVA_CLASSIC_QUOTE_INVOICE_V0_1",
            "document_type": "INVOICE",
            "document_id": "INVOICE519-2024-086",
            "document_number": "519",
            "document_date": "2026-06-22",
            "currency": "KZT",
            "payer": payer_text[len(PAYER_PREFIX) :],
            "object_name": nonempty(sheet["C13"].value, "Лист1!C13"),
            "basis": cast(str, ledger["project_id"]),
            "apparatus_heading": APPARATUS_HEADING,
            "sections": sections,
            "items": items,
            "approved_grand_total_kzt": APPROVED_TOTAL,
            "vat": {
                "rate_percent": 0,
                "included": True,
                "approved_amount_kzt": 0,
                "approved_text": "В том числе НДС 0%",
            },
            "amount_words": {
                "amount_kzt": APPROVED_TOTAL,
                "approved_text": amount_words(),
            },
            "terms": {
                "payment": None,
                "delivery": "EXW г. Астана",
                "manufacturing_lead_time": APPROVED_LEAD_TIME,
                "manufacturing_lead_time_provenance": {
                    "source_order": 6,
                    "source_text": raw_lines[5],
                    "normalized_text": APPROVED_LEAD_TIME,
                    "source_sha256": CANONICAL_INVOICE_SHA256,
                    "source_locator": "Лист1!C121",
                    "normalization_rule": LEAD_TIME_NORMALIZATION_RULE,
                    "approval_reference": LEAD_TIME_APPROVAL_REFERENCE,
                },
                "validity": raw_lines[0],
                "commercial_lines": commercial_lines,
            },
            "signatures": {
                "director_title": nonempty(sheet["C125"].value, "Лист1!C125"),
                "director_name": nonempty(sheet["F125"].value, "Лист1!F125"),
                "executor_label": executor_label,
                "executor_title": "Инженер-Электрик ПТО",
                "executor_name": "Марат А.К.",
                "executor_full_text": executor_full,
            },
        }
        document["document_fingerprint"] = sha256_bytes(canonical_json(document))
        return document
    finally:
        workbook.close()


def utc_now() -> str:
    return datetime.now(UTC).replace(microsecond=0).strftime("%Y-%m-%dT%H:%M:%SZ")


def source_bindings(sources: SourceSet) -> list[dict[str, str]]:
    return [
        {
            "role": "commercial_pricing_ledger",
            "path": str(sources.ledger.path),
            "sha256": LEDGER_SHA256,
        },
        {
            "role": "yauo_enclosure_human_decision",
            "path": str(sources.yauo.path),
            "sha256": YAUO_DECISION_SHA256,
        },
        {
            "role": "canonical_invoice_519",
            "path": str(sources.canonical.path),
            "sha256": CANONICAL_INVOICE_SHA256,
        },
    ]


def build_draft_document(sources: SourceSet) -> dict[str, Any]:
    document = build_document_core(sources)
    bindings = source_bindings(sources)
    document["approval_provenance"] = {
        "status": "DRAFT_UNAPPROVED",
        "authority": None,
        "approval_id": None,
        "approved_at": None,
        "approval_scope": "INVOICE519_DOCUMENT_MODEL_ONLY",
        "approved_document_fingerprint": None,
        "source_bindings": bindings,
        "source_sha256s": [binding["sha256"] for binding in bindings],
        "rendering_authorized": False,
        "client_send_authorized": False,
    }
    return document


def authorization_for(draft: LoadedDraftDocument) -> str:
    return (
        f"{PUBLICATION_AUTHORIZATION_PREFIX}|DRAFT_SHA256={draft.sha256}|"
        f"DOCUMENT_FINGERPRINT={draft.document_fingerprint}"
    )


def approval_id(draft: LoadedDraftDocument) -> str:
    return (
        "IGOR-INVOICE519-DOCUMENT-V0-2-2024-086|"
        f"DRAFT_SHA256={draft.sha256}|"
        f"DOCUMENT_FINGERPRINT={draft.document_fingerprint}"
    )


def build_approved_document(
    draft: LoadedDraftDocument, approved_at: str
) -> dict[str, Any]:
    require(UTC_RE.fullmatch(approved_at) is not None, "approval timestamp is invalid")
    document = copy.deepcopy(draft.payload)
    draft_approval = mapping(
        draft.payload.get("approval_provenance"), "DRAFT approval provenance"
    )
    document["approval_provenance"] = {
        **copy.deepcopy(dict(draft_approval)),
        "status": "APPROVED",
        "authority": "IGOR_DIRECT_HUMAN_APPROVAL",
        "approval_id": approval_id(draft),
        "approved_at": approved_at,
        "approved_document_fingerprint": draft.document_fingerprint,
        "rendering_authorized": True,
    }
    for key in draft.payload:
        if key != "approval_provenance":
            require(
                canonical_json(document[key]) == canonical_json(draft.payload[key]),
                f"approved document changed forbidden field: {key}",
            )
    allowed_approval_changes = {
        "status",
        "authority",
        "approval_id",
        "approved_at",
        "approved_document_fingerprint",
        "rendering_authorized",
    }
    approved_approval = mapping(
        document.get("approval_provenance"), "approved provenance"
    )
    for key in draft_approval:
        if key not in allowed_approval_changes:
            require(
                canonical_json(approved_approval[key])
                == canonical_json(draft_approval[key]),
                f"approved document changed forbidden approval field: {key}",
            )
    return document


def validate_built_document(
    document: Mapping[str, Any], *, expected_status: str = "APPROVED"
) -> None:
    expected_keys = {
        "schema_version",
        "document_family",
        "document_type",
        "document_id",
        "document_number",
        "document_date",
        "currency",
        "payer",
        "object_name",
        "basis",
        "apparatus_heading",
        "sections",
        "items",
        "document_fingerprint",
        "approved_grand_total_kzt",
        "vat",
        "amount_words",
        "terms",
        "signatures",
        "approval_provenance",
    }
    require(set(document) == expected_keys, "built document fields mismatch")
    require(
        document.get("schema_version") == "dinva_quote_invoice_document.v0.2"
        and document.get("document_family") == "DINVA_CLASSIC_QUOTE_INVOICE_V0_1"
        and document.get("document_type") == "INVOICE"
        and document.get("document_number") == "519"
        and document.get("basis") == PROJECT_ID
        and document.get("currency") == "KZT",
        "built document identity mismatch",
    )
    sections = document.get("sections")
    items = document.get("items")
    require(
        isinstance(sections, list) and isinstance(items, list),
        "built collections missing",
    )
    section_list = cast(list[Any], sections)
    item_list = cast(list[Any], items)
    require(len(item_list) == POSITION_COUNT, "built document position count mismatch")
    require(len(section_list) == SECTION_COUNT, "built document section count mismatch")
    next_position = 1
    for raw_section in section_list:
        section = mapping(raw_section, "built section")
        require(
            section.get("first_position") == next_position
            and type(section.get("last_position")) is int
            and cast(int, section["last_position"]) >= next_position,
            "built section coverage mismatch",
        )
        next_position = cast(int, section["last_position"]) + 1
    require(next_position == POSITION_COUNT + 1, "built sections do not cover items")
    running_total = 0
    for expected_position, raw_item in enumerate(item_list, start=1):
        item = mapping(raw_item, "built item")
        require(item.get("position") == expected_position, "built item order mismatch")
        quantity = item.get("quantity")
        unit_price = item.get("approved_unit_price_kzt")
        line_total = item.get("approved_line_total_kzt")
        require(
            type(quantity) is int
            and type(unit_price) is int
            and type(line_total) is int,
            "built item arithmetic fields invalid",
        )
        quantity_int = cast(int, quantity)
        unit_price_int = cast(int, unit_price)
        line_total_int = cast(int, line_total)
        require(
            quantity_int * unit_price_int == line_total_int,
            "built item arithmetic mismatch",
        )
        apparatus = mapping(item.get("apparatus"), "built apparatus")
        require(
            apparatus.get("source_role") == "AUTHORITATIVE_DOCUMENT_TEXT_SOURCE"
            and apparatus.get("source_sha256") == CANONICAL_INVOICE_SHA256
            and apparatus.get("text") == item.get("detailed_technical_composition"),
            "built apparatus provenance mismatch",
        )
        running_total += line_total_int
    require(
        running_total == document.get("approved_grand_total_kzt") == APPROVED_TOTAL,
        "built document total mismatch",
    )
    terms = mapping(document.get("terms"), "built terms")
    require(
        terms.get("payment") is None, "Invoice519 payment lacks governed provenance"
    )
    require(
        terms.get("manufacturing_lead_time") == APPROVED_LEAD_TIME,
        "built document lead time mismatch",
    )
    commercial_lines = terms.get("commercial_lines")
    require(
        isinstance(commercial_lines, list) and len(commercial_lines) == 8,
        "built commercial source line count mismatch",
    )
    commercial_line_list = cast(list[Any], commercial_lines)
    for source_order, raw_line in enumerate(commercial_line_list, start=1):
        line = mapping(raw_line, f"built commercial line {source_order}")
        require(
            line
            == {
                "source_order": source_order,
                "text": line.get("text"),
                "source_sha256": CANONICAL_INVOICE_SHA256,
                "source_locator": f"Лист1!C{115 + source_order}",
            }
            and isinstance(line.get("text"), str)
            and bool(cast(str, line.get("text"))),
            f"built commercial line {source_order} provenance mismatch",
        )
    lead = mapping(
        terms.get("manufacturing_lead_time_provenance"),
        "built lead-time provenance",
    )
    require(
        lead
        == {
            "source_order": 6,
            "source_text": mapping(commercial_line_list[5], "lead source")["text"],
            "normalized_text": APPROVED_LEAD_TIME,
            "source_sha256": CANONICAL_INVOICE_SHA256,
            "source_locator": "Лист1!C121",
            "normalization_rule": LEAD_TIME_NORMALIZATION_RULE,
            "approval_reference": LEAD_TIME_APPROVAL_REFERENCE,
        },
        "built lead-time provenance mismatch",
    )
    require(
        mapping(item_list[YAUO_POSITION - 1], "YAUO item").get("enclosure")
        == YAUO_APPROVED_ENCLOSURE,
        "built YAUO enclosure mismatch",
    )
    approval = mapping(document.get("approval_provenance"), "document approval")
    governed = {
        key: value
        for key, value in document.items()
        if key not in {"approval_provenance", "document_fingerprint"}
    }
    fingerprint = sha256_bytes(canonical_json(governed))
    require(
        document.get("document_fingerprint") == fingerprint,
        "built document fingerprint mismatch",
    )
    require(
        approval.get("approval_scope") == "INVOICE519_DOCUMENT_MODEL_ONLY"
        and approval.get("client_send_authorized") is False,
        "built document approval scope mismatch",
    )
    if expected_status == "DRAFT_UNAPPROVED":
        require(
            approval.get("status") == "DRAFT_UNAPPROVED"
            and approval.get("authority") is None
            and approval.get("approval_id") is None
            and approval.get("approved_at") is None
            and approval.get("approved_document_fingerprint") is None
            and approval.get("rendering_authorized") is False,
            "built DRAFT provenance mismatch",
        )
    else:
        require(expected_status == "APPROVED", "unsupported expected document status")
        require(
            approval.get("status") == "APPROVED"
            and approval.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL"
            and isinstance(approval.get("approval_id"), str)
            and bool(cast(str, approval.get("approval_id")))
            and isinstance(approval.get("approved_at"), str)
            and UTC_RE.fullmatch(cast(str, approval.get("approved_at"))) is not None
            and approval.get("approved_document_fingerprint") == fingerprint
            and approval.get("rendering_authorized") is True,
            "built document approval provenance mismatch",
        )
    bindings = approval.get("source_bindings")
    require(isinstance(bindings, list), "built source bindings missing")
    binding_list = cast(list[Any], bindings)
    require(
        [mapping(value, "built source binding").get("role") for value in binding_list]
        == [
            "commercial_pricing_ledger",
            "yauo_enclosure_human_decision",
            "canonical_invoice_519",
        ],
        "built source binding roles mismatch",
    )
    require(
        approval.get("source_sha256s")
        == [
            mapping(value, "built source binding").get("sha256")
            for value in binding_list
        ],
        "built source SHA list mismatch",
    )


def load_draft_document(source: BoundInput) -> LoadedDraftDocument:
    loaded = load_bound(source, "DRAFT document", json_input=True)
    require(loaded.payload is not None, "DRAFT document payload missing")
    document = cast(dict[str, Any], loaded.payload)
    validate_built_document(document, expected_status="DRAFT_UNAPPROVED")
    fingerprint = cast(str, document["document_fingerprint"])
    approval = mapping(document["approval_provenance"], "DRAFT approval provenance")
    raw_bindings = approval.get("source_bindings")
    require(isinstance(raw_bindings, list), "DRAFT source bindings missing")
    bindings = cast(list[Any], raw_bindings)
    expected_bindings = (
        ("commercial_pricing_ledger", LEDGER_SHA256),
        ("yauo_enclosure_human_decision", YAUO_DECISION_SHA256),
        ("canonical_invoice_519", CANONICAL_INVOICE_SHA256),
    )
    require(
        len(bindings) == len(expected_bindings), "DRAFT source binding count mismatch"
    )
    bound_inputs: dict[str, BoundInput] = {}
    for index, (raw_binding, expected) in enumerate(
        zip(bindings, expected_bindings, strict=True), start=1
    ):
        binding = mapping(raw_binding, f"DRAFT source binding {index}")
        role = binding.get("role")
        path = binding.get("path")
        digest = binding.get("sha256")
        expected_role, expected_digest = expected
        require(role == expected_role, "DRAFT source role invalid")
        require(isinstance(path, str) and bool(path), "DRAFT source path invalid")
        require(
            digest == expected_digest, f"DRAFT source SHA mismatch: {expected_role}"
        )
        bound_inputs[expected_role] = BoundInput(Path(cast(str, path)), expected_digest)
    sources = load_sources(
        bound_inputs["commercial_pricing_ledger"],
        bound_inputs["yauo_enclosure_human_decision"],
        bound_inputs["canonical_invoice_519"],
    )
    expected_draft = build_draft_document(sources)
    require(
        canonical_json(document) == canonical_json(expected_draft),
        "reviewed DRAFT differs from deterministic source-bound document",
    )
    bound_sources = (
        ("commercial_pricing_ledger", sources.ledger),
        ("yauo_enclosure_human_decision", sources.yauo),
        ("canonical_invoice_519", sources.canonical),
    )
    require(loaded.path.read_bytes() == loaded.raw, "DRAFT document TOCTOU mismatch")
    return LoadedDraftDocument(
        loaded.path,
        loaded.raw,
        document,
        source.expected_sha256,
        fingerprint,
        bound_sources,
    )


def recheck_sources(sources: SourceSet) -> None:
    for label, source in (
        ("ledger", sources.ledger),
        ("YAUO decision", sources.yauo),
        ("canonical invoice", sources.canonical),
    ):
        require(source.path.read_bytes() == source.raw, f"{label} TOCTOU mismatch")


def recheck_draft_document(draft: LoadedDraftDocument) -> None:
    require(draft.path.read_bytes() == draft.raw, "DRAFT document TOCTOU mismatch")
    for role, source in draft.bound_sources:
        require(
            source.path.read_bytes() == source.raw,
            f"DRAFT bound source {role} TOCTOU mismatch",
        )


def path_identity(path: Path) -> tuple[int, int]:
    metadata = path.stat()
    return metadata.st_dev, metadata.st_ino


def publish_encoded_document(
    document: Mapping[str, Any],
    recheck_inputs: Callable[[], None],
    output: Path,
    *,
    expected_filename: str,
    expected_status: str,
) -> tuple[Path, str, int]:
    output = output.resolve(strict=False)
    require(output.name == expected_filename, "output filename mismatch")
    require(output.parent.parent.is_dir(), "output directory owner must already exist")
    require(
        not output.parent.exists() and not output.exists(),
        "new output path already exists",
    )
    require(
        not output.is_relative_to(REPO_ROOT.resolve(strict=False)),
        "document output must be outside Git",
    )
    validate_built_document(document, expected_status=expected_status)
    encoded = (json.dumps(document, ensure_ascii=False, indent=2) + "\n").encode(
        "utf-8"
    )
    output.parent.mkdir()
    descriptor = -1
    staging: Path | None = None
    final_link_created = False
    staged_identity: tuple[int, int] | None = None
    try:
        descriptor, raw_staging = tempfile.mkstemp(
            prefix=".invoice519-document-", suffix=".tmp", dir=output.parent
        )
        staging = Path(raw_staging)
        os.chmod(staging, stat.S_IRUSR | stat.S_IWUSR)
        with os.fdopen(descriptor, "wb") as stream:
            descriptor = -1
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())
        staged_identity = path_identity(staging)
        staged_document = load_json_bytes(staging.read_bytes(), "staged document")
        require(staged_document == document, "staged document mismatch")
        validate_built_document(staged_document, expected_status=expected_status)
        recheck_inputs()
        os.link(staging, output)
        final_link_created = True
        require(
            path_identity(output) == staged_identity,
            "published document identity mismatch",
        )
        published_document = load_json_bytes(output.read_bytes(), "published document")
        require(published_document == document, "published document mismatch")
        validate_built_document(published_document, expected_status=expected_status)
        staging.unlink()
        require(
            set(output.parent.iterdir()) == {output},
            "document final directory inventory mismatch",
        )
        return output, sha256_bytes(encoded), len(encoded)
    except BaseException:
        if descriptor >= 0:
            os.close(descriptor)
        if (
            final_link_created
            and staged_identity is not None
            and output.exists()
            and path_identity(output) == staged_identity
        ):
            output.unlink()
        if staging is not None and staging.exists():
            staging.unlink()
        if output.parent.exists() and not any(output.parent.iterdir()):
            output.parent.rmdir()
        raise


def publish_draft_document(
    ledger: BoundInput,
    yauo: BoundInput,
    canonical: BoundInput,
    output: Path,
) -> DraftPublicationResult:
    sources = load_sources(ledger, yauo, canonical)
    document = build_draft_document(sources)
    path, digest, size = publish_encoded_document(
        document,
        lambda: recheck_sources(sources),
        output,
        expected_filename=DRAFT_OUTPUT_FILENAME,
        expected_status="DRAFT_UNAPPROVED",
    )
    return DraftPublicationResult(
        path,
        digest,
        size,
        cast(str, document["document_fingerprint"]),
    )


def publish_document(
    draft_source: BoundInput,
    output: Path,
    authorization: str,
) -> PublicationResult:
    draft = load_draft_document(draft_source)
    require(
        authorization == authorization_for(draft),
        "exact content-bound document approval publication authorization is required",
    )
    approved_at = utc_now()
    document = build_approved_document(draft, approved_at)
    validate_built_document(document, expected_status="APPROVED")
    path, digest, size = publish_encoded_document(
        document,
        lambda: recheck_draft_document(draft),
        output,
        expected_filename=OUTPUT_FILENAME,
        expected_status="APPROVED",
    )
    return PublicationResult(
        path,
        digest,
        size,
        approved_at,
        cast(str, document["document_fingerprint"]),
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("draft", "approved"), required=True)
    parser.add_argument("--commercial-pricing-ledger", type=Path)
    parser.add_argument("--commercial-pricing-ledger-sha256")
    parser.add_argument("--yauo-enclosure-human-decision", type=Path)
    parser.add_argument("--yauo-enclosure-human-decision-sha256")
    parser.add_argument("--canonical-invoice-519", type=Path)
    parser.add_argument("--canonical-invoice-519-sha256")
    parser.add_argument("--draft-document", type=Path)
    parser.add_argument("--draft-document-sha256")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--authorization")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        result: DraftPublicationResult | PublicationResult
        if args.mode == "draft":
            require(
                args.authorization is None
                and args.draft_document is None
                and args.draft_document_sha256 is None,
                "DRAFT mode must not receive approval inputs",
            )
            require(
                isinstance(args.commercial_pricing_ledger, Path)
                and isinstance(args.commercial_pricing_ledger_sha256, str)
                and isinstance(args.yauo_enclosure_human_decision, Path)
                and isinstance(args.yauo_enclosure_human_decision_sha256, str)
                and isinstance(args.canonical_invoice_519, Path)
                and isinstance(args.canonical_invoice_519_sha256, str),
                "DRAFT mode requires exact source paths and SHA-256 values",
            )
            result = publish_draft_document(
                BoundInput(
                    args.commercial_pricing_ledger,
                    args.commercial_pricing_ledger_sha256,
                ),
                BoundInput(
                    args.yauo_enclosure_human_decision,
                    args.yauo_enclosure_human_decision_sha256,
                ),
                BoundInput(
                    args.canonical_invoice_519,
                    args.canonical_invoice_519_sha256,
                ),
                args.output,
            )
        else:
            require(
                args.commercial_pricing_ledger is None
                and args.commercial_pricing_ledger_sha256 is None
                and args.yauo_enclosure_human_decision is None
                and args.yauo_enclosure_human_decision_sha256 is None
                and args.canonical_invoice_519 is None
                and args.canonical_invoice_519_sha256 is None,
                "approved mode accepts only the reviewed DRAFT as its subject",
            )
            require(
                isinstance(args.draft_document, Path)
                and isinstance(args.draft_document_sha256, str)
                and isinstance(args.authorization, str),
                "approved mode requires exact DRAFT path/SHA and authorization",
            )
            result = publish_document(
                BoundInput(args.draft_document, args.draft_document_sha256),
                args.output,
                args.authorization,
            )
    except (OSError, DocumentPublicationError) as exc:
        print(f"HOLD: {exc}")
        return 1
    print(
        "INVOICE519_DOCUMENT="
        + ("DRAFT_UNAPPROVED" if args.mode == "draft" else "APPROVED")
    )
    print(f"DOCUMENT_FINGERPRINT={result.document_fingerprint}")
    print(f"SHA256={result.sha256}")
    print(f"SIZE={result.size}")
    print(f"OUTPUT={result.path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
