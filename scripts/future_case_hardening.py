"""Source-aware PTO and pricing rules for new future Cases only."""

from __future__ import annotations

import copy
import hashlib
import json
from collections.abc import Mapping, Sequence
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from price_baseline_contract import resolve_active_price_baseline, sha256_file

POLICY = "IGOR_POST_WU3_FUTURE_RULES_V1"
FROZEN_CASES = {"CASE-LABORATORY-VRU-20260924-001", "Invoice519"}
K_BY_FAMILY = {
    "ВРУ_FUSE": Decimal("1.05"),
    "АВР": Decimal("1.05"),
    "ШРС": Decimal("1.05"),
    "ВРУ_BREAKER": Decimal("1.20"),
    "ВРУ-ВА": Decimal("1.20"),
    "ЩР": Decimal("1.20"),
    "КРН": Decimal("1.20"),
    "Я5111": Decimal("1.20"),
    "ЯУО": Decimal("1.20"),
}
# Exact approved category/row identities; no SKU search or fuzzy label lookup.
CATEGORIES = {
    "PN-2 100A": ("ВРУ-расп.", 5, "ПН-2 100А"),
    "PN-2 250A": ("ВРУ-расп.", 6, "ПН-2 250А"),
    "PN-2 400A": ("ВРУ-расп.", 7, "ПН-2 400А"),
    "MCCB 3P 16-63A": ("АВР Г-Г", 13, "ВА55/57/59, АМ1 3 полюсные от 16 до 63А"),
    "CONTACTOR 65A": ("АВР Г-Г", 25, "Контактор 65А"),
    "CONTACTOR <=80A": ("АВР Г-Г", 26, "Контактор до 80А"),
    "RKF-8": ("АВР Г-Г", 4, "Реле контроля напряжения RKF-8"),
    "MCB 3P 10A": ("АВР Г-Г", 2, "ВА47 3 полюсный 10А"),
    "MCB 1P 10A": ("АВР Г-Г", 3, "ВА47 1 полюсный 10А"),
    "MCB 1P": ("ВРУ630А", 7, "ВА47 1 полюсный"),
    "Разрядник": ("ВРУ630А", 6, "Разрядник 1Р бочонок"),
    "LAMP": ("АВР Г-Г", 5, "Лампа"),
    "VR32 630A": ("ВРУ630А", 2, "ВР32 39 71240 630А"),
    "CT <=800/5": ("ВРУ630А", 4, "ТТ до 800/5"),
    "AL 5x50": ("ВРУ630А", 5, "5х50 мм (м.п.= 0,7125 кг) - 665А шина АЛ"),
}
FUTURE_INSTALL_TYPES = {
    "fuse_1p",
    "fuse_group_3p",
    "contactor",
    "signal_relay",
    "indicator",
    "busbar",
}
MCCB_INSTALL_TYPES = {"mccb_up_to_100a", "mccb_125_250a", "mccb_400a_plus"}
ROLES = {"POWER_BREAKER", "CONTROL_BREAKER", "CONTACTOR", "FUSE", "FUSE_GROUP", "OTHER"}
CLASS_FIELDS = {
    "family",
    "execution",
    "circuits",
    "roles",
    "categories",
    "fuses",
    "interlocks",
    "overrides",
}


def validate_technical_item(
    context: Mapping[str, Any], item: Mapping[str, Any]
) -> None:
    """Validate explicit normalized facts, never classify by filename/caller flags."""
    future_only(context)
    t = item.get("technical_classification")
    require(
        isinstance(t, Mapping) and set(t) == CLASS_FIELDS,
        "technical classification missing/ambiguous",
    )
    require(t["family"] in K_BY_FAMILY, "technical family missing/ambiguous")
    require(
        t["execution"]
        in (
            {"STANDARD_CONTACTOR_AVR", "BREAKER_AVR"}
            if t["family"] == "АВР"
            else {"STANDARD"}
        ),
        "execution classification missing/ambiguous",
    )
    components = item["components"]
    ids = [c["component_id"] for c in components]
    require(len(ids) == len(set(ids)) and bool(ids), "component identity ambiguous")
    require(
        isinstance(t["roles"], Mapping)
        and set(t["roles"]) == set(ids)
        and set(t["roles"].values()) <= ROLES,
        "component role classification incomplete",
    )
    fuse_ids = {i for i in ids if t["roles"][i] in {"FUSE", "FUSE_GROUP"}}
    if t["family"] == "АВР":
        has_contactors = any(role == "CONTACTOR" for role in t["roles"].values())
        require(
            has_contactors is (t["execution"] == "STANDARD_CONTACTOR_AVR"),
            "AVR execution contradicts switching apparatus roles",
        )
    require(
        isinstance(t["fuses"], Mapping) and set(t["fuses"]) == fuse_ids,
        "fuse technical classification incomplete",
    )
    require(
        isinstance(t["categories"], Mapping)
        and set(t["categories"]) == set(ids) - fuse_ids,
        "pricing category classification incomplete",
    )
    for c in components:
        role = t["roles"][c["component_id"]]
        number(c["quantity"])
        if role in {"POWER_BREAKER", "CONTROL_BREAKER"}:
            kind = c["install_type"]
            require(
                kind in MCCB_INSTALL_TYPES or kind.startswith("modular_"),
                "protection apparatus kind unproven",
            )
            category = t["categories"][c["component_id"]]
            require(
                category in CATEGORIES
                and (
                    (kind in MCCB_INSTALL_TYPES and category == "MCCB 3P 16-63A")
                    or (kind.startswith("modular_") and category.startswith("MCB "))
                ),
                "protection type/category conflict",
            )
        elif role == "CONTACTOR":
            require(
                c["install_type"] == "contactor"
                and t["categories"][c["component_id"]].startswith("CONTACTOR"),
                "contactor execution unproven",
            )
        elif role in {"FUSE", "FUSE_GROUP"}:
            spec = t["fuses"][c["component_id"]]
            require(set(spec) == {"rating_a", "grouping"}, "fuse spec fields mismatch")
            require(
                (
                    role == "FUSE_GROUP"
                    and c["install_type"] == "fuse_group_3p"
                    and spec["grouping"] == "РЩж_3PH"
                )
                or (
                    role == "FUSE"
                    and c["install_type"] == "fuse_1p"
                    and spec["grouping"] == "SINGLE"
                ),
                "fuse group semantics unproven",
            )
            fuse_rule(context, spec["rating_a"], three_phase_rsh=role == "FUSE_GROUP")
    if t["family"] in {"ВРУ_FUSE", "ШРС"}:
        require(bool(fuse_ids), "fuse family without proven fuses")
    if t["family"] in {"ВРУ_BREAKER", "ВРУ-ВА", "ЩР", "КРН", "Я5111", "ЯУО"}:
        require(
            any(v == "POWER_BREAKER" for v in t["roles"].values()) and not fuse_ids,
            "breaker family conflicts with composition",
        )
    circuits = t["circuits"]
    require(isinstance(circuits, list), "circuits must be explicit")
    require(
        t["family"] not in {"АВР", "ВРУ_BREAKER", "ВРУ-ВА", "ЩР", "КРН", "Я5111", "ЯУО"}
        or bool(circuits),
        "power circuit classification missing",
    )
    seen = set()
    for circuit in circuits:
        require(
            set(circuit) == {"circuit_id", "conductor_mm2", "protection_component_ids"},
            "circuit classification fields mismatch",
        )
        require(
            isinstance(circuit["circuit_id"], str)
            and circuit["circuit_id"]
            and circuit["circuit_id"] not in seen,
            "circuit identity ambiguous",
        )
        seen.add(circuit["circuit_id"])
        number(circuit["conductor_mm2"])
        protection = circuit["protection_component_ids"]
        require(
            isinstance(protection, list)
            and bool(protection)
            and len(protection) == len(set(protection))
            and all(i in ids and t["roles"][i] == "POWER_BREAKER" for i in protection),
            "power protection identification missing/ambiguous",
        )
    require(
        isinstance(t["interlocks"], Mapping)
        and set(t["interlocks"]) == {"mechanical", "electrical"}
        and all(type(v) is bool for v in t["interlocks"].values()),
        "interlock classification incomplete",
    )
    require(
        isinstance(t["overrides"], Mapping)
        and set(t["overrides"]) <= {"human", "project", "source"},
        "override authority role unknown",
    )
    for fields in t["overrides"].values():
        require(
            isinstance(fields, Mapping)
            and set(fields) <= {"automation_brand", "enclosure_mm"},
            "unsupported override field; power protection is governed policy",
        )
        for record in fields.values():
            require(
                isinstance(record, Mapping)
                and set(record) == {"path", "sha256", "locator"}
                and all(isinstance(v, str) and bool(v) for v in record.values()),
                "override source binding fields mismatch",
            )


def load_technical_binding(
    context: Mapping[str, Any], binding: Mapping[str, Any] | None
) -> tuple[dict[str, Any], dict[str, Any]]:
    require(
        isinstance(binding, Mapping) and set(binding) == {"path", "sha256", "item_id"},
        "source-bound technical classification required",
    )
    path = Path(binding["path"])
    require(
        path.is_absolute() and sha256_file(path) == binding["sha256"],
        "technical source SHA drift",
    )
    data = json.loads(path.read_bytes(), object_pairs_hook=strict_pairs)
    require(
        data.get("schema_version") == "confirmed_composition_artifact.v0.1"
        and data.get("future_context") == dict(context)
        and data.get("confirmed_by") == "Igor",
        "technical source/context identity mismatch",
    )
    from validate_confirmed_composition_artifact import (
        validate_confirmed_composition_artifact,
    )

    validation = validate_confirmed_composition_artifact(path)
    require(
        validation.status == "PASS",
        "confirmed technical source invalid: " + "; ".join(validation.red_flags),
    )
    items = [i for i in data["items"] if i["item_id"] == binding["item_id"]]
    require(len(items) == 1, "technical item missing/ambiguous")
    require(
        sha256_file(path) == binding["sha256"], "technical source changed during read"
    )
    return data, items[0]


def bound_bom(
    context: Mapping[str, Any], item: Mapping[str, Any]
) -> list[dict[str, Any]]:
    validate_technical_item(context, item)
    t = item["technical_classification"]
    result = []
    for c in item["components"]:
        cid = c["component_id"]
        role = t["roles"][cid]
        category = t["categories"].get(cid)
        qty = number(c["quantity"])
        footprint = None
        if role in {"FUSE", "FUSE_GROUP"}:
            spec = t["fuses"][cid]
            applied = fuse_rule(
                context, spec["rating_a"], three_phase_rsh=role == "FUSE_GROUP"
            )
            category = applied["pricing_category"]
            qty *= applied["quantity_per_group"]
            footprint = applied["physical_footprint"]
        require(category in CATEGORIES, "explicit category mapping required")
        result.append(
            {
                "component_id": cid,
                "category": category,
                "quantity": str(qty),
                "role": role,
                "install_type": c["install_type"],
                "physical_footprint": footprint,
            }
        )
    return result


def future_input_projection(binding: Mapping[str, Any]) -> dict[str, Any]:
    """Normalized technical overlay consumed by the existing input builder/runner."""
    require(
        set(binding) == {"path", "sha256"}, "confirmed source binding fields mismatch"
    )
    path = Path(binding["path"])
    require(
        path.is_absolute() and sha256_file(path) == binding["sha256"],
        "confirmed source SHA drift",
    )
    source = json.loads(path.read_bytes(), object_pairs_hook=strict_pairs)
    context = source.get("future_context")
    future_only(context)
    items, rows, applications = [], [], []
    for original in source["items"]:
        item_binding = {**binding, "item_id": original["item_id"]}
        rules = pto_rules(context, technical_binding=item_binding)
        bom = bound_bom(context, original)
        check_avr_bom(rules, bom)
        item = copy.deepcopy(original)
        item.pop("technical_classification")
        item.pop("confirmation_note", None)
        for c, line in zip(item["components"], bom, strict=True):
            quantity = number(line["quantity"])
            c["component_code"] = line["category"]
            c["quantity"] = (
                int(quantity) if quantity == int(quantity) else float(quantity)
            )
            rows.append(
                {
                    "product_name": item["product_name"],
                    "cabinet_code": item["cabinet"]["cabinet_code"],
                    "consumables_factor": float(Decimal(rules["K"])),
                    "component_code": c["component_code"],
                    "component_qty": c["quantity"],
                    "install_type": c["install_type"],
                }
            )
        items.append(item)
        applications.append(
            {"item_id": original["item_id"], "rules": rules, "bom": bom}
        )
    require(
        sha256_file(path) == binding["sha256"],
        "confirmed source changed during projection",
    )
    return {
        "context": context,
        "source_metadata": {
            key: source[key]
            for key in (
                "confirmation_id",
                "confirmed_by",
                "confirmed_at",
                "source_links",
            )
        },
        "items": items,
        "rows": rows,
        "applications": applications,
    }


def validate_future_draft(data: Mapping[str, Any]) -> dict[str, Any]:
    binding = data["source"].get("future_technical_binding")
    require(isinstance(binding, Mapping), "future technical binding required")
    expected = future_input_projection(binding)
    require(
        data["items"] == expected["items"]
        and data["calculator_input_format"]["rows"] == expected["rows"],
        "completed draft differs from authoritative future composition/K overlay",
    )
    require(
        data["source"].get("future_context") == expected["context"],
        "future draft context mismatch",
    )
    require(
        all(
            data["source"].get(key) == value
            for key, value in expected["source_metadata"].items()
        ),
        "confirmed source metadata differs from authoritative artifact",
    )
    return expected


class RuleError(ValueError):
    """A rule is outside its approved authority."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuleError(message)


def future_only(context: Mapping[str, Any]) -> None:
    require(
        context.get("case_state") == "NEW_FUTURE"
        and context.get("future_cases_only") is True
        and context.get("historical_repricing_authorized") is False
        and bool(context.get("case_id"))
        and context["case_id"] not in FROZEN_CASES,
        "accepted/historical Case is outside future scope",
    )


def number(value: Any, *, zero: bool = False) -> Decimal:
    require(not isinstance(value, bool), "boolean is not a number")
    try:
        result = Decimal(str(value))
    except Exception as exc:
        raise RuleError("invalid number") from exc
    require(
        result.is_finite() and (result >= 0 if zero else result > 0), "invalid number"
    )
    return result


def bound_value(record: Mapping[str, Any]) -> Any:
    """Resolve an exact JSON pointer from unchanged authoritative source bytes."""
    path = Path(record["path"])
    require(path.is_absolute(), "override source must be absolute")
    raw = path.read_bytes()
    require(
        hashlib.sha256(raw).hexdigest() == record["sha256"], "override source drift"
    )
    pointer = record["locator"]
    require(
        isinstance(pointer, str) and pointer.startswith("/"), "JSON pointer required"
    )
    value = json.loads(raw, object_pairs_hook=strict_pairs)
    for key in pointer[1:].split("/"):
        key = key.replace("~1", "/").replace("~0", "~")
        value = value[int(key)] if isinstance(value, list) else value[key]
    return value


def strict_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        require(key not in result, "duplicate source key")
        result[key] = value
    return result


def override(
    field: str, default: Any, sources: Mapping[str, Mapping[str, Mapping[str, Any]]]
) -> tuple[Any, dict[str, Any]]:
    # Explicit Human decisions supersede project requirements, then source facts.
    for role in ("human", "project", "source"):
        if field in sources.get(role, {}):
            record = sources[role][field]
            return bound_value(record), {"role": role, **record}
    return default, {
        "role": "GOVERNED_FUTURE_POLICY",
        "path": str(Path(__file__).resolve()),
        "sha256": sha256_file(Path(__file__)),
        "locator": field,
    }


def pto_rules(
    context: Mapping[str, Any],
    *,
    family: str | None = None,
    conductor_mm2: Any | None = None,
    standard_contactor_avr: bool | None = None,
    sources: Mapping[str, Mapping[str, Mapping[str, Any]]] | None = None,
    technical_binding: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    future_only(context)
    _, item = load_technical_binding(context, technical_binding)
    t = item["technical_classification"]
    require(
        family is None or family == t["family"],
        "caller family contradicts authoritative technical family",
    )
    family = t["family"]
    execution = t["execution"] == "STANDARD_CONTACTOR_AVR"
    require(
        standard_contactor_avr is None or standard_contactor_avr is execution,
        "caller boolean contradicts source-bound AVR execution",
    )
    evidence = t["overrides"]
    require(sources is None or sources == evidence, "unbound caller overrides")
    brand, brand_source = override("automation_brand", "EKF", evidence)
    require(isinstance(brand, str) and bool(brand.strip()), "brand missing")
    areas = [number(c["conductor_mm2"]) for c in t["circuits"]]
    area = max(areas) if areas else None
    require(
        conductor_mm2 is None or number(conductor_mm2) == area,
        "unbound caller conductor claim",
    )
    breaker, breaker_source = override(
        "power_breaker_type", "MCCB" if area is not None and area >= 25 else None, {}
    )
    enclosure, enclosure_source = override(
        "enclosure_mm", [1700, 800, 500] if family.startswith("ВРУ") else None, evidence
    )
    require(
        enclosure is None
        or isinstance(enclosure, list)
        and len(enclosure) == 3
        and all(type(x) is int and x > 0 for x in enclosure),
        "invalid enclosure",
    )
    return {
        "policy": POLICY,
        "automation_brand": brand,
        "power_breaker_type": breaker,
        "enclosure_mm": enclosure,
        "mechanical_interlock_required": execution,
        "mechanical_interlock_separate_priced_line": False,
        "electrical_interlock_required": execution,
        "family": family,
        "context": dict(context),
        "technical_binding": dict(technical_binding),
        "K": str(K_BY_FAMILY[family]),
        "provenance": {
            "automation_brand": brand_source,
            "power_breaker_type": breaker_source,
            "enclosure_mm": enclosure_source,
        },
    }


def fuse_rule(
    context: Mapping[str, Any], rating_a: Any, *, three_phase_rsh: bool = False
) -> dict[str, Any]:
    future_only(context)
    rating = number(rating_a)
    require(rating <= 400, "PPN/PN-2 outside approved range")
    band = 100 if rating <= 100 else 250 if rating <= 250 else 400
    return {
        "physical_footprint": "PPN-33" if rating <= 160 else None,
        "pricing_category": f"PN-2 {band}A",
        "quantity_per_group": 3 if three_phase_rsh else 1,
        "rating_a": str(rating),
    }


def designation_category(value: str) -> str:
    require(value == "КЗ-1000В 0,47 мкФ", "unknown designation; mapping required")
    return "Разрядник"


def check_avr_bom(rules: Mapping[str, Any], bom: Sequence[Mapping[str, Any]]) -> None:
    context = rules["context"]
    binding = rules["technical_binding"]
    require(
        dict(rules) == pto_rules(context, technical_binding=binding),
        "governed rule record modified",
    )
    _, item = load_technical_binding(context, binding)
    expected = bound_bom(context, item)
    require(
        list(bom) == expected,
        "BOM differs from bound technical composition/classification",
    )
    t = item["technical_classification"]
    by_id = {c["component_id"]: c for c in item["components"]}
    for circuit in t["circuits"]:
        if number(circuit["conductor_mm2"]) >= 25:
            require(
                all(
                    by_id[i]["install_type"] in MCCB_INSTALL_TYPES
                    for i in circuit["protection_component_ids"]
                ),
                "conductor rule requires proven industrial power protection",
            )
    require(
        all(line.get("category") != "MECHANICAL_INTERLOCK" for line in bom),
        "mechanical interlock cannot be a separate priced line",
    )
    if rules["mechanical_interlock_required"]:
        require(
            t["interlocks"]["mechanical"] is True
            and t["interlocks"]["electrical"] is True
            and any(line["role"] == "CONTACTOR" for line in bom),
            "standard contactor AVR requires both technical interlocks",
        )


def calculate_future_price(
    context: Mapping[str, Any],
    *,
    family: str | None = None,
    bom: Sequence[Mapping[str, Any]],
    cabinet_cost: Mapping[str, Any],
    work_source: Mapping[str, Any],
    selector: Path,
    rules: Mapping[str, Any],
) -> dict[str, Any]:
    """DRAFT calculation; input cost/flat work must carry exact source bindings."""
    future_only(context)
    require(family is None or family == rules["family"], "technical family conflict")
    family = rules["family"]
    require(rules["K"] == str(K_BY_FAMILY[family]), "K drift")
    require(rules["automation_brand"] == "EKF", "non-EKF category authority required")
    check_avr_bom(rules, bom)
    require(cabinet_cost.get("role") == "ENCLOSURE_COST", "cabinet must be cost")
    cost = number(bound_value(cabinet_cost), zero=True)
    work = number(bound_value(work_source), zero=True)
    baseline = resolve_active_price_baseline(selector)
    rows = []
    material = Decimal(0)
    with_baseline = load_workbook(baseline.path, read_only=True, data_only=False)
    try:
        for line in bom:
            category = line["category"]
            require(category in CATEGORIES, "category mapping required")
            sheet, row, label = CATEGORIES[category]
            ws = with_baseline[sheet]
            require(
                " ".join(str(ws.cell(row, 1).value).split()) == label,
                "category identity drift",
            )
            price = number(ws.cell(row, 2).value)
            quantity = number(line["quantity"])
            material += price * quantity
            rows.append(
                {
                    "category": category,
                    "quantity": str(quantity),
                    "unit_material_kzt": str(price),
                    "locator": f"{sheet}!B{row}",
                }
            )
    finally:
        with_baseline.close()
    require(sha256_file(baseline.path) == baseline.sha256, "workbook drift")
    current = resolve_active_price_baseline(selector)
    require(current == baseline, "baseline changed during calculation")
    require(
        number(bound_value(cabinet_cost), zero=True) == cost
        and number(bound_value(work_source), zero=True) == work,
        "cost/work source changed during calculation",
    )
    from calc_quote_price_draft import canonical_material_work_price

    total = canonical_material_work_price(cost, material, work, K_BY_FAMILY[family])
    return {
        "status": "DRAFT_NOT_APPROVED",
        "policy_id": POLICY,
        "policy_sha256": sha256_file(Path(__file__)),
        "cost_inputs": {"cabinet": dict(cabinet_cost), "work": dict(work_source)},
        "cabinet_cost_kzt": str(cost),
        "material_kzt": str(material),
        "work_kzt": str(work),
        "K": str(K_BY_FAMILY[family]),
        "unit_price_kzt": int(total),
        "rows": rows,
        "workbook_sha256": baseline.sha256,
        "manifest_sha256": baseline.manifest_sha256,
        "client_send_authorized": False,
    }
