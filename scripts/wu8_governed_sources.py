"""WU8 read-only, exact-source resolvers. No activation or publication API."""

from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from price_baseline_contract import resolve_active_price_baseline, sha256_file

MAIN_SHA = "02ca5be9b2eb6775289ee1053c389a659c85e2bd27a2b9867b7d523d6d6e4096"
METAL_SIZES = {
    (500, 400, 250): (22, 2500, 400, 1500),
    (600, 400, 250): (28, 2625, 600, 1500),
    (700, 500, 250): (35, 3250, 600, 1500),
    (1000, 600, 250): (42, 3500, 1150, 1500),
    (1000, 800, 250): (44, 4750, 1150, 1500),
    (1200, 800, 300): (46, 4750, 1150, 2000),
    (1700, 800, 500): (52, 6750, 3500, 2000),
    (1700, 1000, 500): (51, 6750, 3500, 2000),
    (2000, 800, 500): (53, 10000, 2250, 2000),
    (2000, 1000, 500): (54, 10000, 2250, 12000),
}
# Evidence is restricted to breaker types already supported by the future flow.
# Row 9's 3P semantics were explicitly accepted in WU7 and reiterated in WU8.
# Rows with missing poles/characteristic/rating are audited, never inferred.
SHR_BREAKERS = {
    "MCCB SHR 3P 16-63A": (8, "MCCB", 3, None, 16, 63),
    "MCCB 3P 80-100A": (9, "MCCB", 3, None, 80, 100),
    "MCB 3P D 80-100A": (15, "MCB", 3, "D", 80, 100),
    "MCB 3P D 10-63A": (16, "MCB", 3, "D", 10, 63),
    "MCB 1P D 10-63A": (17, "MCB", 1, "D", 10, 63),
    "MCB 3P B <=63A": (18, "MCB", 3, "B", 1, 63),
    "MCB 1P B <=63A": (19, "MCB", 1, "B", 1, 63),
}
SHR_LABELS = {
    8: "ВА55/57/59, АМ1 3 полюсные от 16 до 63А",
    9: "ВА55/57/59, АМ1 от 80 до 100А",
    15: "ВА47 3Р D 80-100A",
    16: "ВА47 3Р D 10-63A",
    17: "ВА47 1Р D 10-63A",
    18: "ВА47 3Р х-ка В до 63А",
    19: "ВА47 1Р х-ка В до 63А",
}
AVR_SHEETS = {
    "Г-Г": ("АВР Г-Г", 2, "K", "L", 27, 33, 20),
    "Г-Д": ("АВР Г-Д", 2, "P", "Q", 5, 11, 23),
    "Г-Г-Д": ("АВР-Г-Г-Д", 3, "J", "K", 25, 31, 22),
}
DRIVE_RATINGS = (100, 160, 250, 400, 630)
CONTACTOR_LABELS = (
    "Контактор до 18А",
    "Контактор до 25А",
    "Контактор до 32А",
    "Контактор 40А",
    "Контактор 50А",
    "Контактор 65А",
    "Контактор до 80А",
    "Контактор до 95А",
    "Контактор до 100А",
    "Контактор до 125А",
    "Контактор до 160А",
    "Контактор до 250А",
    "Контактор до 400А",
    "Контактор до 630А",
    "Контактор до 800А",
    "Контактор до 1000А",
)
CONTACTOR_RATINGS = (
    18,
    25,
    32,
    40,
    50,
    65,
    80,
    95,
    100,
    125,
    160,
    250,
    400,
    630,
    800,
    1000,
)
CONTACTOR_SOURCE_RATINGS = {
    topology: CONTACTOR_RATINGS if topology == "Г-Д" else CONTACTOR_RATINGS[:-1]
    for topology in AVR_SHEETS
}
# Exact cabinet identity, main cost row, dimensions, metal cost row, O/P inputs.
# Main M prices were proved against metal D costs, never metal B/C sales prices.
CABINETS = {
    "CAB-KRN-12": ("КРН-12", 6, (265, 330, 100), 61, 1250, 350),
    "CAB-KRN-18": ("КРН-18", 7, (265, 440, 100), 62, 1250, 350),
    "CAB-KRN-24": ("КРН-24", 8, (395, 330, 100), 63, 1250, 350),
    "CAB-KRN-36": ("КРН-36", 9, (540, 330, 100), 64, 1500, 350),
    "CAB-KRN-48": ("КРН-48", 10, (540, 440, 100), 65, 1800, 350),
    "CAB-KRN-54": ("КРН-54", 11, (540, 440, 100), 66, 1800, 350),
    "CAB-KRN-72": ("КРН-72", 12, (540, 600, 100), 67, 3000, 600),
    "CAB-KRN-96": ("КРН-96", 13, (1000, 600, 250), 68, 3500, 1150),
    "CAB-KRV-12": ("КРВ-12", 18, (275, 320, 125), 74, 1250, 350),
    "CAB-KRV-18": ("КРВ-18", 19, (405, 320, 125), 75, 1250, 350),
    "CAB-KRV-24": ("КРВ-24", 20, (405, 320, 125), 76, 1250, 350),
    "CAB-KRV-36": ("КРВ-36", 21, (550, 320, 125), 77, 1500, 350),
    "CAB-KRV-48": ("КРВ-48", 22, (630, 320, 125), 78, 1800, 350),
    "CAB-KRV-54": ("КРВ-54", 23, (550, 450, 125), 79, 1800, 350),
    "CAB-KRV-72": ("КРВ-72", 24, (700, 450, 125), 80, 3000, 600),
}
AVR_DECISION = {
    "decision_id": "IGOR-WU8-AVR-MOTOR-DRIVE-20261013",
    "authority": "IGOR_DIRECT_HUMAN_DECISION",
    "provenance": "Direct Igor CONTINUE WU8 user instruction, 2026-10-13",
    "workbook_sha256": MAIN_SHA,
    "automatic_price_unit": "one breaker of the exact rating",
    "drive_price_unit": "one drive of the exact rating",
    "quantities": {"Г-Г": 2, "Г-Д": 2, "Г-Г-Д": 3},
    "work": "separate authoritative topology/rating work table; never included",
    "drive_630a": "HOLD until source-bound availability confirmation",
    "case_approval": False,
    "commercial_approval": False,
    "client_send_authorized": False,
}


def require(ok: bool, message: str) -> None:
    if not ok:
        from future_case_hardening import RuleError

        raise RuleError(message)


def literal(value: Any) -> Decimal:
    require(
        type(value) in (int, float) and value > 0, "positive numeric source required"
    )
    result = Decimal(str(value))
    require(result.is_finite(), "finite source required")
    return result


def policy_binding() -> dict[str, str]:
    return {
        "path": str(Path(__file__).resolve()),
        "sha256": sha256_file(Path(__file__)),
    }


def cabinet_display(source: dict[str, Any], index: int) -> str:
    """Resolve display from bound confirmed identity and checked cost, never text."""
    from build_future_case_document_source import (
        CALCULATION_ROLE,
        COMPOSITION_ROLE,
        bound_snapshot,
    )

    original = source["items"][index]
    records = {
        role: [b for b in source["source_bindings"] if b["role"] == role]
        for role in (COMPOSITION_ROLE, CALCULATION_ROLE)
    }
    if not any(records.values()):
        require(
            not re.match(r"^(?:ЩР[нНвВ]|КР[НВ])(?=$|[\s-])", original["enclosure"]),
            "cabinet identity missing: source text cannot select a cabinet",
        )
        return original["enclosure"]
    require(
        all(len(r) == 1 for r in records.values()), "cabinet identity binding ambiguous"
    )
    composition = bound_snapshot(records[COMPOSITION_ROLE][0]).payload()
    calculated = bound_snapshot(records[CALCULATION_ROLE][0]).payload()
    confirmed = [i for i in composition["items"] if i["item_id"] == original["item_id"]]
    prices = [i for i in calculated["items"] if i["item_id"] == original["item_id"]]
    require(len(confirmed) == len(prices) == 1, "cabinet item identity ambiguous")
    cabinet = confirmed[0]["cabinet"]
    require(
        cabinet["cabinet_label"] == original["enclosure"],
        "raw cabinet designation drift",
    )
    code = cabinet["cabinet_code"]
    if code in CABINETS:
        cost = prices[0]["rule_application"]["price"].get("cabinet_source", {})
        require(
            cost.get("cabinet_code") == code
            and cost.get("commercial_identity") == CABINETS[code][0]
            and cost.get("source_designation") == cabinet["cabinet_label"],
            "checked cabinet identity/cost mapping required",
        )
        return cost["commercial_identity"]
    require(
        not code.startswith(("CAB-KRN", "CAB-KRV"))
        and not re.match(r"^ЩР[нНвВ](?=$|[\s-])", cabinet["cabinet_label"]),
        "exact cabinet identity mapping required",
    )
    return original["enclosure"]


def dimensions(source: str) -> tuple[int, int, int]:
    found = re.findall(r"(?<!\d)(\d+)\s*[xх×*]\s*(\d+)\s*[xх×*]\s*(\d+)(?!\d)", source)
    require(len(found) == 1, "exact confirmed enclosure dimensions required")
    return tuple(int(v) for v in found[0])


def metal_cost(size: tuple[int, int, int]) -> dict[str, Any]:
    from governed_metal_cost import GENESIS_PATH, GENESIS_SHA, _inspect_cost

    require(size in METAL_SIZES, "metal dimensions outside approved ten-size whitelist")
    row, fabrication, painting, extra = METAL_SIZES[size]
    suffix = {52: " ВРУ ", 53: " ПР ", 54: " ПР ", 51: ""}.get(row, " ")
    identity = metal_identity(
        row, "х".join(map(str, size)) + suffix, size, fabrication, painting, extra
    )
    result = _inspect_cost(GENESIS_PATH, GENESIS_SHA, identity)
    return {
        **result,
        "dimensions_mm": list(size),
        "identity": identity,
        "policy_binding": policy_binding(),
    }


def metal_identity(
    row: int,
    label: str,
    size: tuple[int, int, int],
    fabrication: int,
    painting: int,
    extra: int,
) -> dict[str, Any]:
    identity = {
        "mapping_id": f"METAL-{size[0]}X{size[1]}X{size[2]}-D{row}-COST",
        "source_role": "ENCLOSURE_COST",
        "sheet": "Лист1",
        "row": row,
        "label": label,
        "formulas": {
            f"D{row}": (
                f"=ROUNDUP((M{row}+N{row})*R{row}" f"+O{row}+P{row}+S{row}+Q{row},0)"
            ),
            f"H{row}": f"=E{row}/1000",
            f"I{row}": f"=F{row}/1000",
            f"J{row}": f"=G{row}/1000",
            f"K{row}": f"=H{row}*I{row}*3+H{row}*J{row}*2+I{row}*J{row}*2",
            f"L{row}": f"=K{row}*$B$1*8.42",
            f"M{row}": f"=L{row}*$B$2",
            f"N{row}": f"=K{row}*2*0.25*$B$3*1.33",
            f"Q{row}": "=1250",
        },
        "identity_values": {
            "B1": 1,
            f"E{row}": size[0],
            f"F{row}": size[1],
            f"G{row}": size[2],
            f"O{row}": fabrication,
            f"P{row}": painting,
            f"R{row}": 1.05,
            f"S{row}": extra,
        },
        "price_inputs": ["B2", "B3"],
        "excluded_sales_outputs": [f"B{row}", f"C{row}"],
    }
    return identity


def cabinet_cost(ws: Any, baseline: Any, code: str) -> dict[str, Any]:
    require(code in CABINETS, "exact cabinet identity mapping required")
    name, row, size, _, _, _ = CABINETS[code]
    label = f"Корпус {name} {'х'.join(map(str, size))}"
    cell = source_cell(ws, f"L{row}", f"M{row}", label)
    entries = [
        e
        for e in baseline.mapping_snapshot
        if e["mapping_id"] == f"CABINET-DYNAMIC-{code}"
    ]
    if code in {"CAB-KRN-12", "CAB-KRN-18", "CAB-KRN-24"}:
        require(
            len(entries) == 1
            and entries[0]["source"]
            == {
                "sheet": "КРН",
                "row": row,
                "label_cell": f"КРН!L{row}",
                "price_cells": [f"КРН!M{row}"],
            }
            and Decimal(str(entries[0]["prices"]["price_kzt"]))
            == Decimal(cell["value"]),
            "cabinet approved snapshot/source mismatch",
        )
    else:
        require(baseline.sha256 == MAIN_SHA, "WU8 authoritative source changed")
    return {
        **cell,
        "cabinet_code": code,
        "commercial_identity": name,
        "dimensions_mm": list(size),
        "source_role": "ENCLOSURE_COST",
        "cost_kzt": cell["value"],
        "workbook_path": str(baseline.path),
        "workbook_sha256": baseline.sha256,
        "manifest_sha256": baseline.manifest_sha256,
        "policy_binding": policy_binding(),
    }


def source_cell(
    ws: Any, label_cell: str, price_cell: str, expected: str
) -> dict[str, Any]:
    label = ws[label_cell].value
    require(
        " ".join(str(label).split()) == expected,
        f"source identity drift: {ws.title}!{label_cell}",
    )
    value = literal(ws[price_cell].value)
    return {
        "label": label,
        "label_locator": f"{ws.title}!{label_cell}",
        "source_locator": f"{ws.title}!{price_cell}",
        "value": str(value),
        "value_kind": "numeric_literal",
    }


def shr_price(ws: Any, category: str) -> dict[str, Any]:
    row = SHR_BREAKERS[category][0]
    material = source_cell(ws, f"A{row}", f"B{row}", SHR_LABELS[row])
    work = source_cell(ws, f"A{row}", f"C{row}", SHR_LABELS[row])
    return {
        "material": material,
        "work": work,
        "signature": list(SHR_BREAKERS[category][1:]),
    }


def avr_category(category: str) -> tuple[str, str, int] | None:
    match = re.fullmatch(r"(MCCB|CONTACTOR) AVR (Г-Г|Г-Д|Г-Г-Д) (\d+)A", category)
    return (match[2], match[1], int(match[3])) if match else None


def avr_price(
    ws: Any,
    topology: str,
    kind: str,
    rating: int,
    quantity: Decimal,
    availability: dict[str, Any] | None = None,
) -> dict[str, Any]:
    sheet, count, label_col, value_col, automatic_row, drive_row, contactor_row = (
        AVR_SHEETS[topology]
    )
    require(ws.title == sheet, "AVR topology sheet mismatch")
    require(kind in {"MCCB", "CONTACTOR"}, "AVR execution unsupported")
    parts = []
    if kind == "MCCB":
        require(rating in DRIVE_RATINGS, "exact motor-drive rating mapping required")
        require(
            quantity == count, "motor breakers/drives must equal switched input count"
        )
        require(
            " ".join(str(ws[f"{label_col}{automatic_row-1}"].value).split())
            == "Автомат",
            "automatic block heading drift",
        )
        require(
            ws[f"{label_col}{drive_row-1}"].value == "Привод",
            "drive block heading drift",
        )
        if rating == 630:
            require(
                isinstance(availability, dict)
                and set(availability)
                == {"rating_a", "available", "confirmed_by", "confirmation_reference"}
                and availability["rating_a"] == 630
                and availability["available"] is True
                and isinstance(availability["confirmed_by"], str)
                and bool(availability["confirmed_by"].strip())
                and isinstance(availability["confirmation_reference"], str)
                and bool(availability["confirmation_reference"].strip()),
                "630A: уточнять наличие привода; bound confirmation required",
            )
        offset = DRIVE_RATINGS.index(rating)
        for role, row in (
            ("AUTOMATIC", automatic_row + offset),
            ("DRIVE", drive_row + offset),
        ):
            parts.append(
                {
                    "role": role,
                    "quantity": str(quantity),
                    **source_cell(
                        ws, f"{label_col}{row}", f"{value_col}{row}", f"{rating}А"
                    ),
                }
            )
    else:
        require(
            rating in CONTACTOR_SOURCE_RATINGS[topology],
            "exact contactor nominal mapping required",
        )
        offset = CONTACTOR_RATINGS.index(rating)
        row = contactor_row + offset
        parts.append(
            {
                "role": "CONTACTOR",
                "quantity": str(quantity),
                **source_cell(ws, f"A{row}", f"B{row}", CONTACTOR_LABELS[offset]),
            }
        )
    # Ordered 'до' rows define disjoint successive work bands.
    ceilings = (
        (95, 100, 160, 250, 400, 630, 800)
        if count == 3
        else (63, 100, 160, 250, 400, 630)
    )
    lower = 0
    matches = []
    for index, upper in enumerate(ceilings):
        if lower < rating <= upper:
            row = index + 2
            lc, vc = ("J", "K") if count == 3 else ("K", "L")
            label = (
                f"АВР 3 ввода до {upper}А с блоком"
                if count == 3
                else f"АВР до {upper}А"
            )
            matches.append(
                {
                    **source_cell(ws, f"{lc}{row}", f"{vc}{row}", label),
                    "lower_exclusive_a": lower,
                    "upper_inclusive_a": upper,
                    "quantity": "1",
                    "role": "AVR_WORK",
                }
            )
        lower = upper
    require(len(matches) == 1, "AVR work rating missing/ambiguous")
    return {
        "topology": topology,
        "input_count": count,
        "execution": kind,
        "rating_a": rating,
        "parts": parts,
        "work": matches[0],
        "material_kzt": str(
            sum(Decimal(p["value"]) * Decimal(p["quantity"]) for p in parts)
        ),
        "decision": AVR_DECISION,
        "policy_binding": policy_binding(),
        "availability": availability if rating == 630 and kind == "MCCB" else None,
    }


def audit_contactor(ws: Any, topology: str, rating: int) -> dict[str, Any]:
    require(
        rating in CONTACTOR_SOURCE_RATINGS[topology],
        "exact contactor nominal mapping required",
    )
    if rating <= (800 if topology == "Г-Г-Д" else 630):
        return avr_price(
            ws, topology, "CONTACTOR", rating, Decimal(AVR_SHEETS[topology][1])
        )
    offset = CONTACTOR_RATINGS.index(rating)
    row = AVR_SHEETS[topology][6] + offset
    return {
        "topology": topology,
        "rating_a": rating,
        "execution": "CONTACTOR",
        "status": "HOLD_MISSING_TOPOLOGY_RATING_WORK_AUTHORITY",
        "material": source_cell(ws, f"A{row}", f"B{row}", CONTACTOR_LABELS[offset]),
        "work": None,
        "fallback_authorized": False,
    }


def audit_sources(selector: Path) -> dict[str, Any]:
    """Reproducible exact SHA audit; excludes undecidable signatures from execution."""
    baseline = resolve_active_price_baseline(selector)
    require(baseline.sha256 == MAIN_SHA, "WU8 authoritative source changed")
    book = load_workbook(
        baseline.path, read_only=True, data_only=False, keep_links=False
    )
    try:
        shr = {category: shr_price(book["ЩР"], category) for category in SHR_BREAKERS}
        excluded = [
            {
                "row": r,
                "label": book["ЩР"].cell(r, 1).value,
                "reason": "HOLD: missing explicit electrical signature",
            }
            for r in (2, 3, 4, 7, 10, 11, 12, 13, 14, 20)
        ]
        avr = [
            avr_price(book[v[0]], topology, "MCCB", rating, Decimal(v[1]))
            for topology, v in AVR_SHEETS.items()
            for rating in DRIVE_RATINGS
            if rating != 630
        ]
        unavailable = [
            {
                "topology": topology,
                "rating_a": 630,
                "status": "HOLD_AVAILABILITY_CONFIRMATION_REQUIRED",
                "automatic": source_cell(
                    book[v[0]], f"{v[2]}{v[4]+4}", f"{v[3]}{v[4]+4}", "630А"
                ),
                "drive": source_cell(
                    book[v[0]], f"{v[2]}{v[5]+4}", f"{v[3]}{v[5]+4}", "630А"
                ),
            }
            for topology, v in AVR_SHEETS.items()
        ]
        contactors = [
            audit_contactor(book[v[0]], topology, rating)
            for topology, v in AVR_SHEETS.items()
            for rating in CONTACTOR_SOURCE_RATINGS[topology]
        ]
        cabinets = [cabinet_cost(book["КРН"], baseline, code) for code in CABINETS]
        from governed_metal_cost import GENESIS_PATH, GENESIS_SHA, _inspect_cost

        for cabinet in cabinets:
            _, _, size, row, fabrication, painting = CABINETS[cabinet["cabinet_code"]]
            identity = metal_identity(
                row,
                " ".join(cabinet["label"].split()) + " ",
                size,
                fabrication,
                painting,
                1500,
            )
            proof = _inspect_cost(GENESIS_PATH, GENESIS_SHA, identity)
            require(
                Decimal(cabinet["cost_kzt"]) == proof["cost_kzt"],
                "cabinet main/metal cost mismatch",
            )
            cabinet["metal_cost_proof"] = {**proof, "identity": identity}
    finally:
        book.close()
    require(
        resolve_active_price_baseline(selector) == baseline,
        "source changed during audit",
    )
    return {
        "status": "READ_ONLY_STRUCTURAL_AUDIT_NOT_CASE_APPROVAL",
        "workbook_path": str(baseline.path),
        "workbook_sha256": baseline.sha256,
        "manifest_sha256": baseline.manifest_sha256,
        "shr_breakers": shr,
        "excluded_breaker_rows": excluded,
        "avr": avr,
        "avr_630a": unavailable,
        "avr_contactors": contactors,
        "cabinets": cabinets,
        "metal": [metal_cost(size) for size in METAL_SIZES],
        "policy_binding": policy_binding(),
    }


if __name__ == "__main__":
    from price_baseline_contract import DEFAULT_ACTIVE_SELECTOR

    print(
        json.dumps(audit_sources(DEFAULT_ACTIVE_SELECTOR), ensure_ascii=False, indent=2)
    )
