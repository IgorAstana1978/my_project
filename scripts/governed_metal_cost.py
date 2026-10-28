"""Read-only D52/D22 cost adapters for exact approved workbook sources.

No materialization or activation entrypoint. Historical D82 resolver is untouched.
"""

from __future__ import annotations

import stat
from decimal import ROUND_UP, Decimal
from pathlib import Path
from typing import Any

from future_case_hardening import number, require
from openpyxl import load_workbook  # type: ignore[import-untyped]
from price_baseline_contract import (
    canonical_json_bytes,
    mapping_identity_fingerprint,
    resolve_active_price_baseline,
    sha256_bytes,
    sha256_file,
    validate_mapping_snapshot,
)

GENESIS_SHA = "b51d7087e0bd8f92e48985294062ead6826c6b50ce3cfacd0f9d0dc22c05f7f2"
GENESIS_PATH = Path(
    r"C:\Users\IgorN\Documents\invoice_quote_filler_data\prices\current"
) / ("прайс_металл_лотки_крышки с 2026.06.18.xlsx")
CURRENT = "CURRENT"
CURRENT_VERSION = "2026.10.21"
METAL_SOURCE_VERSIONS: dict[str, dict[str, Any]] = {
    "2026.06.18": {
        "version": "2026.06.18",
        "path": str(GENESIS_PATH),
        "sha256": GENESIS_SHA,
        "predecessor_version": None,
    },
    "2026.10.21": {
        "version": "2026.10.21",
        "path": str(
            GENESIS_PATH.with_name("прайс_металл_лотки_крышки с 2026.10.21.xlsx")
        ),
        "sha256": "be4a3497cf86ee4e5a13ac55bfeeeb980e6948cd30847ae027f36d6ca599cc71",
        "predecessor_version": "2026.06.18",
    },
}


def resolve_metal_source(version: str = CURRENT) -> dict[str, Any]:
    """Resolve an admitted read-only source; never activate a Case baseline."""
    require(isinstance(version, str), "metal source version required")
    concrete = CURRENT_VERSION if version == CURRENT else version
    require(concrete in METAL_SOURCE_VERSIONS, "unknown metal source version")
    source = dict(METAL_SOURCE_VERSIONS[concrete])
    require(source["version"] == concrete, "mixed metal source version")
    path = Path(str(source["path"]))
    require(path.is_absolute(), "metal source absolute path required")
    for ancestor in (path, *path.parents):
        if ancestor.exists():
            require(
                not ancestor.is_symlink()
                and not getattr(ancestor.lstat(), "st_file_attributes", 0)
                & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0),
                "metal source reparse point",
            )
    require(sha256_file(path) == source["sha256"], "metal source SHA drift")
    return source


LABEL = "1700х800х500 ВРУ "
FORMULAS = {
    "D52": "=ROUNDUP((M52+N52)*R52+O52+P52+S52+Q52,0)",
    "H52": "=E52/1000",
    "I52": "=F52/1000",
    "J52": "=G52/1000",
    "K52": "=H52*I52*3+H52*J52*2+I52*J52*2",
    "L52": "=K52*$B$1*8.42",
    "M52": "=L52*$B$2",
    "N52": "=K52*2*0.25*$B$3*1.33",
    "Q52": "=1250",
}
IDENTITY_VALUES = {
    "B1": 1,
    "E52": 1700,
    "F52": 800,
    "G52": 500,
    "O52": 6750,
    "P52": 3500,
    "R52": 1.05,
    "S52": 2000,
}
IDENTITY = {
    "mapping_id": "METAL-VRU-D52-COST",
    "source_role": "ENCLOSURE_COST",
    "sheet": "Лист1",
    "row": 52,
    "label": LABEL,
    "formulas": FORMULAS,
    "identity_values": IDENTITY_VALUES,
    "price_inputs": ["B2", "B3"],
    "excluded_sales_outputs": ["B52", "C52"],
}
D22_IDENTITY = {
    "mapping_id": "METAL-500X400X250-D22-COST",
    "source_role": "ENCLOSURE_COST",
    "sheet": "Лист1",
    "row": 22,
    "label": "500х400х250 ",
    "formulas": {
        "D22": "=ROUNDUP((M22+N22)*R22+O22+P22+S22+Q22,0)",
        "H22": "=E22/1000",
        "I22": "=F22/1000",
        "J22": "=G22/1000",
        "K22": "=H22*I22*3+H22*J22*2+I22*J22*2",
        "L22": "=K22*$B$1*8.42",
        "M22": "=L22*$B$2",
        "N22": "=K22*2*0.25*$B$3*1.33",
        "Q22": "=1250",
    },
    "identity_values": {
        "B1": 1,
        "E22": 500,
        "F22": 400,
        "G22": 250,
        "O22": 2500,
        "P22": 400,
        "R22": 1.05,
        "S22": 1500,
    },
    "price_inputs": ["B2", "B3"],
    "excluded_sales_outputs": ["B22", "C22"],
}


def inspect_cost(workbook: Path, expected_sha256: str) -> dict[str, Any]:
    return _inspect_cost(workbook, expected_sha256, IDENTITY)


def inspect_cost_d22(workbook: Path, expected_sha256: str) -> dict[str, Any]:
    return _inspect_cost(workbook, expected_sha256, D22_IDENTITY)


def resolve_versioned_metal_cost(
    size: tuple[int, int, int], *, source_version: str = CURRENT
) -> dict[str, Any]:
    """Standalone cost only; legacy Case pricing continues to use GENESIS."""
    from wu8_governed_sources import METAL_SIZES, metal_identity

    require(
        isinstance(size, tuple)
        and len(size) == 3
        and all(type(axis) is int for axis in size)
        and size in METAL_SIZES,
        "metal dimensions outside approved ten-size whitelist",
    )
    source = resolve_metal_source(source_version)
    row, fabrication, painting, extra = METAL_SIZES[size]
    suffix = {52: " ВРУ ", 53: " ПР ", 54: " ПР ", 51: ""}.get(row, " ")
    identity = metal_identity(
        row, "х".join(map(str, size)) + suffix, size, fabrication, painting, extra
    )
    result = _inspect_cost(Path(source["path"]), source["sha256"], identity)
    require(source == resolve_metal_source(source["version"]), "metal source drift")
    return {
        **result,
        "dimensions_mm": list(size),
        "identity": identity,
        "metal_source": source,
    }


def _inspect_cost(
    workbook: Path, expected_sha256: str, identity: dict[str, Any]
) -> dict[str, Any]:
    raw_sha = sha256_file(workbook)
    require(raw_sha == expected_sha256, "metal workbook SHA drift")
    book = load_workbook(workbook, read_only=True, data_only=False, keep_links=False)
    try:
        require(book.sheetnames == ["Лист1"], "metal sheet identity drift")
        s = book["Лист1"]
        row = identity["row"]
        values = identity["identity_values"]
        require(s[f"A{row}"].value == identity["label"], "metal label identity drift")
        for cell, formula in identity["formulas"].items():
            require(s[cell].value == formula, f"metal formula drift: {cell}")
        for cell, value in values.items():
            require(
                number(s[cell].value) == Decimal(str(value)),
                f"metal identity drift: {cell}",
            )
        metal, paint = number(s["B2"].value), number(s["B3"].value)
        h, w, d = (Decimal(values[f"{column}{row}"]) / 1000 for column in "EFG")
        area = h * w * 3 + h * d * 2 + w * d * 2
        mass = area * Decimal("8.42")
        base = (
            (mass * metal + area * 2 * Decimal("0.25") * paint * Decimal("1.33"))
            * Decimal(str(values[f"R{row}"]))
            + values[f"O{row}"]
            + values[f"P{row}"]
            + values[f"S{row}"]
            + 1250
        ).quantize(Decimal(1), rounding=ROUND_UP)
    finally:
        book.close()
    require(sha256_file(workbook) == raw_sha, "metal changed during resolution")
    entry = {
        "mapping_id": identity["mapping_id"],
        "mapping_kind": "cabinet_dynamic",
        "identity": identity,
        "identity_fingerprint": mapping_identity_fingerprint(identity),
        "source": {
            "sheet": "Лист1",
            "row": row,
            "label_cell": f"Лист1!A{row}",
            "price_cells": [f"Лист1!D{row}"],
        },
        "prices": {"price_kzt": int(base)},
    }
    validate_mapping_snapshot([entry])
    return {
        "workbook_path": str(workbook.resolve()),
        "workbook_sha256": raw_sha,
        "structural_fingerprint": mapping_identity_fingerprint(identity),
        "mapping_snapshot": [entry],
        "cost_kzt": int(base),
        "source_role": "ENCLOSURE_COST",
        "source_locator": f"Лист1!D{row}",
        "price_inputs": {"B2": str(metal), "B3": str(paint)},
    }


def resolve_metal_cost(selector: Path) -> dict[str, Any]:
    baseline = resolve_active_price_baseline(selector)
    result = inspect_cost(baseline.path, baseline.sha256)
    require(
        list(baseline.mapping_snapshot) == result["mapping_snapshot"],
        "metal mapping/price drift",
    )
    require(resolve_active_price_baseline(selector) == baseline, "metal selector drift")
    return {
        **result,
        "manifest_sha256": baseline.manifest_sha256,
        "status": "APPROVED_BASELINE_COST_DRAFT",
    }


def workbook_cells(path: Path) -> dict[str, Any]:
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        return {
            s.title: {
                "cells": {
                    c.coordinate: [c.value, c.data_type]
                    for row in s
                    for c in row
                    if c.value is not None
                },
                "merges": sorted(str(r) for r in s.merged_cells.ranges),
            }
            for s in book
        }
    finally:
        book.close()


def audit_metal_candidate(
    candidate: Path,
    expected_sha256: str,
    *,
    selector: Path | None = None,
    predecessor_source: Path | None = None,
    expected_predecessor_sha256: str | None = None,
) -> dict[str, Any]:
    """Price-only candidate audit, never permission to write or activate."""
    predecessor = resolve_active_price_baseline(selector) if selector else None
    if predecessor is None:
        predecessor_path = predecessor_source or GENESIS_PATH
        predecessor_sha = expected_predecessor_sha256 or GENESIS_SHA
        require(
            predecessor_sha == GENESIS_SHA,
            "predecessor needs exact existing genesis metal authority",
        )
    else:
        require(
            predecessor_source is None and expected_predecessor_sha256 is None,
            "selector and genesis predecessor authority cannot mix",
        )
        resolve_metal_cost(selector)  # type: ignore[arg-type]
        predecessor_path, predecessor_sha = predecessor.path, predecessor.sha256
    previous = inspect_cost(predecessor_path, predecessor_sha)
    result = inspect_cost(candidate, expected_sha256)
    before, after = workbook_cells(predecessor_path), workbook_cells(candidate)
    require(set(before) == set(after), "candidate worksheet identity drift")
    changed_inputs = []
    for sheet in before:
        require(
            before[sheet]["merges"] == after[sheet]["merges"],
            "candidate merge identity drift",
        )
        cells_before, cells_after = before[sheet]["cells"], after[sheet]["cells"]
        for cell in cells_before.keys() | cells_after.keys():
            if cells_before.get(cell) == cells_after.get(cell):
                continue
            require(
                sheet == "Лист1" and cell in {"B2", "B3"},
                f"candidate outside price-only footprint: {sheet}!{cell}",
            )
            changed_inputs.append(
                {
                    "cell": cell,
                    "old": cells_before.get(cell),
                    "new": cells_after.get(cell),
                }
            )
    require(
        sha256_file(predecessor_path) == predecessor_sha
        and sha256_file(candidate) == expected_sha256,
        "metal predecessor/candidate changed during audit",
    )
    changes = (
        []
        if previous["cost_kzt"] == result["cost_kzt"]
        else [{"old_kzt": previous["cost_kzt"], "new_kzt": result["cost_kzt"]}]
    )
    payload = {
        "schema_version": "metal_cost_candidate_audit.v0.1",
        "status": "HOLD_EXACT_HUMAN_BASELINE_APPROVAL_REQUIRED",
        "predecessor": (
            {
                "workbook_path": previous["workbook_path"],
                "workbook_sha256": predecessor_sha,
                "manifest_id": predecessor.manifest_id if predecessor else None,
                "manifest_sha256": predecessor.manifest_sha256 if predecessor else None,
            }
        ),
        **result,
        "cost_changes": changes,
        "price_input_changes": sorted(
            changed_inputs, key=lambda change: change["cell"]
        ),
        "future_cases_only": True,
        "historical_repricing_authorized": False,
    }
    if predecessor:
        require(
            resolve_active_price_baseline(selector) == predecessor,  # type: ignore[arg-type]
            "audit predecessor drift",
        )
    return {
        **payload,
        "approval_fingerprint": sha256_bytes(canonical_json_bytes(payload)),
    }
