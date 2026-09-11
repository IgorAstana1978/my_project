"""Build a deterministic DRAFT DINVA classic dynamic profile v0.2."""

from __future__ import annotations

import argparse
import base64
import copy
import hashlib
import json
import os
import re
import stat
import tempfile
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any, NoReturn, cast
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]

REPO_ROOT = Path(__file__).resolve().parents[1]
PROFILE_SCHEMA_VERSION = "dinva_classic_presentation_profile.v0.2"
PROFILE_ID = "DINVA_CLASSIC_QUOTE_INVOICE_V0_1_DYNAMIC_V0_2"
DOCUMENT_FAMILY = "DINVA_CLASSIC_QUOTE_INVOICE_V0_1"
CONTRACT_VERSION = "dinva_classic_presentation_contract.v0.2"
PREDECESSOR_PROFILE_SHA256 = (
    "3c3c448c268e2bc87aa9255720e26e7d26fd5f3faa0bd0a42704ad8f69e3f3ca"
)
CANONICAL_LOGO_DECISION_SHA256 = (
    "e7c043f19b7eb8606f59dd8e7de06b29ca4305cc1fe2362ecb93767dd589f63b"
)
CANONICAL_LOGO_RAW_SHA256 = (
    "28a6a59ae0a5ca274c206c70545f70b333cac0276a7c4dcbebbf9156f88e0fa8"
)
CANONICAL_LOGO_PIXEL_FINGERPRINT = (
    "81d979c4c158452cca8e3b40d23a4fd321538dfcef238b6f8133beb33a122846"
)
REQUIRED_FAMILY_SHA256S = frozenset(
    {
        "8cf9f2b4ecca94e51a9f868891b6bc00151ef4b05b012db0d875862599c5253c",
        "17e31d0312f728800d31fd4f125d285edb1114880500d4833261239b87ab58b5",
        "d8e652325c142a72ffa4aa390197b3e357b5efc317d07f6d763e01d3c1c4fec9",
    }
)
EXPECTED_FIRST_ITEM_ROWS = [16, 16, 17]
APPROVED_V02_SHA256 = "93c41c9f0399e6a6d27e1b075ceb2ab971b345f6c64143c637ac1e2ad02a90b7"
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
FIXED_HEADERS = {
    "B": "№ п/п",
    "C": "Наименование",
    "D": "Ед.",
    "E": "Кол-во",
    "H": "Цена",
    "I": "Сумма",
}


class ProfileV02ExtractionError(ValueError):
    """The governed evidence cannot produce the one v0.2 DRAFT profile."""


@dataclass(frozen=True)
class BoundInput:
    path: Path
    expected_sha256: str


@dataclass(frozen=True)
class LoadedInput:
    path: Path
    raw: bytes
    sha256: str


@dataclass(frozen=True)
class FamilyEvidence:
    path: Path
    sha256: str
    first_item_row: int


def fail(message: str) -> NoReturn:
    raise ProfileV02ExtractionError(message)


def require(condition: bool, message: str) -> None:
    if not condition:
        fail(message)


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def canonical_json(value: object) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")


def reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        require(key not in result, f"duplicate JSON key: {key}")
        result[key] = value
    return result


def load_strict_json(raw: bytes, label: str) -> dict[str, Any]:
    try:
        value = json.loads(
            raw.decode("utf-8", errors="strict"),
            object_pairs_hook=reject_duplicate_keys,
        )
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ProfileV02ExtractionError(
            f"{label} is not strict UTF-8 JSON: {exc}"
        ) from exc
    require(isinstance(value, Mapping), f"{label} root must be an object")
    return dict(cast(Mapping[str, Any], value))


def mapping(value: object, label: str) -> Mapping[str, Any]:
    require(isinstance(value, Mapping), f"{label} must be an object")
    return cast(Mapping[str, Any], value)


def load_bound(source: BoundInput, label: str) -> LoadedInput:
    require(
        SHA256_RE.fullmatch(source.expected_sha256) is not None,
        f"{label} SHA-256 format invalid",
    )
    path = source.path.resolve(strict=True)
    require(
        not path.is_relative_to(REPO_ROOT.resolve(strict=False)),
        f"{label} must be outside Git",
    )
    raw = path.read_bytes()
    digest = sha256_bytes(raw)
    require(digest == source.expected_sha256, f"{label} SHA-256 mismatch")
    require(path.read_bytes() == raw, f"{label} changed during validation")
    return LoadedInput(path, raw, digest)


def normalized(value: object) -> str:
    return re.sub(r"\s+", " ", value.strip()) if isinstance(value, str) else ""


def load_family(source: BoundInput) -> FamilyEvidence:
    loaded = load_bound(source, "classic family evidence")
    require(
        loaded.sha256 in REQUIRED_FAMILY_SHA256S,
        "classic family evidence is not one of the exact 463/519/551 artifacts",
    )
    try:
        workbook = load_workbook(
            loaded.path, data_only=False, read_only=False, keep_links=True
        )
    except (OSError, ValueError, BadZipFile) as exc:
        raise ProfileV02ExtractionError(
            f"classic family evidence open failed: {exc}"
        ) from exc
    try:
        require(
            workbook.index(workbook.active) == 0, "classic active sheet is not first"
        )
        sheet = workbook.active
        require(normalized(sheet["C2"].value) == "ТОО «ДиН ВА-КЭС»", "family mismatch")
        require(sheet["G9"].value == "ВНИМАНИЕ!", "classic warning mismatch")
        header_rows = [
            row
            for row in range(1, 31)
            if all(
                normalized(sheet[f"{column}{row}"].value) == expected
                for column, expected in FIXED_HEADERS.items()
            )
        ]
        require(header_rows == [15], "classic table header geometry mismatch")
        first_item = next(
            (
                row
                for row in range(16, sheet.max_row + 1)
                if type(sheet.cell(row, 2).value) is int
            ),
            None,
        )
        require(type(first_item) is int, "classic first item row missing")
        require(
            sheet.page_setup.orientation == "portrait", "classic orientation mismatch"
        )
        require(str(sheet.page_setup.paperSize) == "9", "classic paper size mismatch")
        require(sheet.page_setup.scale == 54, "classic scale mismatch")
    finally:
        workbook.close()
    require(loaded.path.read_bytes() == loaded.raw, "classic family evidence changed")
    return FamilyEvidence(loaded.path, loaded.sha256, cast(int, first_item))


def style(
    size: float,
    *,
    bold: bool = False,
    italic: bool = False,
    underline: str | None = None,
    horizontal: str | None = None,
    vertical: str | None = None,
    wrap: bool = False,
    number_format: str = "General",
    left: str | None = None,
    right: str | None = None,
    top: str | None = None,
    bottom: str | None = None,
    fill: str | None = None,
    color: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "font": {
            "name": "Times New Roman",
            "size": size,
            "bold": bold,
            "italic": italic,
            "underline": underline,
            "color": color,
        },
        "fill": {
            "type": "solid" if fill else None,
            "foreground": {
                "type": "rgb",
                "value": fill or "00000000",
                "tint": 0.0,
            },
        },
        "border": {"left": left, "right": right, "top": top, "bottom": bottom},
        "alignment": {
            "horizontal": horizontal,
            "vertical": vertical,
            "wrap_text": wrap,
            "shrink_to_fit": False,
        },
        "number_format": number_format,
    }


def dynamic_styles() -> dict[str, dict[str, Any]]:
    red = {"type": "rgb", "value": "FFFF0000", "tint": 0.0}
    body: dict[str, Any] = {
        "vertical": "center",
        "wrap": True,
        "top": "thin",
        "bottom": "thin",
    }
    return {
        "company_title": style(20, bold=True, horizontal="left", top="medium"),
        "country": style(16, bold=True, horizontal="left"),
        "company_info": style(14, bold=True, horizontal="left"),
        "bank_title": style(
            20, bold=True, horizontal="right", right="medium", top="medium"
        ),
        "bank_info": style(14, bold=True, horizontal="right", right="medium"),
        "document_title": style(20, bold=True, top="medium"),
        "payer": style(20, bold=True, horizontal="left"),
        "warning": style(14, color=red, top="medium"),
        "lead_time": style(14, bold=True, underline="single", color=red),
        "object": style(14, bold=True, horizontal="left", vertical="center", wrap=True),
        "table_header": style(
            16,
            bold=True,
            horizontal="center",
            vertical="center",
            wrap=True,
            left="thin",
            right="thin",
            top="medium",
            bottom="medium",
            fill="FFBFBFBF",
        ),
        "table_header_left": style(
            16,
            bold=True,
            horizontal="center",
            vertical="center",
            wrap=True,
            left="medium",
            right="thin",
            top="medium",
            bottom="medium",
            fill="FFBFBFBF",
        ),
        "table_header_right": style(
            16,
            bold=True,
            horizontal="center",
            vertical="center",
            wrap=True,
            left="thin",
            right="medium",
            top="medium",
            bottom="medium",
            fill="FFBFBFBF",
        ),
        "section": style(
            14,
            bold=True,
            horizontal="left",
            vertical="center",
            left="thin",
            right="thin",
            top="medium",
            bottom="medium",
            fill="FFD9D9D9",
        ),
        "position": style(14, horizontal="center", left="medium", right="thin", **body),
        "item_name": style(14, horizontal="left", left="thin", right="thin", **body),
        "unit": style(14, horizontal="center", left="thin", right="thin", **body),
        "quantity": style(14, horizontal="center", left="thin", right="thin", **body),
        "technical_composition": style(
            14,
            horizontal="left",
            vertical="top",
            wrap=True,
            left="thin",
            right="thin",
            top="thin",
            bottom="thin",
        ),
        "enclosure": style(14, horizontal="center", left="thin", right="thin", **body),
        "money": style(
            14,
            horizontal="center",
            number_format="#,##0",
            left="thin",
            right="thin",
            **body,
        ),
        "line_total": style(
            14,
            horizontal="center",
            number_format="#,##0",
            left="thin",
            right="medium",
            **body,
        ),
        "total_label": style(16, bold=True),
        "total_amount": style(14, bold=True, number_format="#,##0_);(#,##0)"),
        "amount_words": style(16, bold=True, horizontal="left"),
        "commercial_line": style(16, bold=True, horizontal="left"),
        "director": style(16, horizontal="right"),
        "director_name": style(16, horizontal="center", vertical="center"),
        "executor": style(8, horizontal="left"),
    }


def validate_predecessor(
    source: BoundInput, family: Sequence[FamilyEvidence]
) -> tuple[LoadedInput, dict[str, Any]]:
    require(
        source.expected_sha256 == PREDECESSOR_PROFILE_SHA256,
        "approved predecessor profile SHA is not the immutable v0.1 SHA",
    )
    loaded = load_bound(source, "approved predecessor profile")
    profile = load_strict_json(loaded.raw, "approved predecessor profile")
    approval = mapping(profile.get("approval_provenance"), "predecessor approval")
    require(
        profile.get("schema_version") == "dinva_classic_presentation_profile.v0.1"
        and profile.get("artifact_status") == "IMMUTABLE_APPROVED_PROFILE"
        and approval.get("status") == "APPROVED"
        and approval.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL",
        "approved predecessor profile state mismatch",
    )
    contract = mapping(profile.get("presentation_contract"), "predecessor contract")
    require(
        profile.get("presentation_contract_fingerprint")
        == sha256_bytes(canonical_json(contract)),
        "approved predecessor profile fingerprint mismatch",
    )
    provenance = profile.get("reference_provenance")
    require(isinstance(provenance, list), "predecessor provenance missing")
    predecessor_family = {
        mapping(item, "predecessor provenance").get("actual_sha256")
        for item in cast(list[Any], provenance)
        if mapping(item, "predecessor provenance").get("role")
        == "CLASSIC_FAMILY_EVIDENCE"
    }
    require(
        predecessor_family == {item.sha256 for item in family},
        "predecessor family evidence set mismatch",
    )
    return loaded, profile


def validate_logo_decision(
    source: BoundInput, family: Sequence[FamilyEvidence]
) -> LoadedInput:
    require(
        source.expected_sha256 == CANONICAL_LOGO_DECISION_SHA256,
        "canonical-logo decision SHA is not the approved SHA",
    )
    loaded = load_bound(source, "canonical-logo Human Decision")
    decision = load_strict_json(loaded.raw, "canonical-logo Human Decision")
    canonical = mapping(decision.get("canonical_logo_decision"), "canonical logo")
    require(
        decision.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL"
        and canonical.get("raw_sha256") == CANONICAL_LOGO_RAW_SHA256
        and canonical.get("normalized_pixel_fingerprint")
        == CANONICAL_LOGO_PIXEL_FINGERPRINT,
        "canonical-logo decision contract mismatch",
    )
    bindings = decision.get("source_bindings")
    require(isinstance(bindings, list), "canonical-logo source bindings missing")
    bound_family = {
        mapping(item, "canonical-logo binding").get("actual_workbook_sha256")
        for item in cast(list[Any], bindings)
        if mapping(item, "canonical-logo binding").get("role")
        == "CLASSIC_FAMILY_EVIDENCE"
    }
    require(
        bound_family == {item.sha256 for item in family},
        "canonical-logo family binding mismatch",
    )
    return loaded


def build_contract(
    predecessor: Mapping[str, Any], family: Sequence[FamilyEvidence]
) -> dict[str, Any]:
    old = mapping(predecessor.get("presentation_contract"), "predecessor contract")
    old_fixed = mapping(old.get("fixed_blocks"), "predecessor fixed blocks")
    old_company = mapping(old_fixed.get("company"), "predecessor company")
    company = {
        "C2": "                     " + cast(str, old_company["C2"]),
        "C3": "                     " + cast(str, old_company["C3"]),
        "B4": old_company["B4"],
        "B5": old_company["B5"],
        "B6": old_company["B6"],
        "I2": old_company["G2"],
        "I3": old_company["G3"],
        "I4": old_company["G4"],
        "I5": old_company["G5"],
        "I6": old_company["G6"],
    }
    old_asset = mapping(cast(list[Any], old["assets"])[0], "predecessor logo")
    require(
        old_asset.get("sha256") == CANONICAL_LOGO_RAW_SHA256,
        "predecessor logo mismatch",
    )
    logo = base64.b64decode(cast(str, old_asset["data_base64"]), validate=True)
    require(
        sha256_bytes(logo) == CANONICAL_LOGO_RAW_SHA256, "canonical logo bytes mismatch"
    )
    old_print = copy.deepcopy(dict(mapping(old["print"], "predecessor print")))
    old_print.update({"fit_to_page": False, "fit_to_width": 0})
    return {
        "contract_version": CONTRACT_VERSION,
        "workbook": {
            "active_sheet_index": 0,
            "extra_sheets_allowed": False,
            "sheets": [{"name": "Счёт-КП", "role": "PRIMARY_DOCUMENT"}],
        },
        "fixed_blocks": {
            "company": company,
            "warning": copy.deepcopy(old_fixed["warning"]),
            "table_headers": copy.deepcopy(old_fixed["table_headers"]),
            "total_label": old_fixed["total_label"],
        },
        "layout": {
            "table_header_row": 15,
            "first_content_row": 16,
            "table_columns": copy.deepcopy(
                mapping(old["layout"], "old layout")["table_columns"]
            ),
            "column_width_rules": {
                "B": {
                    "minimum": 5.42578125,
                    "preferred": 5.42578125,
                    "maximum": 5.42578125,
                    "adaptive": False,
                    "preferred_chars": 1,
                    "maximum_chars": 1,
                },
                "C": {
                    "minimum": 36.0,
                    "preferred": 42.0,
                    "maximum": 44.0,
                    "adaptive": True,
                    "preferred_chars": 24,
                    "maximum_chars": 60,
                },
                "D": {
                    "minimum": 8.140625,
                    "preferred": 8.140625,
                    "maximum": 8.140625,
                    "adaptive": False,
                    "preferred_chars": 1,
                    "maximum_chars": 1,
                },
                "E": {
                    "minimum": 12.5703125,
                    "preferred": 12.5703125,
                    "maximum": 12.5703125,
                    "adaptive": False,
                    "preferred_chars": 1,
                    "maximum_chars": 1,
                },
                "F": {
                    "minimum": 48.0,
                    "preferred": 56.140625,
                    "maximum": 58.0,
                    "adaptive": True,
                    "preferred_chars": 90,
                    "maximum_chars": 260,
                },
                "G": {
                    "minimum": 18.0,
                    "preferred": 19.85546875,
                    "maximum": 22.0,
                    "adaptive": True,
                    "preferred_chars": 26,
                    "maximum_chars": 55,
                },
                "H": {
                    "minimum": 14.42578125,
                    "preferred": 14.42578125,
                    "maximum": 14.42578125,
                    "adaptive": False,
                    "preferred_chars": 1,
                    "maximum_chars": 1,
                },
                "I": {
                    "minimum": 15.5703125,
                    "preferred": 18.0,
                    "maximum": 22.140625,
                    "adaptive": True,
                    "preferred_chars": 7,
                    "maximum_chars": 12,
                },
            },
            "maximum_printable_width": 180.75,
            "row_height_rule": {
                "minimum_height": 37.5,
                "line_height_points": 18.0,
                "vertical_padding_points": 3.0,
                "maximum_height": 408.0,
                "header_minimum_height": 82.0,
                "header_line_height_points": 20.0,
                "section_height": 19.5,
                "total_height": 20.25,
                "amount_height": 20.25,
                "commercial_height": 20.25,
                "signature_height": 20.25,
            },
            "bottom_layout": {"amount_words_offset": 2, "signature_spacer_rows": 1},
            "pagination": {
                "first_page_body_height_points": 900.0,
                "following_page_body_height_points": 1100.0,
                "repeat_rows": "15:15",
                "keep_section_with_first_item": True,
            },
            "gridlines_visible": True,
            "merged_cells": {"mode": "NONE", "ranges": []},
            "top_row_heights": {
                "2": 25.5,
                "3": 25.5,
                "9": 25.5,
                "10": 25.5,
                "13": 63.75,
            },
        },
        "styles": dynamic_styles(),
        "formulas": copy.deepcopy(old["formulas"]),
        "assets": [
            {
                "asset_id": "DINVA_CLASSIC_LOGO_V0_1",
                "media_type": "image/png",
                "sha256": CANONICAL_LOGO_RAW_SHA256,
                "data_base64": base64.b64encode(logo).decode("ascii"),
                "source_reference_sha256s": sorted(item.sha256 for item in family),
                "placement": {
                    "anchor_type": "TWO_CELL",
                    "relative_to": "COMPANY_HEADER_BLOCK",
                    "base_cell": "B2",
                    "from": {
                        "column_delta": 0,
                        "column_offset": 57150,
                        "row_delta": 0,
                        "row_offset": 76200,
                    },
                    "to": {
                        "column_delta": 1,
                        "column_offset": 1200150,
                        "row_delta": 1,
                        "row_offset": 266700,
                    },
                    "protected_first_row": 4,
                },
            }
        ],
        "package": copy.deepcopy(old["package"]),
        "print": old_print,
        "optional_elements": copy.deepcopy(old["optional_elements"]),
        "variable_elements": [
            "item count/order",
            "sections",
            "content",
            "commercial lines",
        ],
    }


def extract_profile_v0_2(
    family_sources: Sequence[BoundInput],
    predecessor_source: BoundInput,
    canonical_logo_decision_source: BoundInput,
) -> dict[str, Any]:
    require(
        len(family_sources) == 3,
        "exactly three 463/519/551 family references are required",
    )
    family = [load_family(source) for source in family_sources]
    require(
        {item.sha256 for item in family} == set(REQUIRED_FAMILY_SHA256S),
        "exact 463/519/551 family evidence set mismatch",
    )
    require(
        sorted(item.first_item_row for item in family) == EXPECTED_FIRST_ITEM_ROWS,
        "classic family first-item geometry mismatch",
    )
    predecessor_loaded, predecessor = validate_predecessor(predecessor_source, family)
    decision_loaded = validate_logo_decision(canonical_logo_decision_source, family)
    contract = build_contract(predecessor, family)
    fingerprint = sha256_bytes(canonical_json(contract))
    provenance = [
        {
            "path": str(item.path),
            "expected_sha256": item.sha256,
            "actual_sha256": item.sha256,
            "role": "CLASSIC_FAMILY_EVIDENCE",
        }
        for item in family
    ] + [
        {
            "path": str(predecessor_loaded.path),
            "expected_sha256": predecessor_loaded.sha256,
            "actual_sha256": predecessor_loaded.sha256,
            "role": "IMMUTABLE_APPROVED_PREDECESSOR_PROFILE",
        },
        {
            "path": str(decision_loaded.path),
            "expected_sha256": decision_loaded.sha256,
            "actual_sha256": decision_loaded.sha256,
            "role": "CANONICAL_LOGO_HUMAN_DECISION",
        },
    ]
    provenance.sort(
        key=lambda item: (item["role"], item["actual_sha256"], item["path"])
    )
    profile = {
        "schema_version": PROFILE_SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "document_family": DOCUMENT_FAMILY,
        "artifact_status": "DRAFT_PROFILE_CANDIDATE",
        "reference_provenance": provenance,
        "presentation_contract": contract,
        "presentation_contract_fingerprint": fingerprint,
        "approval_provenance": {
            "status": "DRAFT_UNAPPROVED",
            "authority": None,
            "approval_id": None,
            "approved_at": None,
            "approved_contract_fingerprint": None,
        },
    }
    for source in (*family_sources, predecessor_source, canonical_logo_decision_source):
        require(
            sha256_bytes(source.path.resolve(strict=True).read_bytes())
            == source.expected_sha256,
            "source TOCTOU mismatch",
        )
    return profile


def family_logo_drawing(
    family: Sequence[BoundInput], asset: Mapping[str, Any]
) -> dict[str, Any]:
    """Preserve family DrawingML, normalizing only the package-local picture ID.

    Markers cannot determine the explicit shape transform, DPI or picture
    extensions. These are evidence, not renderer-derived geometry.
    """
    ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    drawings = []
    for source in family:
        loaded = load_bound(source, "logo family evidence")
        with ZipFile(source.path) as archive:
            root = ElementTree.fromstring(archive.read("xl/drawings/drawing1.xml"))
            rels = ElementTree.fromstring(
                archive.read("xl/drawings/_rels/drawing1.xml.rels")
            )
            raw = archive.read("xl/media/image1.png")
        require(
            sha256_bytes(source.path.read_bytes()) == loaded.sha256,
            "logo evidence TOCTOU mismatch",
        )
        require(
            len(root) == 1 and root[0].tag == f"{{{ns['xdr']}}}twoCellAnchor",
            "family logo anchor count/type differs",
        )
        anchor = root[0]
        require(anchor.attrib == {"editAs": "absolute"}, "family logo editAs differs")
        require(
            len(rels) == 1
            and rels[0].attrib
            == {
                "Id": "rId1",
                "Type": ns["r"] + "/image",
                "Target": "../media/image1.png",
            },
            "family logo relationship differs",
        )
        require(
            raw == base64.b64decode(str(asset["data_base64"]), validate=True)
            and sha256_bytes(raw) == asset["sha256"]
            and raw[:8] == b"\x89PNG\r\n\x1a\n"
            and raw[12:16] == b"IHDR",
            "family logo native bytes differ",
        )
        placement = mapping(asset["placement"], "logo placement")
        for name in ("from", "to"):
            marker = mapping(placement[name], "logo marker")
            expected = {
                "col": 1 + marker["column_delta"],
                "colOff": marker["column_offset"],
                "row": 1 + marker["row_delta"],
                "rowOff": marker["row_offset"],
            }
            require(
                all(
                    anchor.findtext(f"xdr:{name}/xdr:{key}", namespaces=ns)
                    == str(value)
                    for key, value in expected.items()
                ),
                "family logo markers differ from predecessor",
            )
        pic = anchor.find("xdr:pic", ns)
        require(pic is not None, "family logo picture missing")
        pic = cast(ElementTree.Element, pic)
        identity = pic.find("xdr:nvPicPr/xdr:cNvPr", ns)
        require(identity is not None, "family logo identity missing")
        cast(ElementTree.Element, identity).set("id", "1")
        for path in (
            "xdr:spPr/a:xfrm/a:off",
            "xdr:spPr/a:xfrm/a:ext",
            "xdr:blipFill/a:srcRect",
            "xdr:blipFill/a:stretch/a:fillRect",
            "xdr:blipFill/a:blip/a:extLst/a:ext/a14:useLocalDpi",
        ):
            require(
                pic.find(path, ns) is not None, f"family logo geometry missing: {path}"
            )
        for prefix, uri in ns.items():
            ElementTree.register_namespace(prefix, uri)
        # Compare the whole anchor, including locks and hidden shape extensions.
        normalized = ElementTree.canonicalize(
            ElementTree.tostring(anchor, encoding="unicode")
        )
        drawings.append((normalized, ElementTree.tostring(pic, encoding="unicode")))
    require(
        bool(drawings) and all(d == drawings[0] for d in drawings),
        "family logo DrawingML semantics differ",
    )
    return {
        "edit_as": "absolute",
        "picture_xml": drawings[0][1],
        "native_width_px": int.from_bytes(raw[16:20], "big"),
        "native_height_px": int.from_bytes(raw[20:24], "big"),
        "source_sha256s": sorted(s.expected_sha256 for s in family),
        "source_locator": "xl/drawings/drawing1.xml;xl/media/image1.png",
        "normalization": "cNvPr/@id=1 (package-local identity only)",
    }


def extract_profile_v0_3(
    approved_source: BoundInput,
) -> dict[str, Any]:
    """Derive a visual successor from immutable approval and all three sources.

    Top frame coordinates belong to the family header, never an item-count
    template. Borders (including colors) must agree byte-for-byte across 463,
    519 and 551; disagreement requires a new evidence decision.
    """
    require(
        approved_source.expected_sha256 == APPROVED_V02_SHA256,
        "v0.2 predecessor SHA mismatch",
    )
    loaded = load_bound(approved_source, "approved v0.2 predecessor")
    profile = load_strict_json(loaded.raw, "approved v0.2 predecessor")
    approval = mapping(profile.get("approval_provenance"), "approval")
    contract = mapping(profile.get("presentation_contract"), "contract")
    fingerprint = sha256_bytes(canonical_json(contract))
    require(
        profile.get("schema_version") == PROFILE_SCHEMA_VERSION
        and profile.get("profile_id") == PROFILE_ID
        and profile.get("artifact_status") == "IMMUTABLE_APPROVED_PROFILE"
        and approval.get("status") == "APPROVED"
        and approval.get("authority") == "IGOR_DIRECT_HUMAN_APPROVAL"
        and approval.get("approved_contract_fingerprint") == fingerprint
        and profile.get("presentation_contract_fingerprint") == fingerprint,
        "v0.2 predecessor approval mismatch",
    )
    bindings = cast(list[dict[str, Any]], profile["reference_provenance"])
    sources = [BoundInput(Path(b["path"]), b["actual_sha256"]) for b in bindings]
    for binding in bindings:
        require(
            binding["expected_sha256"] == binding["actual_sha256"],
            "source binding mismatch",
        )
    for source in sources:
        load_bound(source, "predecessor evidence")
    family = [
        source
        for source, binding in zip(sources, bindings, strict=True)
        if binding["role"] == "CLASSIC_FAMILY_EVIDENCE"
    ]
    require(
        len(family) == 3
        and {s.expected_sha256 for s in family} == set(REQUIRED_FAMILY_SHA256S),
        "visual family evidence set mismatch",
    )
    border_maps = []
    for source in family:
        load_family(source)
        workbook = load_workbook(source.path, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            border_maps.append(
                {
                    cell.coordinate: ElementTree.tostring(
                        cell.border.to_tree(), encoding="unicode"
                    )
                    for row in sheet.iter_rows(
                        min_row=2, max_row=14, min_col=2, max_col=9
                    )
                    for cell in row
                }
            )
        finally:
            workbook.close()
    require(
        all(value == border_maps[0] for value in border_maps),
        "family top-border geometry differs",
    )
    successor = copy.deepcopy(profile)
    successor["schema_version"] = "dinva_classic_presentation_profile.v0.3"
    successor["profile_id"] = "DINVA_CLASSIC_QUOTE_INVOICE_V0_1_DYNAMIC_V0_3"
    successor["artifact_status"] = "DRAFT_PROFILE_CANDIDATE"
    new = successor["presentation_contract"]
    new["contract_version"] = "dinva_classic_presentation_contract.v0.3"
    new["layout"]["top_block_rules"] = {
        "object_height": {
            "cell": "C13",
            "glyph_width_em": 1.1,
            "line_height_em": 1.5,
            "padding_points": 6.0,
            "maximum_height": 408.0,
        },
        "border_xml_by_cell": border_maps[0],
        "border_source_sha256s": sorted(s.expected_sha256 for s in family),
        "border_source_locator": "Лист1!B2:I14",
    }
    new["assets"][0]["drawing_semantics"] = family_logo_drawing(
        family, new["assets"][0]
    )
    successor["presentation_contract_fingerprint"] = sha256_bytes(canonical_json(new))
    successor["approval_provenance"] = {
        "status": "DRAFT_UNAPPROVED",
        "authority": None,
        "approval_id": None,
        "approved_at": None,
        "approved_contract_fingerprint": None,
    }
    successor["reference_provenance"].append(
        {
            "path": str(loaded.path),
            "expected_sha256": loaded.sha256,
            "actual_sha256": loaded.sha256,
            "role": "IMMUTABLE_APPROVED_V0_2_PREDECESSOR_PROFILE",
        }
    )
    for source in [approved_source, *sources]:
        load_bound(source, "successor final source reread")
    return successor


def publish_draft_profile(
    profile: Mapping[str, Any],
    output: Path,
    *,
    verify_sources: Callable[[], None] | None = None,
) -> Path:
    output = output.resolve(strict=False)
    require(
        output.name
        == (
            (
                "dinva-classic-presentation-profile-v0.5-DRAFT.json"
                if profile.get("schema_version")
                == "dinva_classic_presentation_profile.v0.5"
                else "dinva-classic-presentation-profile-v0.4-DRAFT.json"
            )
            if profile.get("schema_version")
            in {
                "dinva_classic_presentation_profile.v0.4",
                "dinva_classic_presentation_profile.v0.5",
            }
            else (
                "dinva-classic-presentation-profile-v0.3-DRAFT.json"
                if profile.get("schema_version")
                == "dinva_classic_presentation_profile.v0.3"
                else "dinva-classic-presentation-profile-v0.2-DRAFT.json"
            )
        ),
        "profile output filename mismatch",
    )
    require(
        not output.is_relative_to(REPO_ROOT.resolve(strict=False)),
        "profile output must be outside Git",
    )
    require(output.parent.parent.is_dir(), "output directory owner must already exist")
    require(
        not output.parent.exists() and not output.exists(),
        "new profile output path already exists",
    )
    raw = (json.dumps(profile, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if verify_sources is not None:
        verify_sources()
    output.parent.mkdir()
    descriptor = -1
    candidate: Path | None = None
    final_link_created = False
    candidate_identity: tuple[int, int] | None = None
    try:
        descriptor, raw_candidate = tempfile.mkstemp(
            prefix=".dinva-profile-v02-draft-", suffix=".tmp", dir=output.parent
        )
        candidate = Path(raw_candidate)
        os.chmod(candidate, stat.S_IRUSR | stat.S_IWUSR)
        with os.fdopen(descriptor, "wb") as stream:
            descriptor = -1
            stream.write(raw)
            stream.flush()
            os.fsync(stream.fileno())
        require(candidate.read_bytes() == raw, "profile candidate reread mismatch")
        metadata = candidate.stat()
        candidate_identity = (metadata.st_dev, metadata.st_ino)
        if verify_sources is not None:
            verify_sources()
        os.link(candidate, output)
        final_link_created = True
        final_metadata = output.stat()
        require(
            (final_metadata.st_dev, final_metadata.st_ino) == candidate_identity,
            "profile final identity mismatch",
        )
        require(output.read_bytes() == raw, "profile final reread mismatch")
        if verify_sources is not None:
            verify_sources()
        candidate.unlink()
        require(
            set(output.parent.iterdir()) == {output},
            "profile final directory inventory mismatch",
        )
        return output
    except BaseException:
        if descriptor >= 0:
            os.close(descriptor)
        if final_link_created and candidate_identity is not None and output.exists():
            metadata = output.stat()
            if (metadata.st_dev, metadata.st_ino) == candidate_identity:
                output.unlink()
        if candidate is not None and candidate.exists():
            candidate.unlink()
        if output.parent.exists() and not any(output.parent.iterdir()):
            output.parent.rmdir()
        raise


def paired(paths: list[Path], hashes: list[str]) -> list[BoundInput]:
    require(len(paths) == len(hashes), "family path/SHA count mismatch")
    return [
        BoundInput(path, digest) for path, digest in zip(paths, hashes, strict=True)
    ]


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--family-reference", action="append", required=True, type=Path)
    parser.add_argument("--family-reference-sha256", action="append", required=True)
    parser.add_argument("--approved-v0-1-profile", required=True, type=Path)
    parser.add_argument("--approved-v0-1-profile-sha256", required=True)
    parser.add_argument("--canonical-logo-human-decision", required=True, type=Path)
    parser.add_argument("--canonical-logo-human-decision-sha256", required=True)
    parser.add_argument("--output-profile", required=True, type=Path)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        profile = extract_profile_v0_2(
            paired(args.family_reference, args.family_reference_sha256),
            BoundInput(args.approved_v0_1_profile, args.approved_v0_1_profile_sha256),
            BoundInput(
                args.canonical_logo_human_decision,
                args.canonical_logo_human_decision_sha256,
            ),
        )
        output = publish_draft_profile(profile, args.output_profile)
    except (OSError, ProfileV02ExtractionError) as exc:
        print(f"HOLD: {exc}")
        return 1
    raw = output.read_bytes()
    print("DINVA_CLASSIC_PROFILE_V0_2=DRAFT_UNAPPROVED")
    print(
        f"PRESENTATION_CONTRACT_FINGERPRINT={profile['presentation_contract_fingerprint']}"
    )
    print(f"SHA256={sha256_bytes(raw)}")
    print(f"SIZE={len(raw)}")
    print(f"OUTPUT={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
