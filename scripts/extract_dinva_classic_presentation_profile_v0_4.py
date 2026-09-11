"""Build only a source-bound v0.4 DRAFT; no approval or rendering capability."""

from __future__ import annotations

import argparse
import copy
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from dinva_native_text import canonical_spill_fit
from extract_dinva_classic_presentation_profile_v0_2 import (
    REQUIRED_FAMILY_SHA256S,
    BoundInput,
    canonical_json,
    load_bound,
    load_strict_json,
    publish_draft_profile,
    require,
    sha256_bytes,
)
from extract_dinva_classic_presentation_profile_v0_2 import (
    ProfileV02ExtractionError as ProfileV02ExtractionError,
)
from openpyxl import load_workbook  # type: ignore[import-untyped]

APPROVED_V03_SHA256 = "4fab7cb8898dcff55c75bd94b39cc1d8357a73fa59fbd4a0dc52dee5a45c67a8"
COUNTRY_FONT_SHA256 = "54fbe2c70af7c85a97bed0573227e3ccc4b2486012e3ca2a40c6bc77065846f5"
OBJECT_FONT_SHA256 = "3a29d114cb5229e8dbda5bef6c69be4a210a13b7277de4d66e9fc86963226f6c"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def country_runs(path: Path) -> list[dict[str, Any]]:
    with ZipFile(path) as archive:
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cell = sheet.find(f".//{{{S}}}c[@r='C3']")
        require(cell is not None, "family C3 missing")
        assert cell is not None
        node: ET.Element | None
        if cell.get("t") == "s":
            shared = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            node = shared[int(cell.findtext(f"{{{S}}}v", "-1"))]
        else:
            node = cell.find(f"{{{S}}}is")
        require(node is not None, "family C3 rich text missing")
        assert node is not None
        runs = []
        for run in node.findall(f"{{{S}}}r"):
            props = run.find(f"{{{S}}}rPr")
            require(props is not None, "family C3 run properties missing")
            assert props is not None
            values = {
                child.tag.removeprefix(f"{{{S}}}"): child.attrib for child in props
            }
            require(
                len(props) == len(values) == 6
                and set(values) == {"rFont", "family", "charset", "b", "color", "sz"}
                and values["b"] in ({}, {"val": "1"})
                and set(values["color"]) == {"indexed"},
                "unsupported family C3 font semantics",
            )
            runs.append(
                {
                    "text": run.findtext(f"{{{S}}}t", ""),
                    "font": {
                        "name": values["rFont"]["val"],
                        "size": float(values["sz"]["val"]),
                        "bold": True,
                        "family": int(values["family"]["val"]),
                        "charset": int(values["charset"]["val"]),
                        "color_indexed": int(values["color"]["indexed"]),
                    },
                }
            )
        require(len(runs) == 2, "family C3 must retain both canonical runs")
        return runs


def c13_family_evidence(path: Path) -> tuple[dict[str, Any], str | None]:
    """Read only the family C13 semantics that are stable across sources."""

    with ZipFile(path) as archive:
        sheet_xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        styles_xml = ET.fromstring(archive.read("xl/styles.xml"))
        row = sheet_xml.find(f".//{{{S}}}row[@r='13']")
        cell = sheet_xml.find(f".//{{{S}}}c[@r='C13']")
        require(row is not None and cell is not None, "family C13 geometry missing")
        assert row is not None and cell is not None
        xfs = styles_xml.find(f"{{{S}}}cellXfs")
        fonts = styles_xml.find(f"{{{S}}}fonts")
        require(xfs is not None and fonts is not None, "family style tables missing")
        assert xfs is not None and fonts is not None
        style_id = int(cell.get("s", "0"))
        require(0 <= style_id < len(xfs), "family C13 style index invalid")
        xf = xfs[style_id]
        font_id = int(xf.get("fontId", "-1"))
        require(0 <= font_id < len(fonts), "family C13 font index invalid")
        font = fonts[font_id]
        name = font.find(f"{{{S}}}name")
        size = font.find(f"{{{S}}}sz")
        color = font.find(f"{{{S}}}color")
        alignment = xf.find(f"{{{S}}}alignment")
        require(
            name is not None
            and name.get("val") == "Times New Roman"
            and size is not None
            and float(size.get("val", "0")) == 14
            and font.find(f"{{{S}}}b") is not None
            and font.find(f"{{{S}}}i") is not None
            and color is not None
            and color.attrib == {"indexed": "8"},
            "family C13 font semantics differ",
        )
        require(
            alignment is None or not alignment.attrib,
            "family C13 must retain default alignment",
        )
    book = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    try:
        sheet = book["Лист1"]
        require(sheet.row_dimensions[13].height == 26.1, "family row 13 differs")
        require(
            all(sheet[f"{column}13"].value is None for column in "DEFGHI"),
            "family C13 spill cells are occupied",
        )
        require(
            not any(
                rng.min_row <= 13 <= rng.max_row for rng in sheet.merged_cells.ranges
            ),
            "family row 13 must not be merged",
        )
        widths = {column: sheet.column_dimensions[column].width for column in "CDEFGH"}
        value = sheet["C13"].value
        require(value is None or isinstance(value, str), "family C13 value invalid")
    finally:
        book.close()
    return (
        {
            "row_height_pt": 26.1,
            "font": {
                "name": "Times New Roman",
                "size": 14,
                "bold": True,
                "italic": True,
                "color_indexed": 8,
            },
            "stored_alignment": {
                "horizontal": None,
                "vertical": None,
                "wrap_text": False,
                "shrink_to_fit": False,
            },
            "effective_alignment": {"horizontal": "general", "vertical": "bottom"},
            "spill_columns": list("CDEFGH"),
            "required_empty_cells": ["D13", "E13", "F13", "G13", "H13"],
            "column_widths": widths,
        },
        value,
    )


def font_space_advance_em(raw: bytes) -> float:
    """Retain the already-proven C3 spacing binding from Times New Roman Bold."""

    # The exact source was previously proven: timesbd.ttf U+0020 advance=512/2048.
    require(sha256_bytes(raw) == COUNTRY_FONT_SHA256, "country font SHA mismatch")
    return 0.25


def extract_profile_v0_4(
    predecessor: BoundInput,
    measurements: BoundInput,
    country_font_source: BoundInput,
) -> dict[str, Any]:
    require(
        predecessor.expected_sha256 == APPROVED_V03_SHA256,
        "v0.3 predecessor SHA mismatch",
    )
    parent = load_bound(predecessor, "approved v0.3 predecessor")
    profile = load_strict_json(parent.raw, "v0.3 profile")
    contract = profile["presentation_contract"]
    fingerprint = sha256_bytes(canonical_json(contract))
    require(
        profile["schema_version"] == "dinva_classic_presentation_profile.v0.3"
        and profile["profile_id"] == "DINVA_CLASSIC_QUOTE_INVOICE_V0_1_DYNAMIC_V0_3"
        and profile["artifact_status"] == "IMMUTABLE_APPROVED_PROFILE"
        and profile["approval_provenance"]["status"] == "APPROVED"
        and profile["approval_provenance"]["authority"] == "IGOR_DIRECT_HUMAN_APPROVAL"
        and profile["approval_provenance"]["approved_contract_fingerprint"]
        == fingerprint
        and profile["presentation_contract_fingerprint"] == fingerprint,
        "v0.3 predecessor approval mismatch",
    )
    bindings = profile["reference_provenance"]
    require(len(bindings) == 6, "v0.3 source binding count mismatch")
    sources: list[BoundInput] = []
    family: list[BoundInput] = []
    for binding in bindings:
        require(
            binding["expected_sha256"] == binding["actual_sha256"],
            "source binding mismatch",
        )
        source = BoundInput(Path(binding["path"]), binding["actual_sha256"])
        load_bound(source, "v0.3 evidence")
        sources.append(source)
        if binding["role"] == "CLASSIC_FAMILY_EVIDENCE":
            family.append(source)
    require(
        len(family) == 3
        and {s.expected_sha256 for s in family} == set(REQUIRED_FAMILY_SHA256S),
        "v0.4 family evidence mismatch",
    )
    runs = [country_runs(source.path) for source in family]
    require(
        all(value == runs[0] for value in runs), "family C3 rich-text evidence differs"
    )
    require(
        "".join(run["text"] for run in runs[0])
        == contract["fixed_blocks"]["company"]["C3"],
        "family C3 text differs from predecessor",
    )
    c13 = [c13_family_evidence(source.path) for source in family]
    require(
        all(value[0] == c13[0][0] for value in c13),
        "family C13 presentation evidence differs",
    )
    golden_objects = [value for _, value in c13 if value is not None]
    require(len(golden_objects) == 1, "family C13 object evidence is ambiguous")

    measured = load_bound(measurements, "C13 native width measurements")
    data = load_strict_json(measured.raw, "C13 native width measurements")
    font_binding = data["font_source"]
    require(font_binding["sha256"] == OBJECT_FONT_SHA256, "object font SHA mismatch")
    object_font_source = BoundInput(Path(font_binding["path"]), font_binding["sha256"])
    object_font = load_bound(object_font_source, "object font")
    require(
        data["schema_version"] == "dinva_native_horizontal_text_measurements.v0.1"
        and data["font"]
        == {"name": "Times New Roman", "size": 14, "bold": True, "italic": True}
        and data["engine"]["method"]
        == "COLUMN_AUTOFIT_DIFFERENCE_20_60_WITH_0_75_PT_QUANTIZATION_UPPER_BOUND"
        and data["engine"]["quantization_pt"] == 0.75
        and data["normal_font"]
        == {"name": "Calibri", "size": 11, "maximum_digit_width_px": 7}
        and data["horizontal_padding_px"] == 5
        and data["line_break_display"] == "ZERO_ADVANCE_SINGLE_LINE",
        "native horizontal measurement method mismatch",
    )
    advances: dict[str, float] = {}
    for codepoint, glyph in data["glyph_measurements"].items():
        upper = (glyph["width60_pt"] - glyph["width20_pt"] + 0.75) / 40
        require(upper == glyph["upper_advance_pt"], "native glyph evidence mismatch")
        advances[codepoint] = upper
    evidence = c13[0][0]
    rule = {
        "cell": "C13",
        "model": "NATIVE_EXCEL_SINGLE_LINE_SPILL_V1",
        "row_height_pt": evidence["row_height_pt"],
        "font": evidence["font"],
        "stored_alignment": evidence["stored_alignment"],
        "effective_alignment": evidence["effective_alignment"],
        "line_break_display": data["line_break_display"],
        "spill_columns": evidence["spill_columns"],
        "required_empty_cells": evidence["required_empty_cells"],
        "maximum_digit_width_px": data["normal_font"]["maximum_digit_width_px"],
        "horizontal_padding_px": data["horizontal_padding_px"],
        "glyph_upper_advances_pt": advances,
        "measurement_source_sha256": measured.sha256,
        "font_source_sha256": object_font.sha256,
        "source_sha256s": sorted(source.expected_sha256 for source in family),
        "source_locator": "Лист1!C13:H13;xl/worksheets/sheet1.xml;xl/styles.xml",
    }
    fit = canonical_spill_fit(rule, golden_objects[0], evidence["column_widths"])
    calibration = data["calibration"]
    require(
        calibration["corridor_width_pt"] == 804.0
        and calibration["horizontal_padding_pt"] == 3.75
        and calibration["available_width_pt"] == fit["available_pt"] == 800.25
        and abs(calibration["required_upper_bound_pt"] - fit["required_upper_bound_pt"])
        < 1e-9
        and abs(calibration["minimum_margin_pt"] - fit["minimum_margin_pt"]) < 1e-9,
        "canonical Invoice519 spill calibration mismatch",
    )

    country_font = load_bound(country_font_source, "country font")
    require(country_font.sha256 == COUNTRY_FONT_SHA256, "country font SHA mismatch")
    successor = copy.deepcopy(profile)
    successor.update(
        schema_version="dinva_classic_presentation_profile.v0.4",
        profile_id="DINVA_CLASSIC_QUOTE_INVOICE_V0_1_DYNAMIC_V0_4",
        artifact_status="DRAFT_PROFILE_CANDIDATE",
    )
    new = successor["presentation_contract"]
    new["contract_version"] = "dinva_classic_presentation_contract.v0.4"
    new["layout"]["top_row_heights"]["13"] = 26.1
    top = new["layout"]["top_block_rules"]
    top.pop("object_height")
    top["object_presentation"] = rule
    top["country_rich_text"] = {
        "cell": "C3",
        "runs": runs[0],
        "space_advance_em": font_space_advance_em(country_font.raw),
        "font_source_sha256": country_font.sha256,
        "source_sha256s": sorted(source.expected_sha256 for source in family),
        "source_locator": "Лист1!C3",
    }
    object_style = new["styles"]["object"]
    object_style["font"].update(
        italic=True, color={"type": "indexed", "value": 8, "tint": 0.0}
    )
    object_style["alignment"].update(
        horizontal=None, vertical=None, wrap_text=False, shrink_to_fit=False
    )
    successor["presentation_contract_fingerprint"] = sha256_bytes(canonical_json(new))
    successor["approval_provenance"] = {
        "status": "DRAFT_UNAPPROVED",
        "authority": None,
        "approval_id": None,
        "approved_at": None,
        "approved_contract_fingerprint": None,
    }
    for loaded, role in [
        (parent, "IMMUTABLE_APPROVED_V0_3_PREDECESSOR_PROFILE"),
        (measured, "NATIVE_C13_BOLD_ITALIC_WIDTH_MEASUREMENTS"),
        (object_font, "NATIVE_OBJECT_BOLD_ITALIC_FONT_SOURCE"),
        (country_font, "NATIVE_COUNTRY_BOLD_FONT_SOURCE"),
    ]:
        successor["reference_provenance"].append(
            {
                "path": str(loaded.path),
                "expected_sha256": loaded.sha256,
                "actual_sha256": loaded.sha256,
                "role": role,
            }
        )
    for source in [
        predecessor,
        measurements,
        object_font_source,
        country_font_source,
        *sources,
    ]:
        load_bound(source, "v0.4 final source reread")
    return successor


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--approved-v0-3-profile", required=True, type=Path)
    parser.add_argument("--approved-v0-3-profile-sha256", required=True)
    parser.add_argument("--width-measurements", required=True, type=Path)
    parser.add_argument("--width-measurements-sha256", required=True)
    parser.add_argument("--country-font-source", required=True, type=Path)
    parser.add_argument("--country-font-source-sha256", required=True)
    parser.add_argument("--output-profile", required=True, type=Path)
    args = parser.parse_args()
    try:
        result = extract_profile_v0_4(
            BoundInput(args.approved_v0_3_profile, args.approved_v0_3_profile_sha256),
            BoundInput(args.width_measurements, args.width_measurements_sha256),
            BoundInput(args.country_font_source, args.country_font_source_sha256),
        )
        output = publish_draft_profile(result, args.output_profile)
    except (OSError, ValueError, KeyError) as exc:
        print(f"HOLD: {exc}")
        return 1
    print("DINVA_CLASSIC_PROFILE_V0_4=DRAFT_UNAPPROVED")
    print(f"FINGERPRINT={result['presentation_contract_fingerprint']}")
    print(f"SHA256={sha256_bytes(output.read_bytes())}")
    print(f"OUTPUT={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
