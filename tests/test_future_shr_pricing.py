"""Regression for exact ordinary-board apparatus/work and enclosure cost sources."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from typing import Any

import future_case_hardening as f  # type: ignore[import-not-found]
import governed_metal_cost as m  # type: ignore[import-not-found]
import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from price_baseline_contract import (  # type: ignore[import-not-found]
    PriceBaseline,
    sha256_file,
)
from test_future_case_hardening import (
    CONTEXT,
    binding,
    governed,
    save_source,
    source_fixture,
)


def ordinary_source(tmp_path: Path, *, metal: bool = False) -> dict[str, Any]:
    _, source = source_fixture(tmp_path, "ЩР", area=10)
    item = source["items"][0]
    item["cabinet"] = {
        "cabinet_code": "METAL-500X400X250-D22" if metal else "CAB-KRN-18",
        "cabinet_label": "500×400×250 IP31 УХЛ3" if metal else "ЩРн 18з-1 36 УХЛ3 IP31",
    }
    specs = [("P1", "ВА88-32 3P 100A", "mccb_up_to_100a", 1)] if metal else []
    specs += [("M1", "ВА47-29 3P D25", "modular_3p", 6)]
    item["components"] = [
        {
            "component_id": cid,
            "component_code": label.split()[0],
            "component_label": label,
            "install_type": kind,
            "quantity": qty,
        }
        for cid, label, kind, qty in specs
    ]
    t = item["technical_classification"]
    t["roles"] = {cid: "POWER_BREAKER" for cid, *_ in specs}
    t["categories"] = {
        cid: "MCCB 3P 80-100A" if kind == "mccb_up_to_100a" else "MCB 3P D 10-63A"
        for cid, _, kind, _ in specs
    }
    t["circuits"][0]["protection_component_ids"] = [specs[0][0]]
    if metal:
        t["overrides"] = {
            "human": {
                "enclosure_mm": binding(tmp_path, "override.json", [500, 400, 250])
            }
        }
    return source


def metal22(tmp_path: Path) -> Path:
    path = tmp_path / "metal22.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист1"
    sheet["A22"] = "500х400х250 "
    for cell, value in {
        "D22": "=ROUNDUP((M22+N22)*R22+O22+P22+S22+Q22,0)",
        "H22": "=E22/1000",
        "I22": "=F22/1000",
        "J22": "=G22/1000",
        "K22": "=H22*I22*3+H22*J22*2+I22*J22*2",
        "L22": "=K22*$B$1*8.42",
        "M22": "=L22*$B$2",
        "N22": "=K22*2*0.25*$B$3*1.33",
        "Q22": "=1250",
        "E22": 500,
        "F22": 400,
        "G22": 250,
        "B1": 1,
        "B2": 530,
        "B3": 3000,
        "O22": 2500,
        "P22": 400,
        "R22": 1.05,
        "S22": 1500,
        "B22": "=ROUNDUP(D22*$B$5,0)",
        "B5": 1.4375,
    }.items():
        sheet[cell] = value
    book.save(path)
    book.close()
    return path


def prices(tmp_path: Path, monkeypatch: Any) -> Path:
    path = tmp_path / "prices.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "ЩР"
    for row, label, material, work in [
        (9, "ВА55/57/59,  АМ1 от 80 до 100А", 15000, 1800),
        (16, "ВА47 3Р D 10-63A", 4000, 540),
    ]:
        for column, value in enumerate([label, material, work], 1):
            sheet.cell(row, column, value)
    krn = book.create_sheet("КРН")
    krn["L7"] = "Корпус КРН-18 265х440х100 "
    krn["M7"] = 7678
    book.save(path)
    book.close()
    # The production resolver authenticates this snapshot; fixtures inject its result.
    snapshot = (
        {
            "mapping_id": "CABINET-DYNAMIC-CAB-KRN-18",
            "source": {
                "sheet": "КРН",
                "row": 7,
                "label_cell": "КРН!L7",
                "price_cells": ["КРН!M7"],
            },
            "prices": {"price_kzt": 7678},
        },
    )
    baseline = PriceBaseline(
        "active_approved", path, sha256_file(path), mapping_snapshot=snapshot
    )
    monkeypatch.setattr(f, "resolve_active_price_baseline", lambda _: baseline)
    return path


def calculate(
    tmp_path: Path, monkeypatch: Any, *, metal: bool = False
) -> dict[str, Any]:
    source = ordinary_source(tmp_path, metal=metal)
    bound = save_source(tmp_path, source)
    rules, bom = governed(bound)
    prices(tmp_path, monkeypatch)
    if metal:
        path = metal22(tmp_path)
        monkeypatch.setattr(m, "GENESIS_PATH", path)
        monkeypatch.setattr(m, "GENESIS_SHA", sha256_file(path))
    return dict(
        f.calculate_future_price(
            CONTEXT,
            bom=bom,
            rules=rules,
            selector=tmp_path / "selector.json",
            cabinet_cost={
                **binding(tmp_path, "cost.json", 12770 if metal else 7678),
                "role": "ENCLOSURE_COST",
            },
            work_source=binding(tmp_path, "work.json", 5040 if metal else 3240),
        )
    )


@pytest.mark.parametrize("metal,total", [(False, 57095), (True, 92877)])
def test_ordinary_rows_and_per_device_work(
    tmp_path: Path, monkeypatch: Any, metal: bool, total: int
) -> None:
    result = calculate(tmp_path, monkeypatch, metal=metal)
    assert result["status"] == "DRAFT_NOT_APPROVED"
    assert result["unit_price_kzt"] == total
    assert result["work_kzt"] == ("5040" if metal else "3240")
    assert all(row["locator"].startswith("ЩР!") for row in result["rows"])
    assert sum(
        Decimal(r["unit_work_kzt"]) * Decimal(r["quantity"]) for r in result["rows"]
    ) == Decimal(result["work_kzt"])
    assert result["cabinet_source"]["source_locator"] == (
        "Лист1!D22" if metal else "КРН!M7"
    )


@pytest.mark.parametrize(
    "label",
    [
        "ВА47-29 3P D16",
        "ВА47-29 3P D25",
        "ВА47-29 3P D32",
        "ВА47-29 3P D40",
        "ВА47-29 3Р D10",
        "ВА47-29 3P D63",
    ],
)
def test_approved_d_range(tmp_path: Path, label: str) -> None:
    source = ordinary_source(tmp_path)
    source["items"][0]["components"][0]["component_label"] = label
    governed(save_source(tmp_path, source))


@pytest.mark.parametrize(
    "label,kind,category",
    [
        ("ВА47-29 3P C25", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА47-29 3P 25A", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА47-29 3P D9", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА47-29 3P D64", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА47-29 1P D25", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА47-29 3P D25", "modular_1p", "MCB 3P D 10-63A"),
        ("ВА47-29 3P D25", "mccb_up_to_100a", "MCB 3P D 10-63A"),
        ("ВА47-29 3P D25 D32", "modular_3p", "MCB 3P D 10-63A"),
        ("ВА88-32 3P 100A", "mccb_up_to_100a", "MCCB 3P 16-63A"),
        ("ВА88-32 3P 79A", "mccb_up_to_100a", "MCCB 3P 80-100A"),
        ("ВА88-32 3P 101A", "mccb_up_to_100a", "MCCB 3P 80-100A"),
        ("ВА88-32 1P 100A", "mccb_up_to_100a", "MCCB 3P 80-100A"),
        ("ВА88-32 3P D100", "mccb_up_to_100a", "MCCB 3P 80-100A"),
        ("ВА88-32 3P 100A", "mccb_125_250a", "MCCB 3P 80-100A"),
    ],
)
def test_type_poles_range_and_d_fail_closed(
    tmp_path: Path, label: str, kind: str, category: str
) -> None:
    source = ordinary_source(tmp_path)
    comp = source["items"][0]["components"][0]
    comp.update(component_label=label, install_type=kind)
    source["items"][0]["technical_classification"]["categories"]["M1"] = category
    with pytest.raises(f.RuleError):
        governed(save_source(tmp_path, source))


@pytest.mark.parametrize(
    "cell,value",
    [
        ("A9", "wrong MCCB"),
        ("A16", "ВА47 3Р C 10-63A"),
        ("C16", None),
        ("C9", "=1800"),
        ("L7", "Корпус КРН-24"),
        ("M7", "=7678"),
    ],
)
def test_source_cell_drift_hold(
    tmp_path: Path, monkeypatch: Any, cell: str, value: Any
) -> None:
    metal = not cell.startswith(("L", "M"))
    source = ordinary_source(tmp_path, metal=metal)
    bound = save_source(tmp_path, source)
    rules, bom = governed(bound)
    path = prices(tmp_path, monkeypatch)
    book = load_workbook(path)
    book["КРН" if cell.startswith(("L", "M")) else "ЩР"][cell] = value
    book.save(path)
    book.close()
    baseline = f.resolve_active_price_baseline(path)
    monkeypatch.setattr(
        f,
        "resolve_active_price_baseline",
        lambda _: replace(baseline, sha256=sha256_file(path)),
    )
    metal_path = metal22(tmp_path)
    monkeypatch.setattr(m, "GENESIS_PATH", metal_path)
    monkeypatch.setattr(m, "GENESIS_SHA", sha256_file(metal_path))
    with pytest.raises(f.RuleError):
        f.calculate_future_price(
            CONTEXT,
            bom=bom,
            rules=rules,
            selector=path,
            cabinet_cost={
                **binding(tmp_path, "cost.json", 12770 if metal else 7678),
                "role": "ENCLOSURE_COST",
            },
            work_source=binding(tmp_path, "work.json", 5040 if metal else 3240),
        )


@pytest.mark.parametrize(
    "fault",
    ["cost", "work", "unknown_cabinet", "mixed_avr", "price_sha", "override_sha"],
)
def test_bound_cost_work_and_authority_hold(
    tmp_path: Path, monkeypatch: Any, fault: str
) -> None:
    source = ordinary_source(tmp_path)
    item = source["items"][0]
    if fault == "unknown_cabinet":
        item["cabinet"]["cabinet_code"] = "OTHER"
    if fault == "mixed_avr":
        item["components"].append(
            {
                "component_id": "C1",
                "component_code": "CONTROL",
                "component_label": "1P 10A",
                "quantity": 1,
                "install_type": "modular_1p",
            }
        )
        item["technical_classification"]["roles"]["C1"] = "CONTROL_BREAKER"
        item["technical_classification"]["categories"]["C1"] = "MCB 1P 10A"
    rules, bom = governed(save_source(tmp_path, source))
    path = prices(tmp_path, monkeypatch)
    if fault == "price_sha":
        path.write_bytes(path.read_bytes() + b"drift")
    cabinet = {
        **binding(tmp_path, "cost.json", 1 if fault == "cost" else 7678),
        "role": "ENCLOSURE_COST",
    }
    work = binding(tmp_path, "work.json", 1800 if fault == "work" else 3240)
    if fault == "override_sha":
        Path(work["path"]).write_text("{}", encoding="utf-8")
    with pytest.raises(f.RuleError):
        f.calculate_future_price(
            CONTEXT,
            bom=bom,
            rules=rules,
            selector=path,
            cabinet_cost=cabinet,
            work_source=work,
        )


@pytest.mark.parametrize(
    "cell,value",
    [
        ("D22", "=B22"),
        ("A22", "600х400х250 "),
        ("E22", 600),
        ("O22", 2501),
        ("B2", 0),
        ("N22", "=0"),
    ],
)
def test_metal22_identity_formula_hold(tmp_path: Path, cell: str, value: Any) -> None:
    path = metal22(tmp_path)
    book = load_workbook(path)
    book.active[cell] = value
    book.save(path)
    book.close()
    with pytest.raises(f.RuleError):
        m.inspect_cost_d22(path, sha256_file(path))


def test_metal22_recomputed_without_formula_cache(tmp_path: Path) -> None:
    path = metal22(tmp_path)
    raw = path.read_bytes()
    result = m.inspect_cost_d22(path, sha256_file(path))
    assert result["cost_kzt"] == 12770
    assert result["source_role"] == "ENCLOSURE_COST"
    assert result["source_locator"] == "Лист1!D22"
    assert path.read_bytes() == raw
    with pytest.raises(f.RuleError, match="SHA"):
        m.inspect_cost_d22(path, "0" * 64)


def test_protection_cannot_hide_under_other_role(tmp_path: Path) -> None:
    source = ordinary_source(tmp_path, metal=True)
    source["items"][0]["technical_classification"]["roles"]["M1"] = "OTHER"
    with pytest.raises(f.RuleError, match="role conflict"):
        governed(save_source(tmp_path, source))
