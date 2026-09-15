from __future__ import annotations

import copy
import hashlib
import json
import os
import re
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from test_dinva_v0_2_governed_input_bridge import load_file
from test_dinva_v0_4_native_visual import GOLDEN_OBJECT, S, native_successor
from test_render_dinva_classic_quote_invoice import (
    ROOT,
    canonical_file,
    make_case,
    refresh_document_fingerprint,
    render_case,
    reshape_case,
)

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def print_successor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[Any, dict[str, Any], Any]:
    _, parent, _, _, _ = native_successor(tmp_path, monkeypatch)
    replacements = {}
    for binding in parent["reference_provenance"]:
        if binding["role"] != "CLASSIC_FAMILY_EVIDENCE":
            continue
        path = Path(binding["path"])
        book = load_workbook(path, rich_text=True)
        sheet = book.worksheets[0]
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.fitToWidth = None
        for key, value in parent["presentation_contract"]["print"]["margins"].items():
            setattr(sheet.page_margins, key, value)
        book.save(path)
        book.close()
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        replacements[binding["actual_sha256"]] = digest
    # Synthetic evidence identity changes must propagate to every governed locator.
    raw = json.dumps(parent)
    for old, new in replacements.items():
        raw = raw.replace(old, new)
    parent = json.loads(raw)
    for node in [
        parent["presentation_contract"]["assets"][0],
        parent["presentation_contract"]["assets"][0]["drawing_semantics"],
        *parent["presentation_contract"]["layout"]["top_block_rules"].values(),
    ]:
        if isinstance(node, dict):
            for key in ("source_sha256s", "source_reference_sha256s"):
                if key in node:
                    node[key].sort()
    parent["presentation_contract"]["layout"]["top_block_rules"][
        "border_source_sha256s"
    ].sort()
    producer = load_file(
        "print_v05_producer",
        ROOT / "scripts/extract_dinva_classic_presentation_profile_v0_5.py",
    )
    fp = producer.sha256_bytes(producer.canonical_json(parent["presentation_contract"]))
    parent["presentation_contract_fingerprint"] = fp
    parent["artifact_status"] = "IMMUTABLE_APPROVED_PROFILE"
    parent["approval_provenance"].update(
        status="APPROVED",
        authority="IGOR_DIRECT_HUMAN_APPROVAL",
        approved_contract_fingerprint=fp,
    )
    path = tmp_path / "approved-v04.synthetic.json"
    digest = canonical_file(path, parent)
    monkeypatch.setattr(producer, "APPROVED_V04_SHA256", digest)
    monkeypatch.setattr(
        producer, "REQUIRED_FAMILY_SHA256S", frozenset(replacements.values())
    )
    predecessor = producer.BoundInput(path, digest)
    return producer, producer.extract_profile_v0_5(predecessor), predecessor


def lifecycle_schema_accepts(schema: dict[str, Any], profile: dict[str, Any]) -> bool:
    def matches(value: Any, spec: dict[str, Any]) -> bool:
        if "$ref" in spec:
            name = spec["$ref"].removeprefix("#/$defs/")
            return matches(value, schema["$defs"][name])
        if (
            "oneOf" in spec
            and sum(matches(value, branch) for branch in spec["oneOf"]) != 1
        ):
            return False
        if "const" in spec and value != spec["const"]:
            return False
        if "enum" in spec and value not in spec["enum"]:
            return False
        expected_type = spec.get("type")
        if expected_type == "object":
            if not isinstance(value, dict):
                return False
            required = set(spec.get("required", []))
            properties = spec.get("properties", {})
            if not required.issubset(value):
                return False
            if spec.get("additionalProperties") is False and not set(value).issubset(
                properties
            ):
                return False
            return all(
                key not in value or matches(value[key], child)
                for key, child in properties.items()
            )
        if expected_type == "null" and value is not None:
            return False
        if expected_type == "string":
            if not isinstance(value, str):
                return False
            if len(value) < spec.get("minLength", 0):
                return False
            if "pattern" in spec and re.fullmatch(spec["pattern"], value) is None:
                return False
        properties = spec.get("properties", {})
        return all(
            key not in value or matches(value[key], child)
            for key, child in properties.items()
        )

    lifecycle = {
        "type": "object",
        "additionalProperties": False,
        "required": ["artifact_status", "approval_provenance"],
        "properties": {
            key: schema["properties"][key]
            for key in ("artifact_status", "approval_provenance")
        },
        "oneOf": schema["oneOf"],
    }
    subject = {key: profile[key] for key in lifecycle["required"]}
    return matches(subject, lifecycle)


@pytest.fixture
def rendered(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Any, ...]:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    producer, profile, predecessor = print_successor(tmp_path, monkeypatch)
    case = make_case(tmp_path)
    reshape_case(
        case, [15, 6, 17, 6, 14, 6, 16, 2, 6], long_text=False, commercial_line_count=8
    )
    case["document"]["object_name"] = GOLDEN_OBJECT
    refresh_document_fingerprint(case["document"])
    case["document_sha"] = canonical_file(case["document_path"], case["document"])
    case["profile"] = profile
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    output = render_case(case, tmp_path / "synthetic-v05.xlsx")
    with ZipFile(output) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    validator = load_file(
        "print_v05_validator", ROOT / "scripts/validate_dinva_classic_quote_invoice.py"
    )
    return producer, profile, predecessor, case, output, parts, validator


def test_production_successor_synthetic_contract_and_scope(
    rendered: tuple[Any, ...],
) -> None:
    producer, profile, predecessor, case, output, parts, validator = rendered
    assert profile == producer.extract_profile_v0_5(predecessor)
    schema = json.loads(
        (
            ROOT / "schemas/dinva_classic_presentation_profile_v0_5.schema.json"
        ).read_bytes()
    )
    assert profile["schema_version"] == schema["properties"]["schema_version"]["const"]
    assert set(profile) == set(schema["required"])
    old = json.loads(predecessor.path.read_bytes())
    restored = copy.deepcopy(profile["presentation_contract"])
    for key in ("contract_version", "print"):
        restored[key] = old["presentation_contract"][key]
    restored["layout"]["pagination"] = old["presentation_contract"]["layout"][
        "pagination"
    ]
    assert restored == old["presentation_contract"]
    assert profile["reference_provenance"][:10] == old["reference_provenance"]
    assert all(
        v is None for k, v in profile["approval_provenance"].items() if k != "status"
    )
    assert validator.validate_native_print(parts, profile) == pytest.approx(
        1412.0734908136485
    )
    assert (
        validator.validate_country_rich_text(parts, profile)["minimum_gap_pt"] == 10.5
    )
    validator.validate_or_raise(
        output,
        profile,
        case["profile_sha"],
        case["document"],
        case["document_sha"],
        allow_test_profile=True,
    )
    sheet_xml = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    run = sheet_xml.find(f".//{{{S}}}c[@r='C3']/{{{S}}}is/{{{S}}}r/{{{S}}}t")
    assert run is not None and run.text == " " * 21 and run.get(XML_SPACE) == "preserve"
    for text in sheet_xml.findall(f".//{{{S}}}t"):
        if text.text and text.text != text.text.strip():
            assert text.get(XML_SPACE) == "preserve"
    plan = case["renderer"].dynamic_layout_plan(
        profile["presentation_contract"], case["document"]
    )
    assert plan["total_row"] == 113 and plan["page_breaks"] == []
    book = load_workbook(output)
    sheet = book.worksheets[0]
    assert (
        sheet.row_dimensions[13].height == 26.1 and sheet["C13"].value == GOLDEN_OBJECT
    )
    assert sheet["C13"].font.name == "Times New Roman" and sheet["C13"].font.sz == 14
    assert sheet["C13"].font.bold and sheet["C13"].font.italic
    assert not sheet["C13"].alignment.wrap_text and not sheet.merged_cells.ranges
    assert (
        not sheet.print_area and not sheet.print_title_rows and not sheet.row_breaks.brk
    )
    book.close()
    with pytest.raises(ValueError, match="immutable"):
        case["renderer"].validate_profile(profile, allow_test_profile=False)


def test_v05_schema_accepts_only_coherent_approval_lifecycle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, draft, _ = print_successor(tmp_path, monkeypatch)
    schema = json.loads(
        (
            ROOT / "schemas/dinva_classic_presentation_profile_v0_5.schema.json"
        ).read_bytes()
    )
    assert lifecycle_schema_accepts(schema, draft)

    approved = copy.deepcopy(draft)
    approved["artifact_status"] = "IMMUTABLE_APPROVED_PROFILE"
    approved["approval_provenance"] = {
        "status": "APPROVED",
        "authority": "IGOR_DIRECT_HUMAN_APPROVAL",
        "approval_id": "SYNTHETIC-V0-5-CONTENT-BOUND-APPROVAL",
        "approved_at": "2026-09-11T00:00:00Z",
        "approved_contract_fingerprint": draft["presentation_contract_fingerprint"],
    }
    assert lifecycle_schema_accepts(schema, approved)

    invalid = []
    mixed_draft = copy.deepcopy(approved)
    mixed_draft["artifact_status"] = "DRAFT_PROFILE_CANDIDATE"
    invalid.append(mixed_draft)
    mixed_approved = copy.deepcopy(draft)
    mixed_approved["artifact_status"] = "IMMUTABLE_APPROVED_PROFILE"
    invalid.append(mixed_approved)
    null_binding = copy.deepcopy(approved)
    null_binding["approval_provenance"]["approved_contract_fingerprint"] = None
    invalid.append(null_binding)
    populated_draft = copy.deepcopy(draft)
    populated_draft["approval_provenance"]["approval_id"] = "NOT-ALLOWED"
    invalid.append(populated_draft)
    assert all(not lifecycle_schema_accepts(schema, profile) for profile in invalid)


@pytest.mark.parametrize("mutation", ["missing", "default", "spaces", "tab", "edge"])
def test_c3_raw_fail_closed(rendered: tuple[Any, ...], mutation: str) -> None:
    _, profile, _, _, _, parts, validator = rendered
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    nodes = root.findall(f".//{{{S}}}c[@r='C3']/{{{S}}}is/{{{S}}}r/{{{S}}}t")
    if mutation == "missing":
        del nodes[0].attrib[XML_SPACE]
    elif mutation == "default":
        nodes[0].set(XML_SPACE, "default")
    elif mutation == "spaces":
        nodes[0].text = " " * 20
    elif mutation == "tab":
        nodes[0].text = "\t" + " " * 20
    else:
        nodes[1].text = " " + str(nodes[1].text)
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with pytest.raises(ValueError, match="country"):
        validator.validate_country_rich_text(parts, profile)


@pytest.mark.parametrize("mutation", ["break", "area", "titles", "fit", "capacity"])
def test_print_raw_fail_closed(rendered: tuple[Any, ...], mutation: str) -> None:
    _, profile, _, _, _, parts, validator = rendered
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    if mutation == "break":
        ET.SubElement(root, f"{{{S}}}rowBreaks")
    elif mutation in {"area", "titles"}:
        book = ET.fromstring(parts["xl/workbook.xml"])
        names = ET.SubElement(book, f"{{{S}}}definedNames")
        ET.SubElement(
            names,
            f"{{{S}}}definedName",
            {
                "name": (
                    "_xlnm.Print_Area" if mutation == "area" else "_xlnm.Print_Titles"
                )
            },
        ).text = "'Sheet1'!$1:$15"
        parts["xl/workbook.xml"] = ET.tostring(book)
    elif mutation == "fit":
        setup = root.find(f"{{{S}}}pageSetup")
        assert setup is not None
        setup.set("fitToWidth", "0")
    else:
        profile["presentation_contract"]["layout"]["pagination"]["capacity"][
            "nominal_unscaled_height_pt"
        ] = 1100
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with pytest.raises(ValueError):
        validator.validate_native_print(parts, profile)


def test_draft_no_overwrite_and_source_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    producer, profile, predecessor = print_successor(tmp_path, monkeypatch)
    out = tmp_path / "draft" / "dinva-classic-presentation-profile-v0.5-DRAFT.json"
    producer.publish_v0_5_draft(profile, predecessor, out)
    raw = out.read_bytes()
    with pytest.raises(ValueError, match="already exists"):
        producer.publish_v0_5_draft(profile, predecessor, out)
    assert out.read_bytes() == raw
    Path(profile["reference_provenance"][0]["path"]).write_bytes(b"drift")
    other = tmp_path / "drift" / out.name
    with pytest.raises(ValueError, match="SHA"):
        producer.publish_v0_5_draft(profile, predecessor, other)
    assert not other.parent.exists()


@pytest.mark.parametrize("phase", ["before_link", "after_link"])
def test_source_toctou_rolls_back_draft(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, phase: str
) -> None:
    producer, profile, predecessor = print_successor(tmp_path, monkeypatch)

    source = Path(profile["reference_provenance"][0]["path"])
    out = tmp_path / "toctou" / "dinva-classic-presentation-profile-v0.5-DRAFT.json"
    if phase == "after_link":
        original_link = os.link

        def link(src: Any, dst: Any) -> None:
            original_link(src, dst)
            source.write_bytes(b"drift after link")

        monkeypatch.setattr(os, "link", link)
    else:
        original_fsync = os.fsync

        def fsync(fd: int) -> None:
            original_fsync(fd)
            source.write_bytes(b"drift after staging")

        monkeypatch.setattr(os, "fsync", fsync)
    with pytest.raises(ValueError, match="SHA"):
        producer.publish_v0_5_draft(profile, predecessor, out)
    assert not out.parent.exists()


def test_nominal_capacity_derived_not_tuned() -> None:
    contract = load_file(
        "dinva_print_contract_capacity_test",
        ROOT / "scripts/dinva_print_contract.py",
    )

    config = {
        "paper_size": "9",
        "orientation": "portrait",
        "scale": 54,
        "margins": {"top": 0.5, "bottom": 0.5},
    }
    first = contract.nominal_page_capacity(config)
    config["scale"] = 27
    assert contract.nominal_page_capacity(config) == first * 2
    config["margins"] = {"top": 0.5, "bottom": 1.5}
    assert contract.nominal_page_capacity(config) == pytest.approx(
        first * 2 - 72 / 0.27
    )


@pytest.mark.parametrize("sizes", [[1], [40], [130]])
def test_dynamic_cases_no_manual_pagination(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, sizes: list[int]
) -> None:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    _, profile, _ = print_successor(tmp_path, monkeypatch)
    case = make_case(tmp_path)
    reshape_case(case, sizes, long_text=False)
    case["profile"] = profile
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    render_case(case, tmp_path / "synthetic-dynamic-v05.xlsx")
