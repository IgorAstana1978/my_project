from __future__ import annotations

import base64
import copy
import hashlib
import json
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.cell.rich_text import (  # type: ignore[import-untyped]
    CellRichText,
    TextBlock,
)
from openpyxl.cell.text import InlineFont  # type: ignore[import-untyped]
from openpyxl.styles import Border, Color, Side  # type: ignore[import-untyped]
from test_dinva_v0_2_governed_input_bridge import build_profile_sources, load_file
from test_render_dinva_classic_quote_invoice import (
    ROOT,
    canonical_file,
    make_case,
    refresh_document_fingerprint,
    render_case,
    reshape_case,
)

XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SYNTHETIC_DRAWING = f"""<xdr:wsDr xmlns:xdr="{XDR}" xmlns:a="{A}"
xmlns:a14="{A14}" xmlns:r="{R}">
<xdr:twoCellAnchor editAs="absolute">
<xdr:from><xdr:col>1</xdr:col><xdr:colOff>57150</xdr:colOff>
<xdr:row>1</xdr:row><xdr:rowOff>76200</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>2</xdr:col><xdr:colOff>1200150</xdr:colOff>
<xdr:row>2</xdr:row><xdr:rowOff>266700</xdr:rowOff></xdr:to>
<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="5265" name="Synthetic logo"/>
<xdr:cNvPicPr><a:picLocks noChangeAspect="1" noChangeArrowheads="1"/>
</xdr:cNvPicPr></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1">
<a:extLst><a:ext uri="{{28A0092B-C50C-407E-A947-70E740481C1C}}">
<a14:useLocalDpi val="0"/></a:ext></a:extLst></a:blip><a:srcRect/>
<a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr bwMode="auto">
<a:xfrm><a:off x="457200" y="285750"/><a:ext cx="1504950" cy="514350"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln>
<a:extLst><a:ext uri="{{909E8E84-426E-40DD-AFC4-6F175D3DCCD1}}">
<a14:hiddenFill><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
</a14:hiddenFill></a:ext><a:ext uri="{{91240B29-F687-4F45-9708-019B960494DF}}">
<a14:hiddenLine w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>
<a:miter lim="800000"/><a:headEnd/><a:tailEnd/></a14:hiddenLine></a:ext></a:extLst>
</xdr:spPr></xdr:pic><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>"""


def require_element(parent: ElementTree.Element, path: str) -> ElementTree.Element:
    element = parent.find(path)
    assert element is not None, f"missing XML element: {path}"
    return element


def rewrite_drawing(path: Path, transform: Any) -> None:
    with ZipFile(path) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    root = ElementTree.fromstring(parts["xl/drawings/drawing1.xml"])
    transform(root)
    parts["xl/drawings/drawing1.xml"] = ElementTree.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)


def successor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, *, country_rich_text: bool = False
) -> tuple[Any, dict[str, Any], Any]:
    producer, family, old, decision = build_profile_sources(tmp_path, monkeypatch)
    # Synthetic family style evidence, including otherwise empty frame cells.
    hashes = []
    logo = base64.b64decode(
        json.loads(old.path.read_bytes())["presentation_contract"]["assets"][0][
            "data_base64"
        ]
    )
    for index, source in enumerate(family):
        w = load_workbook(source.path)
        if country_rich_text:
            w.active["C3"] = CellRichText(
                [
                    TextBlock(
                        InlineFont(
                            rFont="Times New Roman",
                            sz=20,
                            b=True,
                            family=1,
                            charset=204,
                            color="00000000",
                        ),
                        " " * 21,
                    ),
                    TextBlock(
                        InlineFont(
                            rFont="Times New Roman",
                            sz=16,
                            b=True,
                            family=1,
                            charset=1,
                            color="00000000",
                        ),
                        "Республика Казахстан",
                    ),
                ]
            )
            for run in w.active["C3"].value:
                run.font.color = Color(indexed=8)
        w.active["B2"].border = Border(
            left=Side(style="medium", color="FF123456"), top=Side(style="medium")
        )
        w.active["F3"].border = Border(left=Side(style="thin"))
        w.active["I11"].border = Border(
            right=Side(style="medium"), bottom=Side(style="medium")
        )
        w.save(source.path)
        w.close()
        with ZipFile(source.path, "a") as archive:
            archive.writestr(
                "xl/drawings/drawing1.xml",
                SYNTHETIC_DRAWING.replace('id="5265"', f'id="{5265 + index}"'),
            )
            archive.writestr("xl/media/image1.png", logo)
            archive.writestr(
                "xl/drawings/_rels/drawing1.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f'<Relationship Id="rId1" Type="{R}/image" '
                'Target="../media/image1.png"/></Relationships>',
            )
        hashes.append(hashlib.sha256(source.path.read_bytes()).hexdigest())
    # Regenerate synthetic provenance bindings after changing test evidence.
    # Equal synthetic workbooks may share an old SHA but acquire different ZIP
    # timestamps on save. Bind by source index, never global string replacement.
    old_payload = json.loads(old.path.read_bytes())
    old_payload["reference_provenance"] = [
        {"role": "CLASSIC_FAMILY_EVIDENCE", "actual_sha256": digest}
        for digest in hashes
    ]
    canonical_file(old.path, old_payload)
    decision_payload = json.loads(decision.path.read_bytes())
    decision_payload["source_bindings"] = [
        {"role": "CLASSIC_FAMILY_EVIDENCE", "actual_workbook_sha256": digest}
        for digest in hashes
    ]
    canonical_file(decision.path, decision_payload)
    old = producer.BoundInput(
        old.path, hashlib.sha256(old.path.read_bytes()).hexdigest()
    )
    decision = producer.BoundInput(
        decision.path, hashlib.sha256(decision.path.read_bytes()).hexdigest()
    )
    family = [
        producer.BoundInput(s.path, digest)
        for s, digest in zip(family, hashes, strict=True)
    ]
    monkeypatch.setattr(producer, "REQUIRED_FAMILY_SHA256S", frozenset(hashes))
    monkeypatch.setattr(producer, "PREDECESSOR_PROFILE_SHA256", old.expected_sha256)
    monkeypatch.setattr(
        producer, "CANONICAL_LOGO_DECISION_SHA256", decision.expected_sha256
    )
    approved = producer.extract_profile_v0_2(family, old, decision)
    approved["artifact_status"] = "IMMUTABLE_APPROVED_PROFILE"
    approved["approval_provenance"].update(
        status="APPROVED",
        authority="IGOR_DIRECT_HUMAN_APPROVAL",
        approved_contract_fingerprint=approved["presentation_contract_fingerprint"],
    )
    path = tmp_path / "approved-v02.synthetic.json"
    sha = canonical_file(path, approved)
    monkeypatch.setattr(producer, "APPROVED_V02_SHA256", sha)
    source = producer.BoundInput(path, sha)
    return producer, producer.extract_profile_v0_3(source), source


@pytest.mark.parametrize(
    "sizes", [[1], [10, 10, 10, 10], [15, 6, 17, 6, 14, 6, 16, 2, 6], [130]]
)
def test_visual_successor_production_builder_dynamic_render(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, sizes: list[int]
) -> None:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    producer, profile, source = successor(tmp_path / "evidence", monkeypatch)
    assert profile == producer.extract_profile_v0_3(source)
    assert profile["approval_provenance"]["status"] == "DRAFT_UNAPPROVED"
    original = json.loads(source.path.read_bytes())
    assert (
        profile["presentation_contract"]["assets"][0]["data_base64"]
        == original["presentation_contract"]["assets"][0]["data_base64"]
    )
    assert (
        profile["presentation_contract"]["assets"][0]["placement"]
        == original["presentation_contract"]["assets"][0]["placement"]
    )
    geometry = profile["presentation_contract"]["assets"][0]["drawing_semantics"]
    picture = ElementTree.fromstring(geometry["picture_xml"])
    assert require_element(picture, f"{{{XDR}}}spPr/{{{A}}}xfrm/{{{A}}}ext").attrib == {
        "cx": "1504950",
        "cy": "514350",
    }
    assert require_element(picture, f"{{{XDR}}}spPr/{{{A}}}xfrm/{{{A}}}off").attrib == {
        "x": "457200",
        "y": "285750",
    }
    case = make_case(tmp_path)
    reshape_case(case, sizes, long_text=True, commercial_line_count=8)
    case["document"]["object_name"] = (
        "Объект\n"
        + "Многострочное описание объекта по адресу улица дом " * 3
        + "\nКонец"
    )
    refresh_document_fingerprint(case["document"])
    case["document_sha"] = canonical_file(case["document_path"], case["document"])
    case["profile"] = profile
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    output = render_case(case, tmp_path / "synthetic-visual.xlsx")
    validator = load_file(
        "visual_independent_validator",
        ROOT / "scripts/validate_dinva_classic_quote_invoice.py",
    )
    validator.validate_or_raise(
        output,
        profile,
        case["profile_sha"],
        case["document"],
        case["document_sha"],
        allow_test_profile=True,
    )
    contract = profile["presentation_contract"]
    plan = case["renderer"].dynamic_layout_plan(contract, case["document"])
    assert plan["total_row"] == 16 + len(sizes) + sum(sizes)
    if len(sizes) == 9:
        assert plan["total_row"] == 113
    w = load_workbook(output)
    s = w.worksheets[0]
    assert s["C13"].value == case["document"]["object_name"]
    assert 63.75 < s.row_dimensions[13].height <= 408
    assert s.row_dimensions[13].height == validator.independent_object_height(
        contract, case["document"], plan["widths"]
    )
    assert not s.merged_cells.ranges and s.sheet_view.showGridLines
    assert all(s[f"{col}17"].alignment.wrap_text for col in "CFG")
    for coordinate, xml in contract["layout"]["top_block_rules"][
        "border_xml_by_cell"
    ].items():
        assert s[coordinate].border == Border.from_tree(ElementTree.fromstring(xml))
    w.close()
    with ZipFile(output) as archive:
        assert (
            hashlib.sha256(archive.read("xl/media/image1.png")).hexdigest()
            == contract["assets"][0]["sha256"]
        )
    # Corrupt just the saved frame, then independently reject it.
    w = load_workbook(output)
    w.active["B2"].border = Border()
    tampered = tmp_path / "tampered.synthetic.xlsx"
    w.save(tampered)
    w.close()
    with pytest.raises(validator.ValidationError, match="top border drift"):
        validator.validate_or_raise(
            tampered,
            profile,
            case["profile_sha"],
            case["document"],
            case["document_sha"],
            allow_test_profile=True,
        )


def test_successor_draft_only_no_overwrite_and_source_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    producer, profile, source = successor(tmp_path / "evidence", monkeypatch)
    output = (
        tmp_path / "draft-case" / "dinva-classic-presentation-profile-v0.3-DRAFT.json"
    )
    producer.publish_draft_profile(profile, output)
    renderer = make_case(tmp_path)["renderer"]
    with pytest.raises(renderer.RendererError, match="immutable/approved"):
        renderer.validate_profile(profile, allow_test_profile=False)
    with pytest.raises(producer.ProfileV02ExtractionError, match="already exists"):
        producer.publish_draft_profile(profile, output)
    source.path.write_bytes(source.path.read_bytes() + b" ")
    with pytest.raises(producer.ProfileV02ExtractionError, match="SHA-256 mismatch"):
        producer.extract_profile_v0_3(source)


def test_object_height_bounds_and_logo_marker_drift_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    _, profile, _ = successor(tmp_path / "evidence", monkeypatch)
    case = make_case(tmp_path)
    contract = profile["presentation_contract"]
    renderer = case["renderer"]
    widths = renderer.adaptive_column_widths(contract, case["document"])
    for value in [None, "Короткий объект", "Первая строка\nВторая строка"]:
        case["document"]["object_name"] = value
        assert renderer.object_block_height(contract, case["document"], widths) == 63.75
    case["document"]["object_name"] = "Ш" * 10000
    with pytest.raises(renderer.RendererError, match="object content exceeds"):
        renderer.object_block_height(contract, case["document"], widths)
    case["document"]["object_name"] = "Объект"
    refresh_document_fingerprint(case["document"])
    case["document_sha"] = canonical_file(case["document_path"], case["document"])
    # Frozen picture transform cannot authorize a different marker placement.
    # Compare the original governed profile independently to the tampered file.
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    output = render_case(case, tmp_path / "synthetic-original.xlsx")

    def change_marker(root: ElementTree.Element) -> None:
        require_element(
            root, f"{{{XDR}}}twoCellAnchor/{{{XDR}}}to/{{{XDR}}}rowOff"
        ).text = "999999"

    rewrite_drawing(output, change_marker)
    validator = load_file(
        "logo_marker_validator",
        ROOT / "scripts/validate_dinva_classic_quote_invoice.py",
    )
    with pytest.raises(validator.ValidationError, match="logo anchor drift"):
        validator.validate_or_raise(
            output,
            profile,
            case["profile_sha"],
            case["document"],
            case["document_sha"],
            allow_test_profile=True,
        )
    missing = copy.deepcopy(profile)
    del missing["presentation_contract"]["assets"][0]["drawing_semantics"]
    missing["presentation_contract_fingerprint"] = renderer.sha256_bytes(
        renderer.canonical_json(missing["presentation_contract"])
    )
    case["profile_sha"] = canonical_file(case["profile_path"], missing)
    with pytest.raises(renderer.RendererError, match="logo drawing semantics"):
        render_case(case, tmp_path / "must-not-publish.xlsx")
    assert not (tmp_path / "must-not-publish.xlsx").exists()


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_transform",
        "off",
        "extent",
        "crop",
        "stretch",
        "dpi",
        "locks",
        "hidden_geometry",
        "edit_as",
    ],
)
def test_independent_logo_drawing_mutations_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mutation: str
) -> None:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    _, profile, _ = successor(tmp_path / "evidence", monkeypatch)
    case = make_case(tmp_path)
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    output = render_case(case, tmp_path / "synthetic-logo.xlsx")

    def mutate(root: ElementTree.Element) -> None:
        anchor = root[0]
        pic = require_element(anchor, f"{{{XDR}}}pic")
        sp = require_element(pic, f"{{{XDR}}}spPr")
        if mutation == "missing_transform":
            transform = require_element(sp, f"{{{A}}}xfrm")
            sp.remove(transform)
        elif mutation == "off":
            require_element(sp, f"{{{A}}}xfrm/{{{A}}}off").set("x", "457201")
        elif mutation == "extent":
            require_element(sp, f"{{{A}}}xfrm/{{{A}}}ext").set("cx", "1504949")
        elif mutation == "crop":
            require_element(pic, f"{{{XDR}}}blipFill/{{{A}}}srcRect").set("l", "100")
        elif mutation == "stretch":
            require_element(
                pic, f"{{{XDR}}}blipFill/{{{A}}}stretch/{{{A}}}fillRect"
            ).set("r", "100")
        elif mutation == "dpi":
            require_element(pic, f".//{{{A14}}}useLocalDpi").set("val", "1")
        elif mutation == "locks":
            require_element(pic, f".//{{{A}}}picLocks").set("noChangeAspect", "0")
        elif mutation == "hidden_geometry":
            require_element(pic, f".//{{{A14}}}hiddenLine").set("w", "9526")
        else:
            anchor.set("editAs", "twoCell")

    rewrite_drawing(output, mutate)
    validator = load_file(
        "logo_geometry_validator",
        ROOT / "scripts/validate_dinva_classic_quote_invoice.py",
    )
    with pytest.raises(validator.ValidationError, match="logo DrawingML"):
        validator.validate_or_raise(
            output,
            profile,
            case["profile_sha"],
            case["document"],
            case["document_sha"],
            allow_test_profile=True,
        )


def test_family_logo_geometry_disagreement_is_hold(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    producer, profile, _ = successor(tmp_path / "evidence", monkeypatch)
    family = [
        producer.BoundInput(Path(b["path"]), b["actual_sha256"])
        for b in profile["reference_provenance"]
        if b["role"] == "CLASSIC_FAMILY_EVIDENCE"
    ]

    def drift(root: ElementTree.Element) -> None:
        require_element(root, f".//{{{A}}}xfrm/{{{A}}}ext").set("cy", "514351")

    rewrite_drawing(family[0].path, drift)
    family[0] = producer.BoundInput(
        family[0].path, hashlib.sha256(family[0].path.read_bytes()).hexdigest()
    )
    with pytest.raises(
        producer.ProfileV02ExtractionError, match="DrawingML semantics differ"
    ):
        producer.family_logo_drawing(
            family, profile["presentation_contract"]["assets"][0]
        )
