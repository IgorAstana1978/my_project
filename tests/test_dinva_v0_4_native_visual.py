from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Color,
    Font,
)
from test_dinva_v0_2_governed_input_bridge import load_file
from test_dinva_visual_successor import successor
from test_render_dinva_classic_quote_invoice import (
    ROOT,
    canonical_file,
    make_case,
    refresh_document_fingerprint,
    render_case,
    reshape_case,
)

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
GOLDEN_OBJECT = (
    "«МЖК\nпо адресу: г.Астана, район Есиль, район пересечения улиц "
    "Қазыбек би, Төле би,\nХусейн бен Талал»"
)
CANONICAL_WIDTHS = {
    "C": 42.0,
    "D": 8.140625,
    "E": 12.5703125,
    "F": 56.140625,
    "G": 19.85546875,
    "H": 14.42578125,
}


def native_successor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[Any, dict[str, Any], Any, Any, Any]:
    monkeypatch.syspath_prepend(str(ROOT / "scripts"))
    _, parent, _ = successor(tmp_path / "evidence", monkeypatch, country_rich_text=True)
    family = [
        binding
        for binding in parent["reference_provenance"]
        if binding["role"] == "CLASSIC_FAMILY_EVIDENCE"
    ]
    hashes = []
    for index, binding in enumerate(family):
        path = Path(binding["path"])
        book = load_workbook(path, rich_text=True)
        sheet = book["Лист1"]
        sheet.row_dimensions[13].height = 26.1
        for column, width in CANONICAL_WIDTHS.items():
            sheet.column_dimensions[column].width = width
        sheet["C13"] = GOLDEN_OBJECT if index == 1 else None
        sheet["C13"].font = Font(
            name="Times New Roman",
            size=14,
            bold=True,
            italic=True,
            color=Color(indexed=8),
        )
        sheet["C13"].alignment = Alignment()
        for column in "DEFGHI":
            sheet[f"{column}13"] = None
        book.properties.identifier = f"synthetic-classic-family-{index}"
        book.save(path)
        book.close()
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        binding["expected_sha256"] = digest
        binding["actual_sha256"] = digest
        hashes.append(digest)
    assert len(set(hashes)) == 3
    contract = parent["presentation_contract"]
    asset = contract["assets"][0]
    asset["source_reference_sha256s"] = sorted(hashes)
    asset["drawing_semantics"]["source_sha256s"] = sorted(hashes)
    top = contract["layout"]["top_block_rules"]
    top["border_source_sha256s"] = sorted(hashes)
    fingerprint = hashlib.sha256(
        json.dumps(
            contract, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()
    parent["presentation_contract_fingerprint"] = fingerprint
    parent["artifact_status"] = "IMMUTABLE_APPROVED_PROFILE"
    parent["approval_provenance"].update(
        status="APPROVED",
        authority="IGOR_DIRECT_HUMAN_APPROVAL",
        approved_contract_fingerprint=fingerprint,
    )
    producer = load_file(
        "native_visual_v04_producer",
        ROOT / "scripts/extract_dinva_classic_presentation_profile_v0_4.py",
    )
    predecessor_path = tmp_path / "approved-v03.synthetic.json"
    predecessor_sha = canonical_file(predecessor_path, parent)
    monkeypatch.setattr(producer, "APPROVED_V03_SHA256", predecessor_sha)
    monkeypatch.setattr(producer, "REQUIRED_FAMILY_SHA256S", frozenset(hashes))

    object_font = tmp_path / "timesbi.synthetic.ttf"
    object_font.write_bytes(b"synthetic Times New Roman Bold Italic")
    object_font_sha = hashlib.sha256(object_font.read_bytes()).hexdigest()
    monkeypatch.setattr(producer, "OBJECT_FONT_SHA256", object_font_sha)
    measurements = json.loads(
        (ROOT / "tests/fixtures/dinva_c13_bold_italic_native_widths.json").read_bytes()
    )
    measurements["font_source"] = {
        "path": str(object_font),
        "sha256": object_font_sha,
    }
    measurement_path = tmp_path / "native-widths.synthetic.json"
    measurement_sha = canonical_file(measurement_path, measurements)

    country_font = tmp_path / "timesbd.synthetic.ttf"
    country_font.write_bytes(b"synthetic Times New Roman Bold")
    country_font_sha = hashlib.sha256(country_font.read_bytes()).hexdigest()
    monkeypatch.setattr(producer, "COUNTRY_FONT_SHA256", country_font_sha)
    predecessor = producer.BoundInput(predecessor_path, predecessor_sha)
    measured = producer.BoundInput(measurement_path, measurement_sha)
    country = producer.BoundInput(country_font, country_font_sha)
    profile = producer.extract_profile_v0_4(predecessor, measured, country)
    return producer, profile, predecessor, measured, country


@pytest.mark.parametrize("sizes", [[1], [40], [15, 6, 17, 6, 14, 6, 16, 2, 6], [130]])
def test_production_v04_builder_renderer_validator(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, sizes: list[int]
) -> None:
    monkeypatch.setenv("DINVA_RENDERER_TEST_MODE", "1")
    producer, profile, predecessor, measured, country = native_successor(
        tmp_path, monkeypatch
    )
    assert profile == producer.extract_profile_v0_4(predecessor, measured, country)
    old = json.loads(predecessor.path.read_bytes())
    contract = profile["presentation_contract"]
    assert profile["reference_provenance"][:6] == old["reference_provenance"]
    assert profile["approval_provenance"]["status"] == "DRAFT_UNAPPROVED"
    assert contract["layout"]["top_row_heights"]["13"] == 26.1
    assert "object_height" not in contract["layout"]["top_block_rules"]
    assert contract["styles"]["object"]["alignment"] == {
        "horizontal": None,
        "vertical": None,
        "wrap_text": False,
        "shrink_to_fit": False,
    }
    assert contract["styles"]["object"]["font"]["color"] == {
        "type": "indexed",
        "value": 8,
        "tint": 0.0,
    }

    case = make_case(tmp_path)
    reshape_case(case, sizes, long_text=False, commercial_line_count=8)
    case["document"]["object_name"] = GOLDEN_OBJECT
    refresh_document_fingerprint(case["document"])
    case["document_sha"] = canonical_file(case["document_path"], case["document"])
    case["profile"] = profile
    case["profile_sha"] = canonical_file(case["profile_path"], profile)
    renderer = case["renderer"]
    validator = load_file(
        "native_visual_validator",
        ROOT / "scripts/validate_dinva_classic_quote_invoice.py",
    )
    widths = renderer.adaptive_column_widths(contract, case["document"])
    assert renderer.object_block_height(contract, case["document"], widths) == 26.1
    assert (
        validator.independent_object_height(contract, case["document"], widths) == 26.1
    )
    with pytest.raises(renderer.RendererError, match="immutable/approved"):
        renderer.validate_profile(profile, allow_test_profile=False)
    output = render_case(case, tmp_path / "synthetic-native-visual.xlsx")
    validator.validate_or_raise(
        output,
        profile,
        case["profile_sha"],
        case["document"],
        case["document_sha"],
        allow_test_profile=True,
    )
    plan = renderer.dynamic_layout_plan(contract, case["document"])
    assert plan["total_row"] == 16 + len(sizes) + sum(sizes)
    if len(sizes) == 9:
        assert plan["total_row"] == 113
    book = load_workbook(output)
    sheet = book.worksheets[0]
    assert sheet["C13"].value == GOLDEN_OBJECT
    assert sheet.row_dimensions[13].height == 26.1
    assert sheet["C13"].font.name == "Times New Roman"
    assert sheet["C13"].font.sz == 14
    assert sheet["C13"].font.bold and sheet["C13"].font.italic
    assert sheet["C13"].font.color is not None
    assert sheet["C13"].font.color.indexed == 8
    assert sheet["C13"].alignment.horizontal is None
    assert sheet["C13"].alignment.vertical is None
    assert sheet["C13"].alignment.wrap_text is None
    assert all(sheet[f"{column}13"].value is None for column in "DEFGH")
    assert not sheet.merged_cells.ranges and sheet.sheet_view.showGridLines
    assert all(sheet[f"{column}17"].alignment.wrap_text for column in "CFG")
    for address, border in contract["layout"]["top_block_rules"][
        "border_xml_by_cell"
    ].items():
        assert sheet[address].border == Border.from_tree(ET.fromstring(border))
    book.close()
    with ZipFile(output) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    bounds = validator.validate_country_rich_text(parts, profile)
    assert bounds["minimum_gap_pt"] == 10.5
    validator.validate_object_presentation(parts, profile)
    assert (
        hashlib.sha256(parts["xl/media/image1.png"]).hexdigest()
        == contract["assets"][0]["sha256"]
    )


def test_c13_safe_fit_overflow_unknown_and_null(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, profile, _, _, _ = native_successor(tmp_path, monkeypatch)
    renderer = load_file(
        "native_spill_renderer", ROOT / "scripts/render_dinva_classic_quote_invoice.py"
    )
    validator = load_file(
        "native_spill_validator",
        ROOT / "scripts/validate_dinva_classic_quote_invoice.py",
    )
    contract = profile["presentation_contract"]
    rule = contract["layout"]["top_block_rules"]["object_presentation"]
    assert rule["model"] == "NATIVE_EXCEL_SINGLE_LINE_SPILL_V1"
    assert rule["line_break_display"] == "ZERO_ADVANCE_SINGLE_LINE"
    assert rule["spill_columns"] == list("CDEFGH")
    for calculate in (
        renderer.object_block_height,
        validator.independent_object_height,
    ):
        assert (
            calculate(contract, {"object_name": GOLDEN_OBJECT}, CANONICAL_WIDTHS)
            == 26.1
        )
        assert calculate(contract, {"object_name": None}, CANONICAL_WIDTHS) == 26.1
        with pytest.raises(ValueError, match="spill corridor"):
            calculate(contract, {"object_name": "М" * 500}, CANONICAL_WIDTHS)
        with pytest.raises(ValueError, match="unmeasured"):
            calculate(contract, {"object_name": "🙂"}, CANONICAL_WIDTHS)


def test_family_c13_semantics_fail_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    producer, profile, _, _, _ = native_successor(tmp_path, monkeypatch)
    family_path = Path(
        next(
            binding["path"]
            for binding in profile["reference_provenance"]
            if binding["role"] == "CLASSIC_FAMILY_EVIDENCE"
        )
    )
    book = load_workbook(family_path)
    book["Лист1"].row_dimensions[13].height = 27
    book.save(family_path)
    book.close()
    with pytest.raises(producer.ProfileV02ExtractionError, match="row 13 differs"):
        producer.c13_family_evidence(family_path)


@pytest.mark.parametrize("mutation", ["measurement_sha", "font_drift", "family_drift"])
def test_v04_evidence_fail_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mutation: str
) -> None:
    producer, profile, predecessor, measured, country = native_successor(
        tmp_path, monkeypatch
    )
    if mutation == "measurement_sha":
        measured = producer.BoundInput(measured.path, "0" * 64)
    elif mutation == "font_drift":
        Path(json.loads(measured.path.read_bytes())["font_source"]["path"]).write_bytes(
            b"drift"
        )
    else:
        Path(profile["reference_provenance"][0]["path"]).write_bytes(b"drift")
    with pytest.raises(producer.ProfileV02ExtractionError):
        producer.extract_profile_v0_4(predecessor, measured, country)


def test_country_gap_and_draft_no_overwrite_schema(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    producer, profile, _, _, _ = native_successor(tmp_path, monkeypatch)
    top = profile["presentation_contract"]["layout"]["top_block_rules"]
    assert top["country_rich_text"]["runs"][0]["text"] == " " * 21
    schema = json.loads(
        (
            ROOT / "schemas/dinva_classic_presentation_profile_v0_4.schema.json"
        ).read_bytes()
    )
    assert profile["schema_version"] == schema["properties"]["schema_version"]["const"]
    assert profile["profile_id"] == schema["properties"]["profile_id"]["const"]
    assert set(profile) == set(schema["required"])
    output = (
        tmp_path / "new-draft" / "dinva-classic-presentation-profile-v0.4-DRAFT.json"
    )
    producer.publish_draft_profile(profile, output)
    raw = output.read_bytes()
    with pytest.raises(producer.ProfileV02ExtractionError, match="already exists"):
        producer.publish_draft_profile(profile, output)
    assert output.read_bytes() == raw


def test_only_scoped_v03_contract_fields_change(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, profile, predecessor, _, _ = native_successor(tmp_path, monkeypatch)
    old = json.loads(predecessor.path.read_bytes())["presentation_contract"]
    restored = copy.deepcopy(profile["presentation_contract"])
    restored["contract_version"] = old["contract_version"]
    restored["layout"]["top_row_heights"]["13"] = old["layout"]["top_row_heights"]["13"]
    del restored["layout"]["top_block_rules"]["object_presentation"]
    restored["layout"]["top_block_rules"]["object_height"] = old["layout"][
        "top_block_rules"
    ]["object_height"]
    del restored["layout"]["top_block_rules"]["country_rich_text"]
    restored["styles"]["object"] = old["styles"]["object"]
    assert restored == old
