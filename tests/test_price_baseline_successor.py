"""Exact successor binding and read-only price-ingestion contract tests."""

from __future__ import annotations

import importlib.util
import sys
from dataclasses import replace
from pathlib import Path
from types import ModuleType
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from scripts import price_baseline_contract as binding


def load_script(name: str, filename: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        name, Path(__file__).resolve().parents[1] / "scripts" / filename
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


calculator = cast(
    Any, load_script("successor_calculator_test", "calc_quote_price_draft.py")
)
runner = cast(
    Any,
    load_script(
        "successor_runner_test", "run_checked_price_calculator_from_completed_draft.py"
    ),
)

A_NAMES = {
    "ВА47 1 полюсный",
    "ВА47 2 полюсный",
    "ВА47 3 полюсный до 63А",
}
B_NAMES = {
    "ВА55/57/59, АМ1  3 полюсные от 16 до 63А",
    "УЗО АД-32 1Р+N до 63А EKF",
}
C_NAMES = {
    "ВА47 1 полюсный 10А",
    "ВА47 3 полюсный 10А",
    "ВА55/57/59 400А",
    "ВА55/57/59,  АМ1 от 80 до 100А",
    "Контактор до 250А",
    "Контактор до 400А",
    "Контактор до 630А",
    "ПН-2 100А",
    "ПН-2 250А",
    "ПН-2 400А",
    "Реле времени суточное ТЭ-15",
}
REAL_WORKBOOKS_AVAILABLE = (
    binding.HISTORICAL.path.is_file() and binding.SUCCESSOR.path.is_file()
)


def price_result(version: str, path: Path) -> Any:
    return calculator.PriceCalculationResult(
        price_workbook=path,
        input_csv=Path("unused.csv"),
        price_baseline_version=version,
    )


def price_names(path: Path) -> dict[str, set[int]]:
    """Bounded A/B scan; never traverse a formatted used-range."""
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    found: dict[str, set[int]] = {}
    try:
        for sheet in workbook:
            for label, value in sheet.iter_rows(
                min_row=1, max_row=200, min_col=1, max_col=2, values_only=True
            ):
                if isinstance(label, str) and isinstance(value, int) and value > 0:
                    found.setdefault(label, set()).add(value)
    finally:
        workbook.close()
    return found


def write_price_workbook(path: Path, *, successor: bool) -> None:
    workbook = Workbook()
    krn = workbook.active
    krn.title = "КРН"
    dynamic_rows = [
        (2, "ВА47 1 полюсный", 700, 800, 216),
        (3, "ВА47 2 полюсный", 1400, 1500, 432),
        (4, "ВА47 3 полюсный до 63А", 2300, 2400, 540),
    ]
    for row, label, historical_price, successor_price, work_price in dynamic_rows:
        krn.cell(row=row, column=1, value=label)
        krn.cell(
            row=row,
            column=2,
            value=successor_price if successor else historical_price,
        )
        krn.cell(row=row, column=3, value=work_price)
    krn["A5"] = "УЗО АД-32 1Р+N до 63А EKF"
    krn["B5"] = 4500 if successor else 4100
    krn["C5"] = 432

    shr = workbook.create_sheet("ЩР")
    shr["A8"] = "ВА55/57/59, АМ1  3 полюсные от 16 до 63А"
    shr["B8"] = 15000 if successor else 13000
    shr["C8"] = 1800

    other = workbook.create_sheet("OTHER")
    for row, label in enumerate(sorted(C_NAMES), start=1):
        other.cell(row=row, column=1, value=label)
        other.cell(row=row, column=2, value=(2000 if successor else 1000) + row)
    other.cell(row=20, column=1, value="UNCHANGED")
    other.cell(row=20, column=2, value=999)
    workbook.save(path)
    workbook.close()


@pytest.fixture
def synthetic_baselines(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[binding.PriceBaseline, binding.PriceBaseline]:
    historical_path = tmp_path / "historical.xlsx"
    successor_path = tmp_path / "successor.xlsx"
    write_price_workbook(historical_path, successor=False)
    write_price_workbook(successor_path, successor=True)
    historical = replace(
        binding.HISTORICAL,
        path=historical_path,
        sha256=binding.sha256_file(historical_path),
    )
    successor = replace(
        binding.SUCCESSOR,
        path=successor_path,
        sha256=binding.sha256_file(successor_path),
    )
    monkeypatch.setitem(binding.BASELINES, historical.version, historical)
    monkeypatch.setitem(binding.BASELINES, successor.version, successor)
    return historical, successor


def test_exact_versions_paths_and_shas_fail_closed(
    synthetic_baselines: tuple[binding.PriceBaseline, binding.PriceBaseline],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    historical, successor = synthetic_baselines
    assert (
        binding.require_price_baseline(historical.path, historical.version)
        == historical
    )
    assert (
        binding.require_price_baseline(successor.path, successor.version) == successor
    )
    for path, version in (
        (historical.path, successor.version),
        (successor.path, historical.version),
        (successor.path, "unknown"),
        (Path("other.xlsx"), binding.SUCCESSOR.version),
    ):
        with pytest.raises(ValueError):
            binding.require_price_baseline(path, version)
    monkeypatch.setitem(
        binding.BASELINES,
        successor.version,
        replace(successor, sha256="0" * 64),
    )
    with pytest.raises(ValueError, match="SHA-256"):
        binding.require_price_baseline(successor.path, successor.version)


def test_all_16_changed_names_are_classified_from_controlled_workbooks(
    synthetic_baselines: tuple[binding.PriceBaseline, binding.PriceBaseline],
) -> None:
    historical, successor = synthetic_baselines
    before = price_names(historical.path)
    after = price_names(successor.path)
    assert before.keys() == after.keys()
    changed = {name for name in before if before[name] != after[name]}
    assert len(changed) == 16
    exact = {
        calculator.normalize_workbook_label(mapping.expected_label)
        for mapping in calculator.APPROVED_COMPONENT_PRICE_MAPPINGS
    }
    dynamic = {
        calculator.normalize_workbook_label(definition.workbook_label)
        for definition in calculator.COMPONENT_DEFINITIONS.values()
        if definition.workbook_label is not None
    }
    b_names = {
        name for name in changed if calculator.normalize_workbook_label(name) in exact
    }
    a_names = {
        name
        for name in changed
        if calculator.normalize_workbook_label(name) in dynamic and name not in b_names
    }
    c_names = changed - a_names - b_names
    assert a_names == A_NAMES
    assert b_names == B_NAMES
    assert c_names == C_NAMES
    assert (len(a_names), len(b_names), len(c_names)) == (3, 2, 11)


def test_successor_exact_mappings_and_historical_values_are_disjoint() -> None:
    old = calculator.APPROVED_COMPONENT_PRICE_MAPPINGS
    new = calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS
    assert len(old) == len(new)
    affected = [
        (historical, successor)
        for historical, successor in zip(old, new, strict=True)
        if historical != successor
    ]
    assert len(affected) == 5
    for historical, successor in affected:
        assert (
            calculator.resolve_component_mapping(
                historical.signature, historical.component_code, old
            )
            == historical
        )
        assert (
            calculator.resolve_component_mapping(
                successor.signature, successor.component_code, new
            )
            == successor
        )
    assert [
        (m.sheet_name, m.row, m.expected_material_price, m.expected_work_price)
        for m, _ in affected
    ] == [
        ("ЩР", 8, 13000, 1800),
        ("КРН", 5, 4100, 432),
        ("КРН", 5, 4100, 432),
        ("КРН", 5, 4100, 432),
        ("КРН", 5, 4100, 432),
    ]
    assert [
        (m.sheet_name, m.row, m.expected_material_price, m.expected_work_price)
        for _, m in affected
    ] == [
        ("ЩР", 8, 15000, 1800),
        ("КРН", 5, 4500, 432),
        ("КРН", 5, 4500, 432),
        ("КРН", 5, 4500, 432),
        ("КРН", 5, 4500, 432),
    ]
    assert runner.EXPECTED_PROFILE_INPUTS[1][1:] == (
        str(binding.HISTORICAL.path),
        binding.HISTORICAL.sha256,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_successor_dynamic_and_exact_prices_are_read_from_controlled_workbook(
    synthetic_baselines: tuple[binding.PriceBaseline, binding.PriceBaseline],
) -> None:
    _, successor = synthetic_baselines
    binding.require_price_baseline(successor.path, successor.version)
    workbook = load_workbook(
        successor.path, read_only=True, data_only=False, keep_links=False
    )
    result = price_result(successor.version, successor.path)
    try:
        dynamic = calculator.read_component_prices(
            workbook["КРН"],
            {"EKF-VA47-29-1P", "EKF-VA47-29-2P", "EKF-VA47-29-3P"},
            result,
        )
        assert dynamic == {
            "EKF-VA47-29-1P": (800, 216),
            "EKF-VA47-29-2P": (1500, 432),
            "EKF-VA47-29-3P": (2400, 540),
        }
        for mapping in calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS:
            if (mapping.sheet_name, mapping.row) in {("ЩР", 8), ("КРН", 5)}:
                assert calculator.read_approved_component_price(
                    workbook, mapping, result
                ) == (mapping.expected_material_price, mapping.expected_work_price)
        assert result.red_flags == []
        old_mapping = calculator.APPROVED_COMPONENT_PRICE_MAPPINGS[0]
        assert (
            calculator.read_approved_component_price(workbook, old_mapping, result)
            is None
        )
        assert "cross-version" in result.red_flags[-1]
    finally:
        workbook.close()


@pytest.mark.skipif(
    not REAL_WORKBOOKS_AVAILABLE,
    reason="authoritative price workbooks are unavailable on this host",
)
def test_local_authoritative_workbooks_match_approved_contract_and_delta() -> None:
    binding.require_price_baseline(binding.HISTORICAL.path, binding.HISTORICAL.version)
    binding.require_price_baseline(binding.SUCCESSOR.path, binding.SUCCESSOR.version)
    before = price_names(binding.HISTORICAL.path)
    after = price_names(binding.SUCCESSOR.path)
    assert before.keys() == after.keys()
    assert {name for name in before if before[name] != after[name]} == (
        A_NAMES | B_NAMES | C_NAMES
    )
    workbook = load_workbook(
        binding.SUCCESSOR.path, read_only=True, data_only=False, keep_links=False
    )
    result = price_result(binding.SUCCESSOR.version, binding.SUCCESSOR.path)
    try:
        assert calculator.read_component_prices(
            workbook["КРН"],
            {"EKF-VA47-29-1P", "EKF-VA47-29-2P", "EKF-VA47-29-3P"},
            result,
        ) == {
            "EKF-VA47-29-1P": (800, 216),
            "EKF-VA47-29-2P": (1500, 432),
            "EKF-VA47-29-3P": (2400, 540),
        }
        for mapping in calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS:
            if (mapping.sheet_name, mapping.row) in {("ЩР", 8), ("КРН", 5)}:
                assert calculator.read_approved_component_price(
                    workbook, mapping, result
                ) == (mapping.expected_material_price, mapping.expected_work_price)
        assert result.red_flags == []
    finally:
        workbook.close()


def test_formula_and_duplicate_price_cases_fail_closed(
    tmp_path: Path, monkeypatch: Any
) -> None:
    book_path = tmp_path / "synthetic.xlsx"
    book = Workbook()
    shr = book.active
    shr.title = "ЩР"
    shr["A8"] = "ВА55/57/59, АМ1  3 полюсные от 16 до 63А"
    shr["B8"] = "=15000"
    shr["C8"] = 1800
    krn = book.create_sheet("КРН")
    krn["A2"] = "ВА47 1 полюсный"
    krn["B2"] = 800
    krn["C2"] = 216
    krn["A3"] = "ВА47 1 полюсный"
    krn["B3"] = 800
    krn["C3"] = 216
    book.save(book_path)
    book.close()
    monkeypatch.setattr(
        calculator, "require_price_baseline", lambda path, version: None
    )
    loaded = load_workbook(book_path, read_only=True, data_only=False)
    result = price_result(binding.SUCCESSOR.version, book_path)
    try:
        assert (
            calculator.read_approved_component_price(
                loaded, calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS[0], result
            )
            is None
        )
        assert any("price mismatch" in flag for flag in result.red_flags)
        assert calculator.read_component_prices(
            loaded["КРН"], {"EKF-VA47-29-1P"}, result
        ) == {"EKF-VA47-29-1P": (800, 216)}
        assert any("duplicate component price row" in flag for flag in result.red_flags)
    finally:
        loaded.close()


def test_frozen_invoice519_profile_rejects_successor_version() -> None:
    result = runner.run_checked_price_calculator_from_completed_draft(
        Path("unused.json"),
        binding.SUCCESSOR.path,
        pricing_profile_path=Path("frozen-profile.json"),
        expected_pricing_profile_sha256="a" * 64,
        price_baseline_version=binding.SUCCESSOR.version,
    )
    assert result.status == "FAIL"
    assert result.red_flags == [
        "frozen Invoice519 profile rejects successor price baseline"
    ]


def test_unversioned_and_unknown_entrypoint_calls_fail_before_input_reads() -> None:
    calc = calculator.calculate_price_draft(binding.SUCCESSOR.path, Path("missing.csv"))
    assert calc.status == "FAIL"
    assert calc.red_flags == ["active selector does not exist"]
    calc_unknown = calculator.calculate_price_draft(
        binding.SUCCESSOR.path,
        Path("missing.csv"),
        price_baseline_version="unknown",
    )
    assert calc_unknown.status == "FAIL"
    assert calc_unknown.red_flags == ["unknown price baseline version"]
    checked = runner.run_checked_price_calculator_from_completed_draft(
        Path("missing.json"), binding.SUCCESSOR.path
    )
    assert checked.status == "FAIL"
    assert checked.red_flags == ["active selector does not exist"]
