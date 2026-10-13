"""Governed future source mappings: real semantics, missing authority and drift."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

import dinva_future_content as content  # type: ignore[import-not-found]
import future_case_hardening as f  # type: ignore[import-not-found]
import governed_metal_cost as metal  # type: ignore[import-not-found]
import pytest
import wu8_governed_sources as w  # type: ignore[import-not-found]
from openpyxl import Workbook  # type: ignore[import-untyped]
from price_baseline_contract import (  # type: ignore[import-not-found]
    PriceBaseline,
    sha256_file,
)
from test_future_case_hardening import (
    CONTEXT,
    binding,
    governed,
    save_source,
    source_fixture,
)
from test_future_shr_pricing import metal22, ordinary_source, prices


def avr_sheet(topology: str) -> Any:
    book = Workbook()
    ws = book.active
    name, count, lc, vc, br, dr, cr = w.AVR_SHEETS[topology]
    ws.title = name
    ws[f"{lc}{br-1}"] = "Автомат "
    ws[f"{lc}{dr-1}"] = "Привод"
    for offset, (rating, breaker, drive) in enumerate(
        [
            (100, 15600, 30000),
            (160, 19200, 30000),
            (250, 37200, 48000),
            (400, 74400, 78000),
            (630, 120000, 108000),
        ]
    ):
        ws[f"{lc}{br+offset}"] = f"{rating}А"
        ws[f"{vc}{br+offset}"] = breaker
        ws[f"{lc}{dr+offset}"] = f"{rating}А"
        ws[f"{vc}{dr+offset}"] = drive
    ceilings = (
        [95, 100, 160, 250, 400, 630, 800]
        if count == 3
        else [63, 100, 160, 250, 400, 630]
    )
    work_values = (
        [18000, 22500, 24750, 27000, 32400, 36750, 41440]
        if count == 3
        else [9000, 12600, 16200, 18000, 21600, 25200]
    )
    for row, (rating, work) in enumerate(zip(ceilings, work_values, strict=True), 2):
        work_lc, work_vc = ("J", "K") if count == 3 else ("K", "L")
        ws[f"{work_lc}{row}"] = (
            f"АВР 3 ввода до {rating}А с блоком "
            if count == 3
            else f"АВР до {rating}А "
        )
        ws[f"{work_vc}{row}"] = work
    for offset, label in enumerate(w.CONTACTOR_LABELS):
        if w.CONTACTOR_RATINGS[offset] not in w.CONTACTOR_SOURCE_RATINGS[topology]:
            continue
        ws.cell(cr + offset, 1, label)
        ws.cell(
            cr + offset,
            2,
            282000 if offset == 14 else 310000 if offset == 15 else 19000,
        )
    return ws


@pytest.mark.parametrize(
    "topology,count,material,work",
    [
        ("Г-Г", 2, "91200", "12600"),
        ("Г-Д", 2, "91200", "12600"),
        ("Г-Г-Д", 3, "136800", "22500"),
    ],
)
def test_motor_units_and_separate_work(
    topology: str, count: int, material: str, work: str
) -> None:
    result = w.avr_price(avr_sheet(topology), topology, "MCCB", 100, Decimal(count))
    assert result["material_kzt"] == material
    assert result["work"]["value"] == work
    assert [(p["role"], p["quantity"], p["value"]) for p in result["parts"]] == [
        ("AUTOMATIC", str(count), "15600"),
        ("DRIVE", str(count), "30000"),
    ]
    assert result["work"]["quantity"] == "1"
    assert result["decision"]["authority"] == "IGOR_DIRECT_HUMAN_DECISION"
    assert not result["decision"]["case_approval"]


@pytest.mark.parametrize("topology", list(w.AVR_SHEETS))
@pytest.mark.parametrize("count", [1, 4, Decimal("2.5")])
def test_motor_quantity_cannot_be_inferred(topology: str, count: Any) -> None:
    with pytest.raises(f.RuleError, match="input count"):
        w.avr_price(avr_sheet(topology), topology, "MCCB", 100, Decimal(count))


@pytest.mark.parametrize("topology", list(w.AVR_SHEETS))
@pytest.mark.parametrize("evidence", [None, {}, {"available": True}, True, [], 630])
def test_630_availability_requires_confirmation(topology: str, evidence: Any) -> None:
    with pytest.raises(f.RuleError, match="630A"):
        w.avr_price(
            avr_sheet(topology),
            topology,
            "MCCB",
            630,
            Decimal(w.AVR_SHEETS[topology][1]),
            evidence,
        )


@pytest.mark.parametrize(
    "topology,total,work",
    [
        ("Г-Г", "456000", "25200"),
        ("Г-Д", "456000", "25200"),
        ("Г-Г-Д", "684000", "36750"),
    ],
)
def test_630_explicit_confirmation(topology: str, total: str, work: str) -> None:
    evidence = {
        "rating_a": 630,
        "available": True,
        "confirmed_by": "Synthetic supplier",
        "confirmation_reference": "SYNTHETIC-ONLY",
    }
    result = w.avr_price(
        avr_sheet(topology),
        topology,
        "MCCB",
        630,
        Decimal(w.AVR_SHEETS[topology][1]),
        evidence,
    )
    assert result["material_kzt"] == total and result["work"]["value"] == work
    for key, value in [
        ("rating_a", 400),
        ("available", False),
        ("available", 1),
        ("confirmed_by", ""),
        ("confirmation_reference", ""),
    ]:
        with pytest.raises(f.RuleError):
            w.avr_price(
                avr_sheet(topology),
                topology,
                "MCCB",
                630,
                Decimal(w.AVR_SHEETS[topology][1]),
                {**evidence, key: value},
            )


@pytest.mark.parametrize("rating", [63, 80, 99, 101, 125, 251, 800])
def test_no_nearest_motor_rating(rating: int) -> None:
    with pytest.raises(f.RuleError, match="exact"):
        w.avr_price(avr_sheet("Г-Г"), "Г-Г", "MCCB", rating, Decimal(2))


def test_contactor_work_is_distinct_from_motor_parts() -> None:
    result = w.avr_price(avr_sheet("Г-Г"), "Г-Г", "CONTACTOR", 65, Decimal(2))
    assert result["material_kzt"] == "38000"
    assert result["work"]["value"] == "12600"
    assert [p["role"] for p in result["parts"]] == ["CONTACTOR"]
    with pytest.raises(f.RuleError):
        w.avr_price(avr_sheet("Г-Г"), "Г-Г", "CONTACTOR", 64, Decimal(2))


@pytest.mark.parametrize(
    "cell,value",
    [
        ("K27", "125А"),
        ("L27", "=15600"),
        ("K32", "Drive?"),
        ("K3", "АВР Г-Г 100А"),
        ("L3", None),
    ],
)
def test_avr_source_semantic_drift(cell: str, value: Any) -> None:
    ws = avr_sheet("Г-Г")
    ws[cell] = value
    with pytest.raises(f.RuleError):
        w.avr_price(ws, "Г-Г", "MCCB", 100, Decimal(2))


@pytest.mark.parametrize(
    "size", [(800, 600, 250), (1200, 1000, 300), (1800, 900, 600), (500, 400, 251)]
)
def test_no_nearest_or_bigger_metal_size(size: tuple[int, int, int]) -> None:
    with pytest.raises(f.RuleError, match="whitelist"):
        w.metal_cost(size)


def test_metal_cost_does_not_read_sales_price(tmp_path: Path, monkeypatch: Any) -> None:
    path = metal22(tmp_path)
    monkeypatch.setattr(metal, "GENESIS_PATH", path)
    monkeypatch.setattr(metal, "GENESIS_SHA", sha256_file(path))
    result = w.metal_cost((500, 400, 250))
    assert result["cost_kzt"] == 12770
    assert result["source_locator"] == "Лист1!D22"
    assert result["identity"]["excluded_sales_outputs"] == ["B22", "C22"]


@pytest.mark.parametrize("code", list(w.CABINETS))
def test_cabinet_cost_and_display_use_confirmed_identity(
    tmp_path: Path, monkeypatch: Any, code: str
) -> None:
    import copy
    import json

    from build_future_case_document_source import (  # type: ignore[import-not-found]
        CALCULATION_ROLE,
        COMPOSITION_ROLE,
    )
    from openpyxl import load_workbook

    name, row, size, _, _, _ = w.CABINETS[code]
    raw = ("ЩРв" if "KRV" in code else "ЩРн") + " 18з-1 36 УХЛ3 IP31"
    source = ordinary_source(tmp_path)
    source["items"][0]["cabinet"] = {"cabinet_code": code, "cabinet_label": raw}
    bound = save_source(tmp_path, source)
    rules, bom = governed(bound)
    path = prices(tmp_path, monkeypatch)
    book = load_workbook(path)
    book["КРН"][f"L{row}"] = f"Корпус {name} {'х'.join(map(str, size))} "
    cost = 7678 if code == "CAB-KRN-18" else 9000
    book["КРН"][f"M{row}"] = cost
    book.save(path)
    book.close()
    mapping = (
        {
            "mapping_id": f"CABINET-DYNAMIC-{code}",
            "source": {
                "sheet": "КРН",
                "row": row,
                "label_cell": f"КРН!L{row}",
                "price_cells": [f"КРН!M{row}"],
            },
            "prices": {"price_kzt": cost},
        },
    )
    baseline = PriceBaseline(
        "active_approved", path, sha256_file(path), mapping_snapshot=mapping
    )
    monkeypatch.setattr(f, "resolve_active_price_baseline", lambda _: baseline)
    monkeypatch.setattr(w, "MAIN_SHA", baseline.sha256)
    result = f.calculate_future_price(
        CONTEXT,
        bom=bom,
        rules=rules,
        selector=tmp_path / "selector.json",
        cabinet_cost={**binding(tmp_path, "cost.json", cost), "role": "ENCLOSURE_COST"},
        work_source=binding(tmp_path, "work.json", 3240),
    )
    cabinet = result["cabinet_source"]
    assert cabinet["cost_kzt"] == str(cost)
    assert cabinet["source_locator"] == f"КРН!M{row}"
    assert cabinet["label_locator"] == f"КРН!L{row}"
    assert cabinet["workbook_sha256"] == sha256_file(path)
    assert cabinet["source_designation"] == raw
    assert cabinet["commercial_identity"] == name
    calc = tmp_path / "calc.json"
    calc.write_text(
        json.dumps(
            {
                "items": [
                    {
                        "item_id": source["items"][0]["item_id"],
                        "rule_application": {"price": result},
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    display_source = {
        "items": [{"item_id": source["items"][0]["item_id"], "enclosure": raw}],
        "source_bindings": [
            {
                "role": COMPOSITION_ROLE,
                "path": bound["path"],
                "sha256": bound["sha256"],
            },
            {"role": CALCULATION_ROLE, "path": str(calc), "sha256": sha256_file(calc)},
        ],
    }
    unchanged = copy.deepcopy(display_source)
    assert w.cabinet_display(display_source, 0) == name
    assert display_source == unchanged
    display_source["items"][0]["enclosure"] = name
    with pytest.raises(f.RuleError, match="designation drift"):
        w.cabinet_display(display_source, 0)
    display_source["items"][0]["enclosure"] = raw
    calc.write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="SHA"):
        w.cabinet_display(display_source, 0)
    source["items"][0]["technical_classification"]["overrides"] = {
        "human": {
            "enclosure_mm": binding(tmp_path, "size-override.json", [500, 400, 250])
        }
    }
    conflicting_rules, _ = governed(save_source(tmp_path, source))
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        with pytest.raises(f.RuleError, match="size rule conflict"):
            f.ordinary_enclosure_cost(book, baseline, conflicting_rules)
    finally:
        book.close()


@pytest.mark.parametrize("raw", ["ЩРн-18", "ЩРН 18з-1", "ЩРв-24", "ЩРВ-12", "КРН-18"])
def test_designation_text_cannot_select_cabinet(raw: str) -> None:
    with pytest.raises(f.RuleError, match="identity missing"):
        w.cabinet_display({"items": [{"enclosure": raw}], "source_bindings": []}, 0)


@pytest.mark.parametrize("code", ["CAB-KRV-9", "CAB-KRV-18?", "ЩРн 18з-1", ""])
def test_unproven_cabinet_code_holds(code: str) -> None:
    with pytest.raises(f.RuleError, match="identity mapping"):
        w.cabinet_cost(None, None, code)


def test_800_contactor_three_input_uses_exact_separate_work() -> None:
    result = w.avr_price(avr_sheet("Г-Г-Д"), "Г-Г-Д", "CONTACTOR", 800, Decimal(3))
    assert result["material_kzt"] == "846000"
    assert result["parts"][0]["source_locator"] == "АВР-Г-Г-Д!B36"
    assert result["work"]["source_locator"] == "АВР-Г-Г-Д!K8"
    assert result["work"]["value"] == "41440"
    assert result["work"]["lower_exclusive_a"] == 630
    assert "CONTACTOR AVR Г-Г-Д 800A" in f.CATEGORIES
    for cell, value in [
        ("A36", "Контактор до 630А"),
        ("J8", "АВР 3 ввода до 630А с блоком"),
        ("K8", None),
    ]:
        ws = avr_sheet("Г-Г-Д")
        ws[cell] = value
        with pytest.raises(f.RuleError):
            w.avr_price(ws, "Г-Г-Д", "CONTACTOR", 800, Decimal(3))


@pytest.mark.parametrize(
    "topology,rating,row", [("Г-Г", 800, 34), ("Г-Д", 800, 37), ("Г-Д", 1000, 38)]
)
def test_material_without_work_is_explicit_audited_hold(
    topology: str, rating: int, row: int
) -> None:
    ws = avr_sheet(topology)
    audit = w.audit_contactor(ws, topology, rating)
    assert audit["status"] == "HOLD_MISSING_TOPOLOGY_RATING_WORK_AUTHORITY"
    assert audit["material"]["source_locator"] == f"{ws.title}!B{row}"
    assert audit["work"] is None and audit["fallback_authorized"] is False
    with pytest.raises(f.RuleError, match="work rating missing"):
        w.avr_price(ws, topology, "CONTACTOR", rating, Decimal(2))
    ws[f"A{row}"] = "Unknown contactor"
    with pytest.raises(f.RuleError, match="identity drift"):
        w.audit_contactor(ws, topology, rating)


@pytest.mark.parametrize(
    "category,label,kind",
    [
        ("MCCB SHR 3P 16-63A", "ВА88 3P 63A", "mccb_up_to_100a"),
        ("MCB 3P D 80-100A", "ВА47 3P D80", "modular_3p"),
        ("MCB 1P D 10-63A", "ВА47 1P D10", "modular_1p"),
        ("MCB 3P B <=63A", "ВА47 3P B63", "modular_3p"),
        ("MCB 1P B <=63A", "ВА47 1P B6", "modular_1p"),
    ],
)
def test_exact_shr_signature(category: str, label: str, kind: str) -> None:
    component = {"install_type": kind, "component_label": label}
    f.check_protection_category(component, category)
    with pytest.raises(f.RuleError):
        f.check_protection_category(
            {**component, "component_label": label + " 40A"}, category
        )


def test_ordinary_mccb_cannot_use_small_cabinet(
    tmp_path: Path, monkeypatch: Any
) -> None:
    _, source = source_fixture(tmp_path, "ЩР")
    item = source["items"][0]
    item["cabinet"] = {"cabinet_code": "CAB-KRN-18", "cabinet_label": "ЩРн-18"}
    item["components"][0]["component_label"] = "3P 100A"
    item["technical_classification"]["categories"] = {"P1": "MCCB 3P 80-100A"}
    rules, bom = governed(save_source(tmp_path, source))
    prices(tmp_path, monkeypatch)
    with pytest.raises(f.RuleError, match="metal cabinet"):
        f.calculate_future_price(
            CONTEXT,
            bom=bom,
            rules=rules,
            selector=tmp_path / "selector.json",
            cabinet_cost={
                **binding(tmp_path, "cost.json", 7678),
                "role": "ENCLOSURE_COST",
            },
            work_source=binding(tmp_path, "work.json", 3600),
        )


def test_short_header_preserves_bottom_clause() -> None:
    clause = (
        "Предположительный срок изготовления 10–12 рабочих дней "
        "с первого рабочего дня после дня поступления средств на расчётный счёт; "
        "точный срок изготовления уточнять после оплаты."
    )
    assert (
        content.lead_time_header(clause, future=True)
        == "Срок изготовления 10-12 рабочих дней"
    )
    assert (
        " ".join(
            s.strip()
            for s in content.commercial_block(
                {"manufacturing_lead_time": clause, "delivery": "EXW"}
            )[5:7]
        )
        == clause
    )
    with pytest.raises(f.RuleError):
        content.lead_time_header("когда-нибудь", future=True)
    with pytest.raises(f.RuleError, match="canonical form"):
        content.lead_time_header(
            clause.removeprefix("Предположительный срок изготовления "), future=True
        )
    with pytest.raises(f.RuleError, match="alternative lead clause"):
        content.commercial_block(
            {
                "manufacturing_lead_time": "Срок изготовления 10-12 рабочих дней",
                "delivery": "EXW",
            }
        )


@pytest.mark.parametrize(
    "topology,rating,kind",
    [(topology, rating, "MCCB") for topology in w.AVR_SHEETS for rating in (100, 630)]
    + [("Г-Г-Д", 800, "CONTACTOR")],
)
def test_switching_mapping_in_existing_checked_calculation(
    tmp_path: Path, monkeypatch: Any, topology: str, rating: int, kind: str
) -> None:
    _, source = source_fixture(tmp_path)
    item = source["items"][0]
    power = item["components"][0]
    protection = {**power, "component_id": "P2"}
    power.update(
        component_label=(
            f"MCCB 3P {rating}A с приводом"
            if kind == "MCCB"
            else f"Контактор {rating}А"
        ),
        quantity=w.AVR_SHEETS[topology][1],
        install_type=(
            ("mccb_up_to_100a" if rating == 100 else "mccb_400a_plus")
            if kind == "MCCB"
            else "contactor"
        ),
    )
    item["components"] = [power]
    item["cabinet"] = {
        "cabinet_code": "METAL-500X400X250-D22",
        "cabinet_label": "500×400×250 IP31",
    }
    t = item["technical_classification"]
    t.update(
        execution="BREAKER_AVR" if kind == "MCCB" else "STANDARD_CONTACTOR_AVR",
        roles={"P1": "POWER_BREAKER" if kind == "MCCB" else "CONTACTOR"},
        categories={"P1": f"{kind} AVR {topology} {rating}A"},
    )
    if kind == "CONTACTOR":
        item["components"].append(protection)
        t["roles"]["P2"] = "POWER_BREAKER"
        t["categories"]["P2"] = "MCCB 3P 16-63A"
        t["circuits"][0]["protection_component_ids"] = ["P2"]
    rules, bom = governed(save_source(tmp_path, source))
    ws = avr_sheet(topology)
    if kind == "CONTACTOR":
        _, row, label = f.CATEGORIES["MCCB 3P 16-63A"]
        protection_sheet = ws.parent.create_sheet("АВР Г-Г")
        protection_sheet.cell(row, 1, label)
        protection_sheet.cell(row, 2, 15000)
    path = tmp_path / "main.xlsx"
    ws.parent.save(path)
    ws.parent.close()
    baseline = PriceBaseline("active_approved", path, sha256_file(path))
    monkeypatch.setattr(f, "resolve_active_price_baseline", lambda _: baseline)
    monkeypatch.setattr(w, "MAIN_SHA", baseline.sha256)
    metal_path = metal22(tmp_path)
    monkeypatch.setattr(metal, "GENESIS_PATH", metal_path)
    monkeypatch.setattr(metal, "GENESIS_SHA", sha256_file(metal_path))
    work = (
        (22500 if topology == "Г-Г-Д" else 12600)
        if rating == 100
        else 41440 if rating == 800 else (36750 if topology == "Г-Г-Д" else 25200)
    )
    work_record = binding(tmp_path, "work.json", work)
    cabinet = {**binding(tmp_path, "cabinet.json", 12770), "role": "ENCLOSURE_COST"}

    def calculate() -> Any:
        return f.calculate_future_price(
            CONTEXT,
            bom=bom,
            rules=rules,
            selector=tmp_path / "selector.json",
            cabinet_cost=cabinet,
            work_source=work_record,
        )

    if rating == 630:
        with pytest.raises(f.RuleError, match="availability"):
            calculate()
        import json

        evidence = binding(
            tmp_path,
            "availability.json",
            {
                "rating_a": 630,
                "available": True,
                "confirmed_by": "Synthetic supplier",
                "confirmation_reference": "TEST-ONLY",
            },
        )
        Path(work_record["path"]).write_text(
            json.dumps({"value": work, "avr_drive_availability": evidence}),
            encoding="utf-8",
        )
        work_record["sha256"] = sha256_file(Path(work_record["path"]))
    result = calculate()
    assert result["status"] == "DRAFT_NOT_APPROVED"
    assert result["work_kzt"] == str(work)
    assert result["avr_sources"][0]["input_count"] == w.AVR_SHEETS[topology][1]
    assert result["cabinet_source"]["source_locator"] == "Лист1!D22"
    if rating == 800:
        assert result["avr_sources"][0]["material_kzt"] == "846000"
        assert result["avr_sources"][0]["work"]["source_locator"] == "АВР-Г-Г-Д!K8"
    if rating == 630:
        read_bytes = Path.read_bytes
        work_path = Path(work_record["path"])
        bound_value = f.bound_value
        middle_read = False

        def mark_middle(record: Any) -> Any:
            nonlocal middle_read
            result = bound_value(record)
            if record == work_record:
                middle_read = True
            return result

        def swapped_read(path: Path) -> bytes:
            nonlocal middle_read
            raw = read_bytes(path)
            if path == work_path and middle_read:
                middle_read = False
                return raw + b" "
            return raw

        with monkeypatch.context() as transient:
            transient.setattr(Path, "read_bytes", swapped_read)
            transient.setattr(f, "bound_value", mark_middle)
            with pytest.raises(f.RuleError, match="work-source SHA drift"):
                calculate()
        Path(evidence["path"]).write_text("{}", encoding="utf-8")
        with pytest.raises(f.RuleError, match="drift"):
            calculate()


@pytest.mark.parametrize("suffix", ["D54", "D999"])
def test_metal_row_suffix_must_match_dimensions(tmp_path: Path, suffix: str) -> None:
    _, source = source_fixture(tmp_path, "ЩР")
    source["items"][0]["cabinet"]["cabinet_code"] = f"METAL-500X400X250-{suffix}"
    rules, _ = governed(save_source(tmp_path, source))
    with pytest.raises(f.RuleError, match="source row conflict"):
        f.ordinary_enclosure_cost(None, None, rules)
