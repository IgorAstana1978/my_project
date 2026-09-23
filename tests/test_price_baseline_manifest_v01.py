from __future__ import annotations

import csv
import json
import math
import subprocess
import sys
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from scripts import activate_price_baseline_manifest as activation
from scripts import audit_price_baseline_candidate as auditor
from scripts import calc_quote_price_draft as calculator
from scripts import price_baseline_contract as contract


def write_governed_workbook(
    path: Path,
    *,
    material_delta: int = 0,
    unused_delta: int = 0,
) -> None:
    workbook = Workbook()
    krn = workbook.active
    krn.title = "КРН"
    workbook.create_sheet("ЩР")
    for mapping in calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS:
        sheet = workbook[mapping.sheet_name]
        sheet.cell(mapping.row, 1, mapping.expected_label)
        sheet.cell(
            mapping.row,
            2,
            mapping.expected_material_price
            + (material_delta if mapping.mapping_id == "COMPONENT-MAPPING-012" else 0),
        )
        sheet.cell(mapping.row, 3, mapping.expected_work_price)
    if material_delta:
        krn["B5"] = 4500 + material_delta
    occupied_component_rows = {
        calculator.normalize_workbook_label(krn.cell(row, 1).value): row
        for row in range(1, calculator.MAX_LOOKUP_ROW + 1)
        if calculator.normalize_workbook_label(krn.cell(row, 1).value) is not None
    }
    next_row = 20
    for definition in calculator.COMPONENT_DEFINITIONS.values():
        if definition.workbook_label is None:
            continue
        label = calculator.normalize_workbook_label(definition.workbook_label)
        if label in occupied_component_rows:
            continue
        krn.cell(next_row, 1, definition.workbook_label)
        dynamic_delta = (
            material_delta if definition.workbook_label == "ВА47 1 полюсный" else 0
        )
        krn.cell(next_row, 2, 1000 + next_row + dynamic_delta)
        krn.cell(next_row, 3, 200 + next_row)
        occupied_component_rows[label] = next_row
        next_row += 1
    for cabinet_mapping in calculator.APPROVED_CABINET_PRICE_MAPPINGS:
        sheet = workbook[cabinet_mapping.sheet_name]
        sheet.cell(cabinet_mapping.row, 12, cabinet_mapping.expected_label)
        sheet.cell(cabinet_mapping.row, 13, cabinet_mapping.expected_price)
    occupied_cabinet_labels = {
        calculator.normalize_workbook_label(krn.cell(row, 12).value)
        for row in range(1, calculator.MAX_LOOKUP_ROW + 1)
    }
    cabinet_row = 10
    for code, label in calculator.CABINET_DEFINITIONS.items():
        if code == calculator.INVOICE519_SCHE_CABINET_CODE:
            continue
        normalized = calculator.normalize_workbook_label(label)
        if normalized in occupied_cabinet_labels:
            continue
        krn.cell(cabinet_row, 12, label)
        krn.cell(cabinet_row, 13, 8000 + cabinet_row)
        occupied_cabinet_labels.add(normalized)
        cabinet_row += 1
    krn["A29"] = "UNUSED PRICE ROW"
    krn["B29"] = 7777 + unused_delta
    krn["C29"] = 333
    workbook.save(path)
    workbook.close()


def add_busbar_labor_block(path: Path, *, first_row: int = 49) -> None:
    workbook = load_workbook(path)
    sheet = workbook["ЩР"]
    for offset, size in enumerate(("3х30", "4х40", "5х50")):
        row = first_row + offset
        sheet.cell(row, 1, f"{size} мм шина АЛ")
        sheet.cell(row, 2, 950 + offset * 100)
        sheet.cell(row, 3, 0)
        sheet.cell(row, 6, f"=D{row}*C{row}")
    labor_row = first_row + 3
    sheet.cell(labor_row, 1, "Шина за работу")
    sheet.cell(labor_row, 3, 3000)
    sheet.cell(labor_row, 6, f"=D{labor_row}*C{labor_row}")
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (1, 1),
        (6899.999999999999, 6900),
        (6899.999999999999, 6900),
        (13799.999999999998, 13800),
        (14949.999999999998, 14950),
        (32199.999999999996, 32200),
        (math.nextafter(6900.0, -math.inf), 6900),
        (math.nextafter(6900.0, math.inf), 6900),
        (math.nextafter(math.nextafter(6900.0, math.inf), math.inf), None),
        (math.nextafter(math.nextafter(6900.0, -math.inf), -math.inf), None),
        (6899.99, None),
        (6899.5, None),
        (6900.0001, None),
        (0, None),
        (-1, None),
        (float("nan"), None),
        (float("inf"), None),
        (-float("inf"), None),
        (True, None),
        ("6900", None),
        ("=6900", None),
    ],
)
def test_shared_kzt_literal_normalizer(value: Any, expected: int | None) -> None:
    assert contract.normalize_kzt_literal(value) == expected


@pytest.mark.parametrize("role", ["material", "cabinet", "work"])
def test_zero_outside_structural_work_class_holds(tmp_path: Path, role: str) -> None:
    path = tmp_path / "zero-role.xlsx"
    write_governed_workbook(path)
    workbook = load_workbook(path)
    coordinate = {"material": "B29", "cabinet": "M10", "work": "C29"}[role]
    workbook["КРН"][coordinate] = 0
    workbook.save(path)
    workbook.close()
    inspection = auditor.inspect_workbook(path)
    assert any(coordinate in finding for finding in inspection.invalid_price_cells)


@pytest.mark.parametrize(
    "drift", ["none", "labor_label", "labor_rate", "labor_formula", "busbar_label"]
)
def test_busbar_zero_requires_separate_labor_structure(
    tmp_path: Path, drift: str
) -> None:
    path = tmp_path / "busbar.xlsx"
    write_governed_workbook(path)
    add_busbar_labor_block(path)
    workbook = load_workbook(path)
    sheet = workbook["ЩР"]
    if drift == "labor_label":
        sheet["A52"] = "other labor"
    elif drift == "labor_rate":
        sheet["C52"] = None
    elif drift == "labor_formula":
        sheet["F52"] = "=0"
    elif drift == "busbar_label":
        sheet["A50"] = "ordinary component"
    workbook.save(path)
    workbook.close()
    inspection = auditor.inspect_workbook(path)
    zero_findings = [
        finding for finding in inspection.invalid_price_cells if "ЩР!C49" in finding
    ]
    assert bool(zero_findings) is (drift != "none")
    assert (
        sum(
            cell.endswith("C49") or cell.endswith("C50") or cell.endswith("C51")
            for entry in inspection.price_names
            for cell in entry["price_cells"]
        )
        == 3
    )


def test_auditor_and_activator_share_ulp_literal_semantics(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    workbook = load_workbook(candidate)
    workbook["КРН"]["B5"] = math.nextafter(4600.0, -math.inf)
    workbook.save(candidate)
    workbook.close()
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    assert audit["status"] == "PASS_CANDIDATE_ONLY"
    activation.verify_snapshot_against_workbook(candidate, audit["mapping_snapshot"])
    assert not auditor.inspect_workbook(candidate).invalid_price_cells


@pytest.mark.parametrize("change", ["added", "removed"])
def test_unused_genuine_price_identity_change_holds(
    active_chain: dict[str, Path], tmp_path: Path, change: str
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    workbook = load_workbook(candidate)
    sheet = workbook["КРН"]
    if change == "added":
        sheet["A6"] = "NEW UNUSED PRICE"
        sheet["B6"] = 1234
        sheet["C6"] = 234
    else:
        sheet["A29"] = None
        sheet["B29"] = None
        sheet["C29"] = None
    workbook.save(candidate)
    workbook.close()
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    assert audit["status"] == "HOLD"
    assert "price-bearing names were added or removed" in audit["hold_reasons"]
    assert audit["added_names"] if change == "added" else audit["removed_names"]


def test_unproven_region_values_do_not_define_prices(tmp_path: Path) -> None:
    path = tmp_path / "outside-price-regions.xlsx"
    write_governed_workbook(path)
    workbook = load_workbook(path)
    side = workbook.create_sheet("ВРУ250А")
    side["A2"] = "BASE MATERIAL"
    side["B2"] = 1000
    side["J2"] = "BASE CABINET"
    side["K2"] = 5000
    side["J5"] = "800х600 dimension"
    side["K5"] = "2м"
    side["A16"] = "OUTSIDE MATERIAL"
    side["B16"] = 1200
    side["A17"] = "OUTSIDE FORMULA"
    side["B17"] = "=2*1000"
    side["J12"] = "OUTSIDE CABINET"
    side["K12"] = 5500
    side["J13"] = "OUTSIDE CABINET FORMULA"
    side["K13"] = "=2*1000"
    workbook.save(path)
    workbook.close()
    inspection = auditor.inspect_workbook(path)
    assert [
        (entry["kind"], entry["row"])
        for entry in inspection.price_names
        if entry["sheet"] == "ВРУ250А"
    ] == [("component", 2), ("cabinet", 2)]
    assert not inspection.formula_cells
    assert not inspection.invalid_price_cells


def build_snapshot(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return calculator.build_governed_mapping_snapshot(
            workbook,
            component_mappings=calculator.SUCCESSOR_COMPONENT_PRICE_MAPPINGS,
        )
    finally:
        workbook.close()


def publish_initial_chain(root: Path, workbook: Path) -> tuple[Path, Path]:
    manifests = root / "manifests"
    manifests.mkdir(parents=True)
    manifest_path = manifests / "PBM-INITIAL.json"
    manifest = {
        "schema_version": contract.MANIFEST_SCHEMA,
        "manifest_id": "PBM-INITIAL",
        "created_at": "2099-01-01T00:00:00Z",
        "predecessor": None,
        "workbook": {
            "path": str(workbook.resolve()),
            "sha256": contract.sha256_file(workbook),
            "structural_fingerprint": auditor.inspect_workbook(
                workbook
            ).structural_fingerprint,
        },
        "mapping_snapshot": build_snapshot(workbook),
        "approval": {
            "schema_version": contract.APPROVAL_SCHEMA,
            "authority": contract.APPROVAL_AUTHORITY,
            "approval_id": "IGOR-BASELINE-INITIAL",
            "approved_by": "Igor",
            "approved_at": "2099-01-01T00:00:00Z",
            "candidate_audit_sha256": "a" * 64,
            "approval_fingerprint": "b" * 64,
        },
        "scope": {
            "future_cases_only": True,
            "historical_repricing_authorized": False,
        },
    }
    contract.validate_manifest(manifest)
    manifest_raw = contract.canonical_json_bytes(manifest)
    manifest_path.write_bytes(manifest_raw)
    selector_path = root / "active-price-baseline.json"
    selector = {
        "schema_version": contract.SELECTOR_SCHEMA,
        "manifest_id": "PBM-INITIAL",
        "manifest_path": str(manifest_path.resolve()),
        "manifest_sha256": contract.sha256_bytes(manifest_raw),
        "activated_at": "2099-01-01T00:00:00Z",
        "approval_id": "IGOR-BASELINE-INITIAL",
    }
    selector_path.write_bytes(contract.canonical_json_bytes(selector))
    return selector_path, manifest_path


@pytest.fixture
def active_chain(tmp_path: Path) -> dict[str, Path]:
    root = tmp_path / "prices"
    current = root / "current"
    current.mkdir(parents=True)
    workbook = current / "approved.xlsx"
    write_governed_workbook(workbook)
    selector, manifest = publish_initial_chain(root, workbook)
    return {
        "root": root,
        "workbook": workbook,
        "selector": selector,
        "manifest": manifest,
    }


def write_candidate(active_chain: dict[str, Path], tmp_path: Path) -> Path:
    candidate = active_chain["root"] / "current" / "candidate.xlsx"
    write_governed_workbook(candidate, material_delta=100)
    return candidate


def test_active_selector_manifest_workbook_chain_is_exact_and_fail_closed(
    active_chain: dict[str, Path],
) -> None:
    resolved = contract.resolve_active_price_baseline(active_chain["selector"])
    assert resolved.version == contract.ACTIVE_VERSION
    assert resolved.path == active_chain["workbook"].resolve()
    assert resolved.manifest_id == "PBM-INITIAL"
    assert resolved.mapping_snapshot

    selector = json.loads(active_chain["selector"].read_text(encoding="utf-8"))
    selector["manifest_sha256"] = "0" * 64
    active_chain["selector"].write_bytes(contract.canonical_json_bytes(selector))
    with pytest.raises(contract.BaselineContractError, match="manifest SHA-256"):
        contract.resolve_active_price_baseline(active_chain["selector"])


def test_corrupt_manifest_and_missing_selector_fail_closed(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    active_chain["manifest"].write_text("corrupt", encoding="utf-8")
    with pytest.raises(contract.BaselineContractError, match="manifest SHA-256"):
        contract.resolve_active_price_baseline(active_chain["selector"])
    with pytest.raises(contract.BaselineContractError, match="does not exist"):
        contract.resolve_active_price_baseline(tmp_path / "missing-selector.json")


def test_candidate_audit_is_read_only_and_separates_identity_from_prices(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    candidate_before = candidate.read_bytes()
    selector_before = active_chain["selector"].read_bytes()
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    assert audit["status"] == "PASS_CANDIDATE_ONLY"
    assert audit["hold_reasons"] == []
    assert audit["changed_prices"]
    assert all(change["currently_governed"] for change in audit["changed_prices"])
    assert "COMPONENT-MAPPING-012" in {
        mapping_id
        for change in audit["changed_prices"]
        for mapping_id in change["governed_mapping_ids"]
    }
    assert audit["mapping_snapshot"] != []
    assert all(
        "material_kzt" not in entry["identity"]
        and "work_kzt" not in entry["identity"]
        and "price_kzt" not in entry["identity"]
        for entry in audit["mapping_snapshot"]
    )
    assert candidate.read_bytes() == candidate_before
    assert active_chain["selector"].read_bytes() == selector_before


def test_unused_price_row_change_is_visible_without_becoming_mapping_authority(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = active_chain["root"] / "current" / "unused-change.xlsx"
    write_governed_workbook(candidate, unused_delta=100)
    predecessor = contract.resolve_active_price_baseline(active_chain["selector"])
    audit = auditor.audit_candidate(active_chain["selector"], candidate)

    assert audit["status"] == "PASS_CANDIDATE_ONLY"
    assert audit["structural_compatibility"] is True
    assert audit["added_names"] == []
    assert audit["removed_names"] == []
    assert audit["changed_prices"] == [
        {
            "kind": "component",
            "sheet": "КРН",
            "row": 29,
            "label": "UNUSED PRICE ROW",
            "label_cell": "КРН!A29",
            "price_cells": ["КРН!B29", "КРН!C29"],
            "before": {"КРН!B29": 7777, "КРН!C29": 333},
            "after": {"КРН!B29": 7877, "КРН!C29": 333},
            "currently_governed": False,
            "governed_mapping_ids": [],
        }
    ]
    assert audit["mapping_snapshot"] == [
        dict(entry) for entry in predecessor.mapping_snapshot
    ]
    assert audit["approval_payload"]["candidate_workbook_sha256"] == (
        contract.sha256_file(candidate)
    )

    audit_path = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    approval_path = tmp_path / "unused-row-approval.json"
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    workbook = load_workbook(candidate)
    workbook["КРН"]["B29"] = 7977
    workbook.save(candidate)
    workbook.close()
    with pytest.raises(activation.ActivationError, match="SHA-256 drifted"):
        activation.activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=active_chain["selector"],
        )


@pytest.mark.parametrize(
    "invalid_value",
    [None, 0, -1, 1.5, "not a price"],
)
def test_unused_price_row_invalid_value_holds_without_mapping_expansion(
    active_chain: dict[str, Path], tmp_path: Path, invalid_value: Any
) -> None:
    candidate = active_chain["root"] / "current" / "unused-invalid.xlsx"
    write_governed_workbook(candidate)
    workbook = load_workbook(candidate)
    workbook["КРН"]["B29"] = invalid_value
    workbook.save(candidate)
    workbook.close()

    predecessor = contract.resolve_active_price_baseline(active_chain["selector"])
    audit = auditor.audit_candidate(active_chain["selector"], candidate)

    assert audit["status"] == "HOLD"
    assert "missing conflicts detected" in audit["hold_reasons"]
    assert any(
        finding.startswith("КРН!B29:") for finding in audit["conflicts"]["missing"]
    )
    assert audit["changed_prices"][0]["kind"] == "component"
    assert audit["changed_prices"][0]["row"] == 29
    assert audit["changed_prices"][0]["after"]["КРН!B29"] == invalid_value
    assert audit["changed_prices"][0]["currently_governed"] is False
    assert audit["changed_prices"][0]["governed_mapping_ids"] == []
    assert audit["mapping_snapshot"] == [
        dict(entry) for entry in predecessor.mapping_snapshot
    ]


@pytest.mark.parametrize("drift", ["formula", "missing", "identity", "duplicate"])
def test_candidate_audit_holds_formula_identity_and_duplicate_drift(
    active_chain: dict[str, Path], tmp_path: Path, drift: str
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    workbook = load_workbook(candidate)
    krn = workbook["КРН"]
    if drift == "formula":
        krn["B5"] = "=4500"
    elif drift == "missing":
        krn["B5"] = None
    elif drift == "identity":
        krn["A5"] = "different technical identity"
    else:
        krn["A6"] = krn["A5"].value
        krn["B6"] = krn["B5"].value
        krn["C6"] = krn["C5"].value
    workbook.save(candidate)
    workbook.close()
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    assert audit["status"] == "HOLD"
    assert audit["hold_reasons"]
    if drift == "identity":
        assert audit["conflicts"]["mapping_identity_drift"]
    else:
        assert audit["conflicts"][drift]


def approval_for(audit_path: Path, audit: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema_version": contract.APPROVAL_SCHEMA,
        "authority": contract.APPROVAL_AUTHORITY,
        "approval_id": "IGOR-BASELINE-CANDIDATE",
        "approved_by": "Igor",
        "approved_at": "2099-01-02T00:00:00Z",
        "candidate_audit_sha256": contract.sha256_file(audit_path),
        "approval_fingerprint": audit["approval_fingerprint"],
    }


def activate_candidate(
    active_chain: dict[str, Path], tmp_path: Path
) -> tuple[Path, dict[str, Any], activation.ActivationResult]:
    candidate = write_candidate(active_chain, tmp_path)
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    audit_path = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    approval_path = tmp_path / "approval.json"
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    result = activation.activate(
        candidate_audit_json=audit_path,
        approval_json=approval_path,
        active_selector=active_chain["selector"],
    )
    return candidate, audit, result


def test_activation_is_immutable_atomic_and_exactly_approved(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate, audit, result = activate_candidate(active_chain, tmp_path)
    assert result.manifest_path.is_file()
    assert contract.sha256_file(result.manifest_path) == result.manifest_sha256
    resolved = contract.resolve_active_price_baseline(active_chain["selector"])
    assert resolved.path == candidate.resolve()
    assert resolved.manifest_id == audit["manifest_id"]
    assert resolved.approval_id == "IGOR-BASELINE-CANDIDATE"
    with pytest.raises(activation.ActivationError):
        activation.activate(
            candidate_audit_json=contract.audit_artifact_path(
                active_chain["selector"], audit
            ),
            approval_json=tmp_path / "approval.json",
            active_selector=active_chain["selector"],
        )


def test_activation_rejects_approval_not_bound_to_exact_audit(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    audit_path = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    approval = approval_for(audit_path, audit)
    approval["candidate_audit_sha256"] = "0" * 64
    approval_path = tmp_path / "approval.json"
    approval_path.write_bytes(contract.canonical_json_bytes(approval))
    with pytest.raises(activation.ActivationError, match="exact candidate audit"):
        activation.activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=active_chain["selector"],
        )


def test_ordinary_activation_rejects_arbitrary_audit_path(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    canonical = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    wrong = tmp_path / "arbitrary-audit.json"
    wrong.write_bytes(canonical.read_bytes())
    approval_path = tmp_path / "approval.json"
    approval_path.write_bytes(contract.canonical_json_bytes(approval_for(wrong, audit)))
    with pytest.raises(activation.ActivationError, match="canonical path"):
        activation.activate(
            candidate_audit_json=wrong,
            approval_json=approval_path,
            active_selector=active_chain["selector"],
        )
    assert (
        contract.resolve_active_price_baseline(active_chain["selector"]).manifest_id
        == "PBM-INITIAL"
    )


def test_activation_rejects_mapping_identity_change_even_when_approved(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    audit["mapping_snapshot"][0]["identity"]["expected_label"] = "new identity"
    audit["mapping_snapshot"][0]["identity_fingerprint"] = (
        contract.mapping_identity_fingerprint(audit["mapping_snapshot"][0]["identity"])
    )
    audit["approval_payload"]["mapping_snapshot_sha256"] = contract.sha256_bytes(
        contract.canonical_json_bytes(audit["mapping_snapshot"])
    )
    audit["approval_fingerprint"] = contract.sha256_bytes(
        contract.canonical_json_bytes(audit["approval_payload"])
    )
    audit_path = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    approval_path = tmp_path / "approval.json"
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    with pytest.raises(activation.ActivationError, match="mapping identity drift"):
        activation.activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=active_chain["selector"],
        )


def test_activation_rejects_workbook_toctou_drift(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    audit = auditor.audit_candidate(active_chain["selector"], candidate)
    audit_path = auditor.materialize_audit_artifact(active_chain["selector"], audit)
    approval_path = tmp_path / "approval.json"
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    workbook = load_workbook(candidate)
    workbook["КРН"]["B5"] = 9999
    workbook.save(candidate)
    workbook.close()
    with pytest.raises(activation.ActivationError, match="SHA-256 drifted"):
        activation.activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=active_chain["selector"],
        )


def write_input_csv(path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, delimiter=";", lineterminator="\n")
        writer.writerow(calculator.REQUIRED_COLUMNS)
        writer.writerow(
            [
                "TEST",
                "CAB-KRN-24",
                "1.20",
                "EKF-VA47-29-1P",
                "1",
                "modular_1p",
            ]
        )


def test_future_calculation_follows_active_manifest_without_python_edit(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    input_csv = tmp_path / "input.csv"
    write_input_csv(input_csv)
    before = calculator.calculate_price_draft(
        None,
        input_csv,
        active_selector_path=active_chain["selector"],
    )
    assert before.status == "PASS"
    candidate, _, _ = activate_candidate(active_chain, tmp_path)
    after = calculator.calculate_price_draft(
        None,
        input_csv,
        active_selector_path=active_chain["selector"],
    )
    assert after.status == "PASS"
    assert after.price_workbook == candidate.resolve()
    assert after.price_baseline_manifest_id is not None
    assert after.total_preliminary_price != before.total_preliminary_price


def test_explicit_historical_binding_never_reads_active_selector(
    active_chain: dict[str, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    historical = tmp_path / "historical.xlsx"
    historical.write_bytes(b"frozen historical bytes")
    frozen = contract.PriceBaseline(
        contract.HISTORICAL.version,
        historical,
        contract.sha256_file(historical),
    )
    monkeypatch.setitem(contract.BASELINES, contract.HISTORICAL.version, frozen)
    active_chain["selector"].write_text("corrupt", encoding="utf-8")
    assert (
        contract.require_price_baseline(historical, contract.HISTORICAL.version)
        == frozen
    )


def test_manifest_schema_file_is_valid_json() -> None:
    schema = (
        Path(__file__).resolve().parents[1]
        / "schemas"
        / "price_baseline_manifest_v0_1.schema.json"
    )
    payload = json.loads(schema.read_text(encoding="utf-8"))
    assert payload["properties"]["schema_version"]["const"] == (
        contract.MANIFEST_SCHEMA
    )


def bootstrap_case(tmp_path: Path) -> dict[str, Any]:
    root = tmp_path / "bootstrap-prices"
    current = root / "current"
    current.mkdir(parents=True)
    workbook = current / "approved-successor.xlsx"
    write_governed_workbook(workbook)
    approved = contract.PriceBaseline(
        contract.SUCCESSOR.version,
        workbook.resolve(),
        contract.sha256_file(workbook),
    )
    return {
        "root": root,
        "workbook": workbook,
        "selector": root / "active-price-baseline.json",
        "approved": approved,
    }


def write_bootstrap_audit_and_approval(
    case: dict[str, Any],
    tmp_path: Path,
) -> tuple[dict[str, Any], Path, Path]:
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=case["approved"],
    )
    audit_path = auditor.materialize_audit_artifact(case["selector"], audit)
    approval_path = tmp_path / "bootstrap-approval.json"
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    return audit, audit_path, approval_path


def test_bootstrap_audit_artifact_is_deterministic_immutable_and_canonical(
    tmp_path: Path,
) -> None:
    case = bootstrap_case(tmp_path)
    audit = auditor.audit_bootstrap_candidate(
        case["selector"], case["workbook"], approved_successor=case["approved"]
    )
    expected = contract.audit_artifact_path(case["selector"], audit)
    assert expected.parent == case["root"] / "audits"
    assert audit["manifest_id"] in expected.name
    assert audit["candidate_workbook"]["sha256"] in expected.name
    assert not expected.exists()

    path = auditor.materialize_audit_artifact(case["selector"], audit)
    raw = contract.canonical_json_bytes(audit)
    assert path == expected
    assert path.read_bytes() == raw
    assert contract.sha256_file(path) == contract.sha256_bytes(raw)
    assert auditor.materialize_audit_artifact(case["selector"], audit) == path
    assert path.read_bytes() == raw

    different = json.loads(raw)
    different["approval_fingerprint"] = "0" * 64
    with pytest.raises(
        auditor.CandidateAuditError, match="existing audit artifact differs"
    ):
        auditor.materialize_audit_artifact(case["selector"], different)
    assert path.read_bytes() == raw
    assert not case["selector"].exists()
    assert not (case["root"] / "manifests").exists()


def test_bootstrap_activation_rejects_noncanonical_audit_path_and_bytes(
    tmp_path: Path,
) -> None:
    case = bootstrap_case(tmp_path)
    audit, audit_path, approval_path = write_bootstrap_audit_and_approval(
        case, tmp_path
    )
    wrong_path = tmp_path / "arbitrary-audit.json"
    wrong_path.write_bytes(audit_path.read_bytes())
    with pytest.raises(activation.ActivationError, match="canonical path"):
        activation.bootstrap_activate(
            candidate_audit_json=wrong_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
        )

    audit_path.write_bytes(audit_path.read_bytes() + b" ")
    approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(audit_path, audit))
    )
    with pytest.raises(activation.ActivationError, match="canonical JSON bytes"):
        activation.bootstrap_activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
        )
    assert not case["selector"].exists()
    assert not (case["root"] / "manifests").exists()


def test_candidate_auditor_cli_materializes_only_canonical_bytes(
    active_chain: dict[str, Path], tmp_path: Path
) -> None:
    candidate = write_candidate(active_chain, tmp_path)
    script = (
        Path(__file__).resolve().parents[1]
        / "scripts"
        / "audit_price_baseline_candidate.py"
    )
    command = [
        sys.executable,
        str(script),
        "--active-selector",
        str(active_chain["selector"]),
        "--candidate-workbook",
        str(candidate),
    ]
    preview = subprocess.run(
        command,
        capture_output=True,
        check=False,
    )
    assert preview.returncode == 0, preview.stderr.decode(errors="replace")
    audit = json.loads(preview.stdout)
    path = contract.audit_artifact_path(active_chain["selector"], audit)
    assert preview.stdout == contract.canonical_json_bytes(audit)
    assert not path.exists()

    result = subprocess.run(
        command + ["--materialize-audit"], capture_output=True, check=False
    )
    assert result.returncode == 0, result.stderr.decode(errors="replace")
    assert result.stdout == preview.stdout == path.read_bytes()
    assert audit["status"] == "PASS_CANDIDATE_ONLY"
    assert not (
        active_chain["root"] / "manifests" / f"{audit['manifest_id']}.json"
    ).exists()


def bootstrap_synthetic_case(
    case: dict[str, Any],
    tmp_path: Path,
) -> tuple[dict[str, Any], activation.ActivationResult]:
    audit, audit_path, approval_path = write_bootstrap_audit_and_approval(
        case, tmp_path
    )
    result = activation.bootstrap_activate(
        candidate_audit_json=audit_path,
        approval_json=approval_path,
        active_selector=case["selector"],
        approved_successor=case["approved"],
    )
    return audit, result


def test_bootstrap_audit_passes_only_exact_static_successor(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=case["approved"],
    )
    assert audit["schema_version"] == auditor.BOOTSTRAP_AUDIT_SCHEMA
    assert audit["intent"] == auditor.BOOTSTRAP_INTENT
    assert audit["status"] == "PASS_CANDIDATE_ONLY"
    assert audit["hold_reasons"] == []
    assert audit["genesis_state"]["selector_exists"] is False
    assert audit["genesis_state"]["manifest_entries"] == []
    assert audit["candidate_workbook"]["sha256"] == case["approved"].sha256
    assert audit["mapping_snapshot"] == build_snapshot(case["workbook"])
    assert audit["approval_payload"]["intent"] == "FIRST_CHAIN_BOOTSTRAP"
    assert audit["approval_payload"]["future_cases_only"] is True
    assert audit["approval_payload"]["historical_repricing_authorized"] is False
    assert case["selector"].parent != contract.DEFAULT_ACTIVE_SELECTOR.parent


def test_bootstrap_creates_genesis_chain_then_ordinary_activation_works(
    tmp_path: Path,
) -> None:
    case = bootstrap_case(tmp_path)
    bootstrap_audit, bootstrap_result = bootstrap_synthetic_case(case, tmp_path)
    manifest = json.loads(bootstrap_result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["predecessor"] is None
    resolved = contract.resolve_active_price_baseline(case["selector"])
    assert resolved.path == case["workbook"].resolve()
    assert resolved.manifest_id == bootstrap_audit["manifest_id"]

    candidate = case["root"] / "current" / "ordinary-successor.xlsx"
    write_governed_workbook(candidate, material_delta=100)
    ordinary_audit = auditor.audit_candidate(case["selector"], candidate)
    assert ordinary_audit["status"] == "PASS_CANDIDATE_ONLY"
    ordinary_audit_path = auditor.materialize_audit_artifact(
        case["selector"], ordinary_audit
    )
    ordinary_approval_path = tmp_path / "ordinary-approval.json"
    ordinary_approval_path.write_bytes(
        contract.canonical_json_bytes(approval_for(ordinary_audit_path, ordinary_audit))
    )
    activation.activate(
        candidate_audit_json=ordinary_audit_path,
        approval_json=ordinary_approval_path,
        active_selector=case["selector"],
    )
    assert contract.resolve_active_price_baseline(case["selector"]).path == (
        candidate.resolve()
    )


@pytest.mark.parametrize("existing_kind", ["selector", "manifest"])
def test_bootstrap_audit_holds_nonempty_genesis_state(
    tmp_path: Path, existing_kind: str
) -> None:
    case = bootstrap_case(tmp_path)
    if existing_kind == "selector":
        case["selector"].write_text("existing", encoding="utf-8")
    else:
        manifests = case["root"] / "manifests"
        manifests.mkdir()
        (manifests / "existing.json").write_text("{}", encoding="utf-8")
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=case["approved"],
    )
    assert audit["status"] == "HOLD"
    assert any(existing_kind in reason for reason in audit["hold_reasons"])


def test_bootstrap_audit_holds_wrong_workbook_path_and_sha(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    wrong_path = case["root"] / "current" / "other.xlsx"
    write_governed_workbook(wrong_path)
    wrong_path_audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        wrong_path,
        approved_successor=case["approved"],
    )
    assert wrong_path_audit["status"] == "HOLD"
    assert any("path differs" in reason for reason in wrong_path_audit["hold_reasons"])

    wrong_sha = contract.PriceBaseline(
        case["approved"].version,
        case["approved"].path,
        "0" * 64,
    )
    wrong_sha_audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=wrong_sha,
    )
    assert wrong_sha_audit["status"] == "HOLD"
    assert any(
        "SHA-256 differs" in reason for reason in wrong_sha_audit["hold_reasons"]
    )


def test_bootstrap_audit_holds_mapping_identity_drift(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    workbook = load_workbook(case["workbook"])
    workbook["КРН"]["A5"] = "different mapping identity"
    workbook.save(case["workbook"])
    workbook.close()
    drift_approved = contract.PriceBaseline(
        case["approved"].version,
        case["approved"].path,
        contract.sha256_file(case["workbook"]),
    )
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=drift_approved,
    )
    assert audit["status"] == "HOLD"
    assert audit["conflicts"]["mapping_identity_drift"]


@pytest.mark.parametrize(
    "invalid_value",
    ["=4500", None, 0, -1, 1.5, "not a price"],
)
def test_bootstrap_audit_holds_formula_and_invalid_prices(
    tmp_path: Path, invalid_value: Any
) -> None:
    case = bootstrap_case(tmp_path)
    workbook = load_workbook(case["workbook"])
    workbook["КРН"]["B29"] = invalid_value
    workbook.save(case["workbook"])
    workbook.close()
    changed_approved = contract.PriceBaseline(
        case["approved"].version,
        case["approved"].path,
        contract.sha256_file(case["workbook"]),
    )
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=changed_approved,
    )
    assert audit["status"] == "HOLD"
    category = "formula" if invalid_value == "=4500" else "missing"
    assert audit["conflicts"][category]


def test_bootstrap_audit_holds_duplicate_price_identity(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    workbook = load_workbook(case["workbook"])
    krn = workbook["КРН"]
    krn["A6"] = krn["A29"].value
    krn["B6"] = 7777
    krn["C6"] = 333
    workbook.save(case["workbook"])
    workbook.close()
    changed_approved = contract.PriceBaseline(
        case["approved"].version,
        case["approved"].path,
        contract.sha256_file(case["workbook"]),
    )
    audit = auditor.audit_bootstrap_candidate(
        case["selector"],
        case["workbook"],
        approved_successor=changed_approved,
    )
    assert audit["status"] == "HOLD"
    assert audit["conflicts"]["duplicate"]


def test_bootstrap_activation_rejects_approval_mismatch(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    audit, audit_path, approval_path = write_bootstrap_audit_and_approval(
        case, tmp_path
    )
    approval = json.loads(approval_path.read_text(encoding="utf-8"))
    approval["approval_fingerprint"] = "0" * 64
    approval_path.write_bytes(contract.canonical_json_bytes(approval))
    with pytest.raises(activation.ActivationError, match="candidate content"):
        activation.bootstrap_activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
        )
    assert audit["status"] == "PASS_CANDIDATE_ONLY"


@pytest.mark.parametrize("drift", ["workbook", "audit", "approval", "selector"])
def test_bootstrap_activation_rejects_toctou_drift(
    tmp_path: Path, drift: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    case = bootstrap_case(tmp_path)
    audit, audit_path, approval_path = write_bootstrap_audit_and_approval(
        case, tmp_path
    )
    if drift == "workbook":
        workbook = load_workbook(case["workbook"])
        workbook["КРН"]["B29"] = 8888
        workbook.save(case["workbook"])
        workbook.close()
    elif drift == "audit":
        audit_path.write_text(
            audit_path.read_text(encoding="utf-8") + " ", encoding="utf-8"
        )
    elif drift == "approval":
        original_snapshot_builder = activation.build_successor_mapping_snapshot
        mutated = False

        def mutate_approval_during_activation(path: Path) -> list[dict[str, Any]]:
            nonlocal mutated
            snapshot = cast(list[dict[str, Any]], original_snapshot_builder(path))
            if not mutated:
                mutated = True
                approval_path.write_text(
                    approval_path.read_text(encoding="utf-8") + " ",
                    encoding="utf-8",
                )
            return snapshot

        monkeypatch.setattr(
            activation,
            "build_successor_mapping_snapshot",
            mutate_approval_during_activation,
        )
    else:
        case["selector"].write_text("appeared", encoding="utf-8")
    with pytest.raises((activation.ActivationError, contract.BaselineContractError)):
        activation.bootstrap_activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
        )
    assert audit["status"] == "PASS_CANDIDATE_ONLY"


def test_bootstrap_exact_orphan_retry_is_deterministic(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    audit, audit_path, approval_path = write_bootstrap_audit_and_approval(
        case, tmp_path
    )

    def interrupted_selector_publish(path: Path, content: bytes) -> None:
        del path, content
        raise activation.ActivationError("synthetic interruption")

    with pytest.raises(activation.ActivationError, match="synthetic interruption"):
        activation.bootstrap_activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
            selector_publish_fn=interrupted_selector_publish,
        )
    assert not case["selector"].exists()
    manifest_path = case["root"] / "manifests" / f"{audit['manifest_id']}.json"
    orphan_before = manifest_path.read_bytes()

    result = activation.bootstrap_activate(
        candidate_audit_json=audit_path,
        approval_json=approval_path,
        active_selector=case["selector"],
        approved_successor=case["approved"],
    )
    assert result.manifest_path.read_bytes() == orphan_before
    assert contract.resolve_active_price_baseline(case["selector"]).path == (
        case["workbook"].resolve()
    )


def test_bootstrap_rejects_unrelated_orphan_manifest(tmp_path: Path) -> None:
    case = bootstrap_case(tmp_path)
    _, audit_path, approval_path = write_bootstrap_audit_and_approval(case, tmp_path)
    manifests = case["root"] / "manifests"
    manifests.mkdir()
    (manifests / "unrelated.json").write_text("{}", encoding="utf-8")
    with pytest.raises(activation.ActivationError, match="not empty"):
        activation.bootstrap_activate(
            candidate_audit_json=audit_path,
            approval_json=approval_path,
            active_selector=case["selector"],
            approved_successor=case["approved"],
        )


def test_historical_binding_ignores_bootstrap_selector(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    case = bootstrap_case(tmp_path)
    bootstrap_synthetic_case(case, tmp_path)
    historical = tmp_path / "historical.xlsx"
    historical.write_bytes(b"frozen historical bytes")
    frozen = contract.PriceBaseline(
        contract.HISTORICAL.version,
        historical,
        contract.sha256_file(historical),
    )
    monkeypatch.setitem(contract.BASELINES, contract.HISTORICAL.version, frozen)
    case["selector"].write_text("corrupt bootstrap selector", encoding="utf-8")
    assert (
        contract.require_price_baseline(historical, contract.HISTORICAL.version)
        == frozen
    )
