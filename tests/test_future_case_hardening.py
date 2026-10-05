from __future__ import annotations

import json
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import build_price_calculator_input_draft_from_confirmed_composition as builder  # type: ignore[import-not-found]
import future_case_hardening as f  # type: ignore[import-not-found]
import governed_metal_cost as m  # type: ignore[import-not-found]
import pytest
import run_checked_price_calculator_from_completed_draft as runner  # type: ignore[import-not-found]
import validate_completed_price_calculator_input_draft as completed_validator  # type: ignore[import-not-found]
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from price_baseline_contract import (  # type: ignore[import-not-found]
    PriceBaseline,
    sha256_file,
)
from test_build_price_calculator_input_draft_from_confirmed_composition import (
    valid_data,
)
from test_validate_completed_price_calculator_input_draft import (
    valid_data as legacy_completed_data,
)

CONTEXT = {
    "case_id": "CASE-SYNTHETIC-FUTURE",
    "case_state": "NEW_FUTURE",
    "future_cases_only": True,
    "historical_repricing_authorized": False,
}


def binding(tmp_path: Path, name: str, value: Any) -> dict[str, Any]:
    p = tmp_path / name
    p.write_text(json.dumps({"value": value}), encoding="utf-8")
    return {"path": str(p), "sha256": sha256_file(p), "locator": "/value"}


def source_fixture(
    tmp_path: Path, family: str = "АВР", area: int = 25
) -> tuple[dict[str, Any], dict[str, Any]]:
    source = valid_data()
    source["future_context"] = dict(CONTEXT)
    item = source["items"][0]
    item["item_id"] = "ITEM-1"
    item["product_name"] = "Synthetic " + family
    item["quantity"] = 1
    power = {
        "component_id": "P1",
        "component_code": "SOURCE-MCCB",
        "component_label": "Автоматический выключатель 3P 63А",
        "quantity": 2,
        "install_type": "mccb_up_to_100a",
    }
    control = {
        "component_id": "C1",
        "component_code": "SOURCE-CONTROL",
        "component_label": "Автоматический выключатель 1P 10А",
        "quantity": 1,
        "install_type": "modular_1p",
    }
    contractor = {
        "component_id": "K1",
        "component_code": "SOURCE-CONTACTOR",
        "component_label": "Контактор 65А",
        "quantity": 2,
        "install_type": "contactor",
    }
    fuse = {
        "component_id": "F1",
        "component_code": "SOURCE-PPN",
        "component_label": "ППН 125А",
        "quantity": 1,
        "install_type": "fuse_group_3p",
    }
    if family in {"ВРУ_FUSE", "ШРС"}:
        comps = [fuse]
        roles = {"F1": "FUSE_GROUP"}
        cats = {}
        fuses = {"F1": {"rating_a": 125, "grouping": "РЩж_3PH"}}
        circuits = []
    else:
        comps = [power, control]
        roles = {"P1": "POWER_BREAKER", "C1": "CONTROL_BREAKER"}
        cats = {"P1": "MCCB 3P 16-63A", "C1": "MCB 1P 10A"}
        fuses = {}
        circuits = [
            {
                "circuit_id": "INPUT1",
                "conductor_mm2": area,
                "protection_component_ids": ["P1"],
            }
        ]
    if family == "АВР":
        comps.append(contractor)
        roles["K1"] = "CONTACTOR"
        cats["K1"] = "CONTACTOR 65A"
    item["components"] = comps
    item["technical_classification"] = {
        "family": family,
        "execution": "STANDARD_CONTACTOR_AVR" if family == "АВР" else "STANDARD",
        "circuits": circuits,
        "roles": roles,
        "categories": cats,
        "fuses": fuses,
        "interlocks": {"mechanical": True, "electrical": True},
        "overrides": {},
    }
    source["items"] = [item]
    return save_source(tmp_path, source), source


def save_source(tmp_path: Path, source: dict[str, Any]) -> dict[str, Any]:
    p = tmp_path / "confirmed.json"
    p.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
    return {
        "path": str(p),
        "sha256": sha256_file(p),
        "item_id": source["items"][0]["item_id"],
    }


def governed(binding: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    rules = f.pto_rules(CONTEXT, technical_binding=binding)
    _, item = f.load_technical_binding(CONTEXT, binding)
    return rules, f.bound_bom(CONTEXT, item)


@pytest.mark.parametrize("area,kind", [(24, None), (25, "MCCB"), (26, "MCCB")])
def test_source_bound_threshold(area: int, kind: str | None, tmp_path: Path) -> None:
    bound, _ = source_fixture(tmp_path, area=area)
    rules, bom = governed(bound)
    assert rules["power_breaker_type"] == kind
    assert rules["automation_brand"] == "EKF"
    f.check_avr_bom(rules, bom)  # Control MCB remains distinct and allowed.


@pytest.mark.parametrize(
    "bypass",
    [
        "only_mcb",
        "no_roles",
        "no_circuits",
        "empty_protection",
        "control_role",
        "optional_marker",
        "forged_rules",
    ],
)
def test_conductor25_cannot_bypass(tmp_path: Path, bypass: str) -> None:
    bound, source = source_fixture(tmp_path)
    item = source["items"][0]
    t = item["technical_classification"]
    if bypass == "only_mcb":
        item["components"][0]["install_type"] = "modular_1p"
        t["categories"]["P1"] = "MCB 1P 10A"
    elif bypass == "no_roles":
        t["roles"].pop("P1")
    elif bypass == "no_circuits":
        t["circuits"] = []
    elif bypass == "empty_protection":
        t["circuits"][0]["protection_component_ids"] = []
    elif bypass == "control_role":
        t["roles"]["P1"] = "CONTROL_BREAKER"
    bound = save_source(tmp_path, source)
    with pytest.raises((f.RuleError, KeyError)):
        rules, bom = governed(bound)
        if bypass == "optional_marker":
            bom = [
                {"category": "MCB 1P 10A", "quantity": 1}
            ]  # No power_breaker marker.
        if bypass == "forged_rules":
            rules["power_breaker_type"] = None
        f.check_avr_bom(rules, bom)


@pytest.mark.parametrize(
    "bypass",
    [
        "missing_execution",
        "wrong_execution",
        "contradictory_breaker_execution",
        "naked_false",
        "mechanical_false",
        "electrical_false",
        "priced_mechanical",
    ],
)
def test_avr_rule_not_optional(tmp_path: Path, bypass: str) -> None:
    bound, source = source_fixture(tmp_path)
    t = source["items"][0]["technical_classification"]
    if bypass == "missing_execution":
        t.pop("execution")
    elif bypass == "wrong_execution":
        t["execution"] = "UNKNOWN"
    elif bypass == "contradictory_breaker_execution":
        t["execution"] = "BREAKER_AVR"
    elif bypass.endswith("_false") and bypass != "naked_false":
        t["interlocks"][bypass.split("_")[0]] = False
    bound = save_source(tmp_path, source)
    with pytest.raises((f.RuleError, KeyError)):
        rules, bom = governed(bound)
        if bypass == "naked_false":
            f.pto_rules(CONTEXT, technical_binding=bound, standard_contactor_avr=False)
        if bypass == "priced_mechanical":
            bom.append({"category": "MECHANICAL_INTERLOCK", "quantity": 1})
        f.check_avr_bom(rules, bom)


def test_family_and_source_override(tmp_path: Path) -> None:
    bound, source = source_fixture(tmp_path)
    with pytest.raises(f.RuleError):
        f.pto_rules(CONTEXT, family="КРН", technical_binding=bound)
    with pytest.raises(f.RuleError):
        f.pto_rules(CONTEXT, family="АВР")
    human = binding(tmp_path, "human.json", "EKF")
    project = binding(tmp_path, "project.json", "CHINT")
    source["items"][0]["technical_classification"]["overrides"] = {
        "human": {"automation_brand": human},
        "project": {"automation_brand": project},
    }
    bound = save_source(tmp_path, source)
    assert governed(bound)[0]["automation_brand"] == "EKF"
    source["items"][0]["technical_classification"]["overrides"].pop("human")
    bound = save_source(tmp_path, source)
    assert governed(bound)[0]["automation_brand"] == "CHINT"
    Path(project["path"]).write_text("{}", encoding="utf-8")
    with pytest.raises(f.RuleError, match="drift"):
        governed(bound)


@pytest.mark.parametrize("role", ["human", "project", "source"])
@pytest.mark.parametrize("area", [25, 26])
@pytest.mark.parametrize("attempt", ["MCB", None, "MCCB"])
def test_power_protection_has_no_generic_override(
    tmp_path: Path, role: str, area: int, attempt: Any
) -> None:
    _, source = source_fixture(tmp_path, area=area)
    source["items"][0]["technical_classification"]["overrides"] = {
        role: {"power_breaker_type": binding(tmp_path, "override.json", attempt)}
    }
    with pytest.raises(f.RuleError, match="unsupported override field"):
        governed(save_source(tmp_path, source))


@pytest.mark.parametrize(
    "fault", ["role", "field", "record_extra", "record_typo", "role_type", "shadowed"]
)
def test_nested_override_fields_fail_closed(tmp_path: Path, fault: str) -> None:
    _, source = source_fixture(tmp_path)
    record = binding(tmp_path, "override.json", "EKF")
    overrides: dict[str, Any] = {"source": {"automation_brand": record}}
    if fault == "role":
        overrides = {"sorce": {"automation_brand": record}}
    elif fault == "field":
        overrides = {"source": {"automation_brad": record}}
    elif fault == "record_extra":
        record["unexpected"] = True
    elif fault == "record_typo":
        record["locatr"] = record.pop("locator")
    elif fault == "role_type":
        overrides = {"source": []}
    elif fault == "shadowed":
        overrides = {
            "human": {"automation_brand": record},
            "source": {"automation_brad": record},
        }
    source["items"][0]["technical_classification"]["overrides"] = overrides
    with pytest.raises(f.RuleError, match="override"):
        governed(save_source(tmp_path, source))


@pytest.mark.parametrize("role", ["human", "project", "source"])
def test_supported_brand_enclosure_overrides(tmp_path: Path, role: str) -> None:
    _, source = source_fixture(tmp_path, "ВРУ_BREAKER")
    source["items"][0]["technical_classification"]["overrides"] = {
        role: {
            "automation_brand": binding(tmp_path, "brand.json", "CHINT"),
            "enclosure_mm": binding(tmp_path, "enclosure.json", [1800, 900, 600]),
        }
    }
    rules, bom = governed(save_source(tmp_path, source))
    assert rules["automation_brand"] == "CHINT"
    assert rules["enclosure_mm"] == [1800, 900, 600]
    assert rules["power_breaker_type"] == "MCCB"
    f.check_avr_bom(rules, bom)


def test_below25_allows_proven_power_mcb(tmp_path: Path) -> None:
    _, source = source_fixture(tmp_path, area=24)
    item = source["items"][0]
    item["components"][0]["install_type"] = "modular_3p"
    item["technical_classification"]["categories"]["P1"] = "MCB 3P 10A"
    rules, bom = governed(save_source(tmp_path, source))
    assert rules["power_breaker_type"] is None
    f.check_avr_bom(rules, bom)


@pytest.mark.parametrize("claim", [None, True, False])
def test_future_k_does_not_claim_separate_human_approval(
    tmp_path: Path, claim: bool | None
) -> None:
    bound, _ = source_fixture(tmp_path)
    path = tmp_path / "draft.json"
    assert (
        builder.build_price_calculator_input_draft(Path(bound["path"]), path).status
        == "PASS"
    )
    data = json.loads(path.read_bytes())
    data["operator_completion"] = {
        "completed_by": "Synthetic operator",
        "completed_at": "2026-10-05T10:00:00+05:00",
        "completion_note": "Technical completion; K is governed policy",
    }
    if claim is not None:
        data["operator_completion"]["consumables_factor_confirmed_by_igor"] = claim
    path.write_text(json.dumps(data), encoding="utf-8")
    result = completed_validator.validate_completed_price_calculator_input_draft(path)
    assert result.future_bound
    assert result.status == ("PASS" if claim is None else "FAIL"), result.red_flags


@pytest.mark.parametrize("claim", [None, True, False])
def test_legacy_k_approval_contract_unchanged(
    tmp_path: Path, claim: bool | None
) -> None:
    data = legacy_completed_data()
    if claim is None:
        data["operator_completion"].pop("consumables_factor_confirmed_by_igor")
    else:
        data["operator_completion"]["consumables_factor_confirmed_by_igor"] = claim
    path = tmp_path / "legacy.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    result = completed_validator.validate_completed_price_calculator_input_draft(path)
    assert not result.future_bound
    assert result.status == ("PASS" if claim is True else "FAIL"), result.red_flags


@pytest.mark.parametrize(
    "rating,band,footprint",
    [
        (1, 100, "PPN-33"),
        (100, 100, "PPN-33"),
        (101, 250, "PPN-33"),
        (125, 250, "PPN-33"),
        (160, 250, "PPN-33"),
        (161, 250, None),
        (250, 250, None),
        (251, 400, None),
        (400, 400, None),
    ],
)
def test_fuse_boundaries(rating: int, band: int, footprint: str | None) -> None:
    value = f.fuse_rule(CONTEXT, rating, three_phase_rsh=True)
    assert (
        value["pricing_category"] == f"PN-2 {band}A"
        and value["quantity_per_group"] == 3
        and value["physical_footprint"] == footprint
    )


@pytest.mark.parametrize("rating", [0, -1, 401, True, "NaN"])
def test_fuse_range_hold(rating: Any) -> None:
    with pytest.raises(f.RuleError):
        f.fuse_rule(CONTEXT, rating)


def test_frozen_and_exact_designation() -> None:
    for cid in ["Invoice519", "CASE-LABORATORY-VRU-20260924-001"]:
        with pytest.raises(f.RuleError):
            f.future_only({**CONTEXT, "case_id": cid})
    assert f.designation_category("КЗ-1000В 0,47 мкФ") == "Разрядник"
    with pytest.raises(f.RuleError):
        f.designation_category("КЗ-1000В 0,50 мкФ")


def price_chain(tmp_path: Path, monkeypatch: Any) -> Path:
    p = tmp_path / "prices.xlsx"
    w = Workbook()
    w.remove(w.active)
    for sheet, row, label in f.CATEGORIES.values():
        s = w[sheet] if sheet in w.sheetnames else w.create_sheet(sheet)
        s.cell(row, 1, label)
        s.cell(row, 2, 2100)
    w.save(p)
    w.close()
    baseline = PriceBaseline("active_approved", p, sha256_file(p))
    monkeypatch.setattr(f, "resolve_active_price_baseline", lambda _: baseline)
    monkeypatch.setattr(
        runner, "resolve_price_baseline", lambda *_args, **_kwargs: baseline
    )
    monkeypatch.setattr(
        runner, "require_price_baseline", lambda *_args, **_kwargs: baseline
    )
    return p


@pytest.mark.parametrize("family,k", list(f.K_BY_FAMILY.items()))
def test_existing_composition_to_checked_price_path(
    tmp_path: Path, monkeypatch: Any, family: str, k: Decimal
) -> None:
    bound, _ = source_fixture(tmp_path, family)
    source_path = Path(bound["path"])
    source_before = source_path.read_bytes()
    draft_path = tmp_path / "draft.json"
    built = builder.build_price_calculator_input_draft(source_path, draft_path)
    assert built.status == "PASS", built.red_flags
    data = json.loads(draft_path.read_bytes())
    data["operator_completion"] = {
        "completed_by": "Synthetic operator",
        "completed_at": "2026-10-02T10:00:00+05:00",
        "completion_note": "Governed source review",
    }
    draft_path.write_text(json.dumps(data), encoding="utf-8")
    prices = price_chain(tmp_path, monkeypatch)
    cabinet = {**binding(tmp_path, "cost.json", 10000), "role": "ENCLOSURE_COST"}
    work = binding(tmp_path, "work.json", 6000)
    costs_path = tmp_path / "inputs.json"
    costs_path.write_text(
        json.dumps(
            {
                "case_id": CONTEXT["case_id"],
                "items": [
                    {"item_id": "ITEM-1", "cabinet_cost": cabinet, "work_source": work}
                ],
            }
        ),
        encoding="utf-8",
    )
    run = runner.run_checked_price_calculator_from_completed_draft(
        draft_path,
        prices,
        active_selector_path=tmp_path / "selector.json",
        future_cost_inputs=costs_path,
        expected_future_cost_inputs_sha256=sha256_file(costs_path),
    )
    assert run.status == "PASS", run.red_flags
    applied = run.future_rule_applications[0]
    assert applied["price"]["K"] == str(k)
    material = sum(Decimal(c["quantity"]) * 2100 for c in applied["bom"])
    expected = int(
        ((10000 + material * k + 6000) * Decimal("1.25") * Decimal("1.15")).quantize(
            Decimal(1), rounding=ROUND_HALF_UP
        )
    )
    assert run.overall_preliminary_total == expected
    if family in {"ВРУ_FUSE", "ШРС"}:
        assert (
            applied["bom"][0]["quantity"] == "3"
            and applied["bom"][0]["category"] == "PN-2 250A"
            and applied["bom"][0]["physical_footprint"] == "PPN-33"
        )
    assert source_path.read_bytes() == source_before
    data["calculator_input_format"]["rows"][0]["consumables_factor"] = 9
    draft_path.write_text(json.dumps(data), encoding="utf-8")
    bad = runner.run_checked_price_calculator_from_completed_draft(
        draft_path,
        prices,
        active_selector_path=tmp_path / "selector.json",
        future_cost_inputs=costs_path,
        expected_future_cost_inputs_sha256=sha256_file(costs_path),
    )
    assert bad.status == "FAIL"


def metal_fixture(tmp_path: Path, name: str = "metal.xlsx") -> Path:
    p = tmp_path / name
    w = Workbook()
    s = w.active
    s.title = "Лист1"
    s["A52"] = m.LABEL
    for cell, value in {
        **m.FORMULAS,
        **m.IDENTITY_VALUES,
        "B2": 530,
        "B3": 3000,
        "B52": "=ROUNDUP(D52*$B$5,0)",
        "B5": 1.4375,
    }.items():
        s[cell] = value
    w.save(p)
    w.close()
    return p


def test_metal_first_successor_has_separate_authority(
    tmp_path: Path, monkeypatch: Any
) -> None:
    previous = metal_fixture(tmp_path)
    original = previous.read_bytes()
    sha = sha256_file(previous)
    monkeypatch.setattr(m, "GENESIS_SHA", sha)
    candidate = tmp_path / "candidate.xlsx"
    candidate.write_bytes(original)
    w = load_workbook(candidate)
    w.active["B2"] = 600
    w.save(candidate)
    w.close()
    before = {p.name for p in tmp_path.iterdir()}
    audit = m.audit_metal_candidate(
        candidate,
        sha256_file(candidate),
        predecessor_source=previous,
        expected_predecessor_sha256=sha,
    )
    assert (
        audit["predecessor"]["workbook_sha256"] == sha
        and audit["workbook_sha256"] != sha
    )
    assert audit["cost_changes"] and audit["price_input_changes"][0]["cell"] == "B2"
    assert len(audit["approval_fingerprint"]) == 64 and audit["status"].startswith(
        "HOLD"
    )
    assert (
        before == {p.name for p in tmp_path.iterdir()}
        and previous.read_bytes() == original
    )
    assert audit == m.audit_metal_candidate(
        candidate,
        sha256_file(candidate),
        predecessor_source=previous,
        expected_predecessor_sha256=sha,
    )


@pytest.mark.parametrize(
    "cell,value",
    [
        ("D52", "=B52"),
        ("A52", "OTHER"),
        ("B1", 1.2),
        ("B2", 0),
        ("D82", "=123"),
        ("B5", 2),
    ],
)
def test_metal_drift_hold(
    tmp_path: Path, monkeypatch: Any, cell: str, value: Any
) -> None:
    previous = metal_fixture(tmp_path)
    sha = sha256_file(previous)
    monkeypatch.setattr(m, "GENESIS_SHA", sha)
    candidate = tmp_path / "candidate.xlsx"
    candidate.write_bytes(previous.read_bytes())
    w = load_workbook(candidate)
    w.active[cell] = value
    w.save(candidate)
    w.close()
    with pytest.raises(f.RuleError):
        m.audit_metal_candidate(
            candidate,
            sha256_file(candidate),
            predecessor_source=previous,
            expected_predecessor_sha256=sha,
        )
