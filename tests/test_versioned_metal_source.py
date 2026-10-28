"""Public successor seams; synthetic sources grant no production authority."""

import copy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import governed_metal_cost as metal  # type: ignore[import-not-found]
import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]


@pytest.fixture
def synthetic_sources(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    from wu8_governed_sources import (  # type: ignore[import-not-found]
        METAL_SIZES,
        metal_identity,
    )

    registry = {}
    for version, rate in (("2026.06.18", 100), ("2026.10.21", 200)):
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Лист1"
        sheet["B2"], sheet["B3"] = rate, 50
        for size, (row, fabrication, painting, extra) in METAL_SIZES.items():
            suffix = {52: " ВРУ ", 53: " ПР ", 54: " ПР ", 51: ""}.get(row, " ")
            identity = metal_identity(
                row,
                "х".join(map(str, size)) + suffix,
                size,
                fabrication,
                painting,
                extra,
            )
            sheet[f"A{row}"] = identity["label"]
            for cell, value in {
                **identity["formulas"],
                **identity["identity_values"],
            }.items():
                sheet[cell] = value
            sheet[f"B{row}"] = "SALES MUST NOT BE READ"
            sheet[f"C{row}"] = "SALES MUST NOT BE READ"
        path = tmp_path / (version + ".xlsx")
        book.save(path)
        book.close()
        registry[version] = {
            "version": version,
            "path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "predecessor_version": None if version == "2026.06.18" else "2026.06.18",
        }
    monkeypatch.setattr(metal, "METAL_SOURCE_VERSIONS", registry)
    return registry


@pytest.mark.skipif(
    not metal.GENESIS_PATH.with_name(
        "прайс_металл_лотки_крышки с 2026.10.21.xlsx"
    ).exists(),
    reason="Pinned local approved workbooks are not distributed",
)
def test_current_source_is_concrete_and_historical_source_stays_pinned() -> None:
    assert hasattr(metal, "resolve_metal_source"), "public version resolver missing"
    current = metal.resolve_metal_source("CURRENT")
    assert current["version"] == "2026.10.21"
    assert current["sha256"] == (
        "be4a3497cf86ee4e5a13ac55bfeeeb980e6948cd30847ae027f36d6ca599cc71"
    )
    assert current["predecessor_version"] == "2026.06.18"
    historical = metal.resolve_metal_source("2026.06.18")
    assert historical["path"] == str(metal.GENESIS_PATH)
    assert historical["sha256"] == metal.GENESIS_SHA
    assert Path(current["path"]).name.endswith("2026.10.21.xlsx")


@pytest.fixture
def synthetic_catalogue(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    synthetic_sources: Any,
) -> Any:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]

    for version, source in synthetic_sources.items():
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Лист1"
        for cell, value in catalogue.HEADERS.items():
            sheet[cell] = value
        for _, first, last in catalogue.RANGES:
            for row in range(first, last + 1):
                size = [700, 500, 200] if row == 34 else [540, 330, 100]
                sheet[f"A{row}"] = "х".join(map(str, size))
                for axis, axis_value in zip("EFG", size, strict=True):
                    sheet[f"{axis}{row}"] = axis_value
                sheet[f"B{row}"] = "SALES NOT AUTHORITY"
        sheet["A64"] = "Корпус КРН-36 540х330х100 "
        if version == "2026.06.18":
            sheet["G34"] = "=400"
        path = Path(source["path"])
        book.save(path)
        book.close()
        source["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    active = tmp_path / "main.xlsx"
    active.write_bytes(b"synthetic main identity")
    active_sha = hashlib.sha256(active.read_bytes()).hexdigest()
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps({"workbook": {"path": str(active)}}), encoding="utf-8"
    )
    selector = tmp_path / "selector.json"
    selector.write_text(json.dumps({"manifest_path": str(manifest)}), encoding="utf-8")
    manifest_sha = hashlib.sha256(manifest.read_bytes()).hexdigest()
    monkeypatch.setattr(catalogue, "DEFAULT_ACTIVE_SELECTOR", selector)
    monkeypatch.setattr(
        catalogue, "SELECTOR_SHA", hashlib.sha256(selector.read_bytes()).hexdigest()
    )
    monkeypatch.setattr(catalogue, "MANIFEST_SHA", manifest_sha)
    monkeypatch.setattr(catalogue, "ACTIVE_SHA", active_sha)
    baseline = SimpleNamespace(
        sha256=active_sha,
        path=active,
        manifest_sha256=manifest_sha,
        manifest_path=manifest,
        mapping_snapshot=[
            {
                "mapping_id": "CABINET-PRICE-KRN-36",
                "identity": {
                    "authority": "TECHNICAL_MAPPING_IDENTITY",
                    "cabinet_code": "КРН-36",
                    "height_mm": 330,
                    "width_mm": 540,
                    "depth_mm": 100,
                    "material": "metal",
                    "expected_label": "Корпус КРН-36 540х330х100",
                    "mapping_id": "CABINET-PRICE-KRN-36",
                    "row": 9,
                    "sheet": "КРН",
                    "strict_label": False,
                },
                "identity_fingerprint": (
                    "4eda337c68656d43d1ae64b378dffd20949aa06b5f940a6dba32f8fa1ce859d1"
                ),
                "source": {
                    "label_cell": "КРН!L9",
                    "price_cells": ["КРН!M9"],
                    "row": 9,
                    "sheet": "КРН",
                },
            }
        ],
    )
    monkeypatch.setattr(catalogue, "resolve_active_price_baseline", lambda: baseline)
    monkeypatch.setattr(
        catalogue,
        "KRN36_METAL_SHA",
        synthetic_sources["2026.10.21"]["sha256"],
        raising=False,
    )
    return baseline, synthetic_sources


def test_synthetic_current_inventory_parses_all_rows_without_financial_reads(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    synthetic_catalogue: Any,
) -> None:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]
    import cabinet_sizing as sizing  # type: ignore[import-not-found]

    original = catalogue.load_workbook

    class Sheet:
        def __init__(self, inner: Any):
            self.inner = inner

        def __getitem__(self, cell: str) -> Any:
            assert cell[0] in "AEFG", "catalogue read financial data"
            return self.inner[cell]

    class Book:
        def __init__(self, inner: Any):
            self.inner = inner

        def __getitem__(self, name: str) -> Any:
            return Sheet(self.inner[name])

        def close(self) -> None:
            self.inner.close()

    monkeypatch.setattr(
        catalogue, "load_workbook", lambda *a, **kw: Book(original(*a, **kw))
    )
    old = catalogue.inventory(source_version="2026.06.18")
    new = catalogue.inventory(source_version="2026.10.21")
    assert len(old["records"]) == len(new["records"]) == 70
    before = next(r for r in old["records"] if r["row"] == 34)
    after = next(r for r in new["records"] if r["row"] == 34)
    assert before["source_issues"] == ["HOLD_LABEL_DIMENSION_SOURCE_CONFLICT"]
    assert after["dimensions_mm"] == [700, 500, 200] and after["source_issues"] == []
    assert next(r for r in new["records"] if r["row"] == 64)["source_issues"] == [
        "HOLD_ACTIVE_MAPPING_AXIS_CONFLICT"
    ]
    assert all(r["source"]["version"] == "2026.10.21" for r in new["records"])
    monkeypatch.setattr(catalogue, "HEADERS", dict(catalogue.HEADERS, E7="Wrong axis"))
    with pytest.raises(sizing.SizingHold, match="HOLD_CATALOGUE_HEADERS"):
        catalogue.inventory(source_version="2026.10.21")
    monkeypatch.setattr(Path, "is_symlink", lambda _: True)
    with pytest.raises(ValueError, match="reparse"):
        metal.resolve_metal_source("2026.10.21")


def test_ten_versioned_costs_execute_identity_formula_without_sales(
    synthetic_sources: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from wu8_governed_sources import METAL_SIZES  # type: ignore[import-not-found]

    assert hasattr(
        metal, "resolve_versioned_metal_cost"
    ), "public versioned cost missing"
    original = metal.load_workbook
    accesses = []

    class Sheet:
        def __init__(self, inner: Any):
            self.inner = inner

        def __getitem__(self, cell: str) -> Any:
            accesses.append(cell)
            assert cell in {"B1", "B2", "B3"} or cell[0] not in "BC"
            assert cell != "D34"
            return self.inner[cell]

    class Book:
        def __init__(self, inner: Any):
            self.inner = inner
            self.sheetnames = inner.sheetnames

        def __getitem__(self, name: str) -> Any:
            return Sheet(self.inner[name])

        def close(self) -> None:
            self.inner.close()

    monkeypatch.setattr(
        metal, "load_workbook", lambda *a, **kw: Book(original(*a, **kw))
    )
    rows = []
    for size in METAL_SIZES:
        old = metal.resolve_versioned_metal_cost(size, source_version="2026.06.18")
        new = metal.resolve_versioned_metal_cost(size)
        assert old["cost_kzt"] < new["cost_kzt"]
        assert old["identity"] == new["identity"]
        assert new["metal_source"] == synthetic_sources["2026.10.21"]
        rows.append(new["source_locator"])
    assert set(rows) == {
        f"Лист1!D{r}" for r in (22, 28, 35, 42, 44, 46, 51, 52, 53, 54)
    }
    # ceil((1.05*8.42*200 + 1.05*2*.25*50*1.33)*1.05 + 5650)
    assert metal.resolve_versioned_metal_cost((500, 400, 250))["cost_kzt"] == 7544
    with pytest.raises(ValueError, match="ten-size"):
        metal.resolve_versioned_metal_cost((700, 500, 200))
    assert accesses
    with pytest.raises(ValueError, match="unknown metal source"):
        metal.resolve_metal_source("unknown")
    source_path = Path(synthetic_sources["2026.10.21"]["path"])
    raw = source_path.read_bytes()
    original_sha = synthetic_sources["2026.10.21"]["sha256"]
    altered = original(source_path)
    altered["Лист1"]["D22"] = "=1"
    altered.save(source_path)
    altered.close()
    synthetic_sources["2026.10.21"]["sha256"] = hashlib.sha256(
        source_path.read_bytes()
    ).hexdigest()
    with pytest.raises(ValueError, match="formula drift: D22"):
        metal.resolve_versioned_metal_cost((500, 400, 250))
    source_path.write_bytes(raw)
    synthetic_sources["2026.10.21"]["sha256"] = original_sha
    source_path.write_bytes(raw + b"drift")
    with pytest.raises(ValueError, match="SHA drift"):
        metal.resolve_versioned_metal_cost((500, 400, 250))


def test_technical_krn36_binding_preserves_legacy_and_other_records(
    synthetic_catalogue: Any,
) -> None:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]

    historical = catalogue.inventory(source_version="2026.10.21")
    current = catalogue.inventory(source_version="2026.10.21", technical_binding=True)
    for old, new in zip(historical["records"], current["records"], strict=True):
        if old["row"] != 64:
            assert old == new
            continue
        assert new["dimensions_mm"] == [540, 330, 100]
        assert new["active_mapping"] == old["active_mapping"]
        assert old["source_issues"] == ["HOLD_ACTIVE_MAPPING_AXIS_CONFLICT"]
        assert new["source_issues"] == []
        binding = new["technical_binding"]
        assert binding["rail_count"] == 3
        assert binding["dimensions_mm"] == [540, 330, 100]
        assert binding["metal_source"] == new["source"]
        assert binding["wu8_source"]["key"] == "CAB-KRN-36"
        assert binding["legacy_price_identity"]["identity"]["height_mm"] == 330
        assert binding["legacy_price_identity"]["physical_authority"] is False
        assert binding["human_decision"]["authority"] == "IGOR_DIRECT_HUMAN_DECISION"
        assert not {"zones", "usable_mm", "ip", "installation"} & binding.keys()
    assert catalogue.MOUNTING_SOURCES == {}


def test_v4_capacity_binding_is_additive_and_uses_exact_inherited_authority(
    synthetic_catalogue: Any,
) -> None:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]

    v3 = catalogue.inventory(source_version="2026.10.21", technical_binding=True)
    v4 = catalogue.inventory(
        source_version="2026.10.21", technical_binding=True, capacity_binding=True
    )
    assert len(v4["records"]) == 70
    for old, new in zip(v3["records"], v4["records"], strict=True):
        if old["row"] != 64:
            assert old == new
        else:
            capacity = new["din_capacity_binding"]
            assert capacity["rail_count"] == 3
            assert capacity["modules_per_rail"] == 12
            assert capacity["module_pitch_mm"] == 18
            assert capacity["rail_capacity_mm"] == 216
            assert capacity["total_capacity_mm"] == 648
            assert capacity["capacity_only"] is True
            assert (
                capacity["human_decision"]["authority"] == "IGOR_DIRECT_HUMAN_DECISION"
            )
            assert {k: v for k, v in new.items() if k != "din_capacity_binding"} == old


@pytest.mark.parametrize(
    "fault",
    [
        "dimensions",
        "label",
        "sha",
        "legacy",
        "fingerprint",
        "locator",
        "wu8",
        "wu8_sha",
        "version",
    ],
)
def test_technical_krn36_binding_rejects_mixed_or_drifted_evidence(
    synthetic_catalogue: Any, monkeypatch: pytest.MonkeyPatch, fault: str
) -> None:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]
    import cabinet_sizing as sizing  # type: ignore[import-not-found]
    import wu8_governed_sources as wu8  # type: ignore[import-not-found]

    baseline, registry = synthetic_catalogue
    version = "2026.10.21"
    if fault in {"dimensions", "label"}:
        path = Path(registry[version]["path"])
        book = metal.load_workbook(path)
        book["Лист1"]["E64" if fault == "dimensions" else "A64"] = (
            541 if fault == "dimensions" else "Корпус КРН-48 540х330х100 "
        )
        book.save(path)
        book.close()
        registry[version]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        monkeypatch.setattr(
            catalogue, "KRN36_METAL_SHA", registry[version]["sha256"], raising=False
        )
    elif fault == "sha":
        monkeypatch.setattr(catalogue, "KRN36_METAL_SHA", "0" * 64, raising=False)
    elif fault == "legacy":
        baseline.mapping_snapshot[0]["identity"]["height_mm"] = 540
    elif fault == "fingerprint":
        baseline.mapping_snapshot[0]["identity_fingerprint"] = "0" * 64
    elif fault == "locator":
        baseline.mapping_snapshot[0]["source"]["label_cell"] = "КРН!L10"
    elif fault == "wu8":
        monkeypatch.setitem(
            wu8.CABINETS, "CAB-KRN-36", ("КРН-36", 9, (330, 540, 100), 64, 1500, 350)
        )
    elif fault == "wu8_sha":
        monkeypatch.setattr(catalogue, "WU8_SHA", "0" * 64, raising=False)
    else:
        version = "2026.06.18"
    with pytest.raises(sizing.SizingHold, match="HOLD_.*(KRN36|DRIFT)"):
        catalogue.inventory(source_version=version, technical_binding=True)


def test_current_catalogue_schema_pins_source_and_replays_history(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    synthetic_sources: Any,
    capsys: pytest.CaptureFixture[str],
) -> None:
    import cabinet_catalogue as catalogue  # type: ignore[import-not-found]
    import cabinet_sizing as sizing  # type: ignore[import-not-found]
    import cabinet_sizing_catalogue as selector  # type: ignore[import-not-found]
    from test_cabinet_sizing_catalogue import record
    from test_cabinet_sizing_pragmatic import provision

    item, legacy = provision(tmp_path, monkeypatch)
    legacy["schema_version"] = selector.REQUEST_SCHEMA
    assert hasattr(
        selector, "CURRENT_REQUEST_SCHEMA"
    ), "current catalogue schema missing"
    current = dict(
        legacy,
        schema_version="cabinet_sizing.catalogue_request.v0.2",
        metal_source=metal.resolve_metal_source(),
    )
    calls = []

    def inventory(*, source_version: str | None = None) -> Any:
        calls.append(source_version)
        result: dict[str, Any] = {"records": [record("SMALL")], "policy": "SYNTHETIC"}
        if source_version is not None:
            result["metal_source"] = metal.resolve_metal_source(source_version)
        return result

    monkeypatch.setattr(catalogue, "inventory", inventory)
    old = sizing.size_composition("CASE", [item], legacy)
    report = sizing.size_composition("CASE", [item], current)
    assert report["schema_version"] == "cabinet_sizing.catalogue.v0.2"
    assert report["sources"]["metal_source"] == synthetic_sources["2026.10.21"]
    assert "2026.10.21" in calls and None in calls
    item["cabinet"] = report["items"][0]["cabinet"]
    sizing.verify_report(
        report, [item], expected_case_id="CASE", before_human_approval_1=True
    )
    monkeypatch.setattr(metal, "CURRENT_VERSION", "2026.06.18")
    assert sizing.size_composition("CASE", [item], current) == report
    assert sizing.size_composition("CASE", [item], legacy) == old
    for key, value in (
        ("version", "CURRENT"),
        ("version", "unknown"),
        ("path", synthetic_sources["2026.06.18"]["path"]),
        ("sha256", "0" * 64),
    ):
        damaged = copy.deepcopy(current)
        damaged["metal_source"][key] = value
        with pytest.raises(sizing.SizingHold, match="HOLD_METAL_SOURCE"):
            sizing.size_composition("CASE", [item], damaged)
    composition = tmp_path / "composition.json"
    request_path = tmp_path / "request.json"
    composition.write_text(json.dumps({"items": [item]}), encoding="utf-8")
    request_path.write_text(json.dumps(current), encoding="utf-8")
    assert (
        sizing.main(
            [
                "--composition-json",
                str(composition),
                "--case-id",
                "CASE",
                "--request-json",
                str(request_path),
                "--request-sha256",
                hashlib.sha256(request_path.read_bytes()).hexdigest(),
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["results"][0]["report"] == report
    monkeypatch.setattr(selector, "request_for", lambda *a: current)
    assert sizing.size_composition("CASE", [item]) == report
    source_path = Path(current["metal_source"]["path"])
    source_path.write_bytes(source_path.read_bytes() + b"drift")
    with pytest.raises(sizing.SizingHold, match="HOLD_METAL_SOURCE"):
        sizing.verify_report(
            report, [item], expected_case_id="CASE", before_human_approval_1=True
        )
